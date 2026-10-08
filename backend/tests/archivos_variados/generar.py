"""Banco de archivos variados para probar los importadores sin archivos reales (spec v2.2 · 4.5).

Cada archivo imita un desorden real de un negocio pequeño. Las cifras esperadas
están calculadas A MANO en `tests/test_archivos_variados.py`, no con el motor:
si el motor se equivoca, la prueba lo atrapa.

Uso:  python tests/archivos_variados/generar.py
Los archivos quedan en esta misma carpeta. Es reproducible: no hay azar.
"""
from __future__ import annotations

import csv
import io
from datetime import date
from pathlib import Path

from openpyxl import Workbook

AQUI = Path(__file__).resolve().parent


def _libro(hojas: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _guardar(nombre: str, datos: bytes) -> None:
    (AQUI / nombre).write_bytes(datos)


# 1 · listas mensuales en bloques, encabezados en distinto orden y con otros nombres
def ventas_compras_bloques() -> bytes:
    return _libro({
        "VENTAS": [
            ["VENTAS ENERO 2026"],
            ["FECHA", "CLIENTE", "PRODUCTO", "CANTIDAD", "VALOR UNITARIO", "TOTAL"],
            [date(2026, 1, 10), "JUAN PEREZ", "ARROZ", 10, 3000, 30000],
            [date(2026, 1, 20), "ANA GOMEZ", "FRIJOL", 5, 8000, 40000],
            ["TOTAL", None, None, None, None, 70000],
            [],
            ["VENTAS FEBRERO 2026"],
            ["Cliente", "Fecha", "Total", "Producto", "Cant.", "V/R Unit"],
            ["LUIS DIAZ", date(2026, 2, 5), 45000, "ARROZ", 15, 3000],
        ],
        "COMPRAS": [
            ["COMPRAS ENERO 2026"],
            ["FECHA", "PROVEEDOR", "PRODUCTO", "CANTIDAD", "COSTO UNITARIO", "VALOR TOTAL"],
            [date(2026, 1, 2), "DISTRIBUIDORA SUR", "ARROZ", 40, 2000, 80000],
            [date(2026, 1, 3), "GRANOS SA", "FRIJOL", 10, 5000, 50000],
        ],
    })


# 2 · columnas mal rotuladas: «REFERENCIA» trae la cantidad; en el otro bloque
#     «CANTIDAD» y «VALOR UNITARIO» están cruzadas.
def columnas_mal_rotuladas() -> bytes:
    return _libro({
        "VENTAS": [
            ["VENTAS MARZO 2026"],
            ["FECHA", "CLIENTE", "PRODUCTO", "REFERENCIA", "PRECIO", "TOTAL"],
            [date(2026, 3, 3), "MARIA", "QUESO", 2, 12000, 24000],
            [date(2026, 3, 4), "PEDRO", "QUESO", 3, 12000, 36000],
            [date(2026, 3, 9), "MARIA", "LECHE", 4, 2500, 10000],
            [],
            ["VENTAS ABRIL 2026"],
            ["FECHA", "CLIENTE", "PRODUCTO", "CANTIDAD", "VALOR UNITARIO", "TOTAL"],
            [date(2026, 4, 2), "MARIA", "QUESO", 12500, 2, 25000],
            [date(2026, 4, 6), "PEDRO", "LECHE", 2600, 5, 13000],
        ],
    })


# 3 · fechas copiadas del mes anterior y una fecha futura (día y mes cruzados)
def fechas_copiadas() -> bytes:
    return _libro({
        "VENTAS": [
            ["VENTAS ENERO 2026"],
            ["FECHA", "CLIENTE", "DETALLE", "VALOR"],
            [date(2026, 1, 5), "CLIENTE A", "SERVICIO", 100000],
            [date(2026, 1, 15), "CLIENTE B", "SERVICIO", 50000],
            [],
            ["VENTAS FEBRERO 2026"],
            ["FECHA", "CLIENTE", "DETALLE", "VALOR"],
            [date(2026, 1, 5), "CLIENTE C", "SERVICIO", 70000],
            [date(2026, 1, 15), "CLIENTE D", "SERVICIO", 30000],
            [date(2026, 2, 20), "CLIENTE E", "SERVICIO", 20000],
        ],
        "GASTOS": [
            ["GASTOS DEL TRIMESTRE"],
            ["FECHA", "CONCEPTO", "VALOR"],
            [date(2026, 1, 31), "ARRIENDO", 80000],
            [date(2026, 11, 2), "SERVICIOS PUBLICOS", 15000],
        ],
    })


# 4 · totales mezclados, filas vacías, filas en cero, números como texto con $ y comas
def totales_mezclados() -> bytes:
    return _libro({
        "Hoja1": [
            ["MI NEGOCIO"],
            ["RELACION DE COMPRAS 2026"],
            ["Fecha", "Proveedor", "Concepto", "Valor"],
            [date(2026, 4, 2), "FERRETERIA X", "CEMENTO", "$ 1.200.000"],
            [],
            [date(2026, 4, 10), "FERRETERIA X", "ARENA", "350.000,50"],
            [date(2026, 4, 12), None, None, 0],
            [None, None, None, 1550000.50],
            [date(2026, 4, 20), "DEPOSITO Y", "VARILLA", 800000],
            ["TOTAL", None, None, 2350000.50],
        ],
    })


# 5 · hojas vacías y una hoja que promete compras y no las trae
def hojas_vacias() -> bytes:
    return _libro({
        "Hoja1": [],
        "COMPRAS": [["COMPRAS"]],
        "VENTAS": [
            ["FECHA", "CLIENTE", "PRODUCTO", "VALOR"],
            [date(2026, 5, 2), "X", "PAN", 5000],
            [date(2026, 5, 3), "Y", "PAN", 7000],
        ],
        "Hoja3": [],
    })


# 6 · cartera y cuentas por pagar con terceros que no están en las ventas ni en las compras
def cartera_terceros() -> bytes:
    return _libro({
        "VENTAS": [
            ["FECHA", "CLIENTE", "PRODUCTO", "TOTAL", "FORMA DE PAGO"],
            [date(2026, 6, 1), "TIENDA AZUL", "GASEOSA", 100000, "CREDITO"],
            [date(2026, 6, 2), "JUAN RUIZ", "GASEOSA", 50000, "CONTADO"],
            [date(2026, 6, 3), "JUAN RUIZ", "AGUA", 30000, "CONTADO"],
        ],
        "CARTERA": [
            ["CARTERA JUNIO 2026"],
            ["CLIENTE", "VALOR", "ABONO", "SALDO"],
            ["TIENDA AZUL", 100000, 40000, 60000],
            ["JUAN RUIZ", 30000, 0, 30000],
            ["DOÑA ROSA", 80000, 20000, 60000],
        ],
        "PROVEEDORES": [
            ["CUENTAS POR PAGAR"],
            ["PROVEEDOR", "DEUDA", "ABONOS", "SALDO"],
            ["POSTOBON SA", 200000, 50000, 150000],
        ],
    })


# 7 · CSV de gastos sin NIT ni nombre del negocio
def gastos_csv() -> bytes:
    filas = [["fecha", "concepto", "valor"],
             ["2026-07-01", "Arriendo local", "900.000"],
             ["2026-07-05", "Energía", "120.500"],
             ["2026-07-05", "Agua", "45.300"],
             ["2026-07-10", "Pago empleado", "1.300.000"],
             ["2026-07-15", "Papelería", "23.000"]]
    buf = io.StringIO()
    csv.writer(buf, delimiter=";").writerows(filas)
    return buf.getvalue().encode("utf-8")


# 8 · balance de prueba en PDF con texto (incluye cuentas padre que no deben sumarse)
def balance_pdf() -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    lineas = [
        "COMERCIAL DEMO PDF S.A.S.",
        "BALANCE DE PRUEBA A 31 DE MARZO DE 2026",
        "1105 CAJA 1.500.000",
        "110505 CAJA GENERAL 1.500.000",
        "1435 MERCANCIAS NO FABRICADAS 4.000.000",
        "2205 PROVEEDORES NACIONALES 2.000.000",
        "3105 CAPITAL SUSCRITO Y PAGADO 3.000.000",
        "4135 COMERCIO AL POR MAYOR Y AL POR MENOR 2.500.000",
        "5105 GASTOS DE PERSONAL 1.200.000",
        "6135 COSTO DE VENTAS 800.000",
        "Página 1 de 1",
    ]
    y = 740
    for t in lineas:
        c.drawString(60, y, t)
        y -= 22
    c.showPage()
    c.save()
    return buf.getvalue()


# 8b · listado de ventas en PDF, como tabla con líneas
def ventas_pdf() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    datos = [["FECHA", "CLIENTE", "PRODUCTO", "VALOR"],
             ["2026-09-14", "RESTAURANTE EL PASO", "TOMATE", "25.000"],
             ["2026-09-21", "TIENDA LA 5", "CEBOLLA", "35.000"]]
    t = Table(datos)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([Paragraph("VENTAS DE SEPTIEMBRE", getSampleStyleSheet()["Title"]), t])
    return buf.getvalue()


# 9 · tabla de ventas pegada en un Word, con la identidad en los párrafos
def ventas_word() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("PANADERIA EL TRIGAL S.A.S.")
    d.add_paragraph("NIT 900.555.111-6")
    d.add_paragraph("Relación de ventas de agosto de 2026")
    filas = [["FECHA", "CLIENTE", "PRODUCTO", "CANTIDAD", "PRECIO", "TOTAL"],
             ["2026-08-03", "CAFE LA 14", "PAN", "100", "500", "50.000"],
             ["2026-08-04", "HOTEL SOL", "TORTA", "2", "35.000", "70.000"],
             ["2026-08-05", "CAFE LA 14", "PAN", "80", "500", "40.000"]]
    t = d.add_table(rows=len(filas), cols=len(filas[0]))
    for i, f in enumerate(filas):
        for j, v in enumerate(f):
            t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# 10 · PDF escaneado: solo dibujo, sin una letra de texto
def escaneado_pdf() -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for i in range(12):
        c.rect(60, 700 - i * 40, 480, 26, stroke=1, fill=0)
    c.showPage()
    c.save()
    return buf.getvalue()


# 11 · libro diario suelto en CSV, encabezados en desorden (P01)
def libro_diario_csv() -> bytes:
    filas = [["Descripción", "Crédito", "Cuenta", "Fecha", "Comprobante", "Tercero", "Débito"],
             ["Aporte inicial", "0", "110505", "2026-01-02", "CE-1", "SOCIO UNO", "5.000.000"],
             ["Aporte inicial", "5.000.000", "3105", "2026-01-02", "CE-1", "SOCIO UNO", "0"],
             ["Venta", "0", "110505", "2026-01-15", "FV-1", "CLIENTE A", "1.190.000"],
             ["Venta", "1.000.000", "4135", "2026-01-15", "FV-1", "CLIENTE A", "0"],
             ["IVA venta", "190.000", "240805", "2026-01-15", "FV-1", "CLIENTE A", "0"],
             ["Arriendo", "0", "512010", "2026-02-01", "CE-2", "ARRENDADOR", "600.000"],
             ["Arriendo", "600.000", "110505", "2026-02-01", "CE-2", "ARRENDADOR", "0"]]
    buf = io.StringIO()
    csv.writer(buf, delimiter=";").writerows(filas)
    return buf.getvalue().encode("utf-8")


# 12 · se vende más de lo que se compró, y un producto que se compra por bulto y se vende por kilo
def inventario_insuficiente() -> bytes:
    return _libro({
        "COMPRAS": [
            ["FECHA", "PROVEEDOR", "PRODUCTO", "CANTIDAD", "COSTO UNITARIO", "TOTAL"],
            [date(2026, 9, 1), "PROVEEDOR A", "PAPA", 2, 80000, 160000],
            [date(2026, 9, 1), "PROVEEDOR A", "CEBOLLA", 1, 60000, 60000],
        ],
        "VENTAS": [
            ["FECHA", "CLIENTE", "PRODUCTO", "CANTIDAD", "PRECIO", "TOTAL"],
            [date(2026, 9, 10), "CLIENTE X", "PAPA", 3, 100000, 300000],
            [date(2026, 9, 11), "CLIENTE Y", "CEBOLLA", 1, 2500, 2500],
        ],
    })


# 13 · la forma de pago y el proveedor solo están en la primera fila
def repetidos() -> bytes:
    return _libro({
        "COMPRAS": [
            ["FECHA", "PROVEEDOR", "PRODUCTO", "CANTIDAD", "COSTO UNITARIO", "TOTAL", "FORMA DE PAGO"],
            [date(2026, 3, 1), "DISTRI NORTE", "JABON", 10, 2000, 20000, "CREDITO"],
            [date(2026, 3, 1), None, "DETERGENTE", 5, 6000, 30000, None],
            [date(2026, 3, 1), None, "CLORO", 4, 3000, 12000, None],
            [date(2026, 3, 1), None, "ESCOBA", 2, 5000, 10000, None],
        ],
    })


# 14 · títulos con errores de escritura y la fecha como solo el día
def titulos_mal_escritos() -> bytes:
    return _libro({
        "VENTAS": [
            ["VENTAS ENRO 2O26"],
            ["DIA", "CLIENTE", "VALOR"],
            [5, "UNO", 10000],
            [20, "DOS", 20000],
            [],
            ["VENTAS FEBERO 2026"],
            ["DIA", "CLIENTE", "VALOR"],
            [3, "TRES", 30000],
        ],
    })


# 15 · inventario inicial, ventas y conteo físico al cierre
def inventario_inicial_conteo() -> bytes:
    return _libro({
        "INVENTARIO": [
            ["INVENTARIO INICIAL ENERO 2026"],
            ["PRODUCTO", "CANTIDAD", "COSTO UNITARIO", "TOTAL"],
            ["TORNILLO", 100, 200, 20000],
            ["TUERCA", 50, 100, 5000],
            [],
            ["CONTEO FISICO 31 DE ENERO 2026"],
            ["PRODUCTO", "CANTIDAD"],
            ["TORNILLO", 65],
            ["TUERCA", 50],
        ],
        "VENTAS": [
            ["FECHA", "CLIENTE", "PRODUCTO", "CANTIDAD", "PRECIO", "TOTAL"],
            [date(2026, 1, 10), "C1", "TORNILLO", 30, 500, 15000],
        ],
    })


# 16 · balance en Excel con saldo anterior, débitos, créditos y saldo final, y cuentas padre
def balance_excel() -> bytes:
    return _libro({
        "Balance": [
            ["BALANCE DE PRUEBA - FEBRERO 2026"],
            ["CODIGO", "NOMBRE", "SALDO ANTERIOR", "DEBITOS", "CREDITOS", "SALDO FINAL"],
            ["1", "ACTIVO", 1000000, 500000, 300000, 1200000],
            ["1105", "CAJA", 1000000, 500000, 300000, 1200000],
            ["110505", "CAJA GENERAL", 1000000, 500000, 300000, 1200000],
            ["3105", "CAPITAL", 1000000, 0, 0, 1000000],
            ["4135", "VENTAS", 0, 0, 500000, 500000],
            ["5120", "ARRIENDOS", 0, 300000, 0, 300000],
        ],
    })


# 17 · una hoja por mes, sin títulos, con ventas y compras lado a lado
def hoja_por_mes_lado_a_lado() -> bytes:
    def mes(ventas: list, compras: list) -> list[list]:
        filas = [["FECHA", "CLIENTE", "PRODUCTO", "VALOR", None, "FECHA", "PROVEEDOR", "PRODUCTO", "VALOR"]]
        for i in range(max(len(ventas), len(compras))):
            v = ventas[i] if i < len(ventas) else [None] * 4
            c = compras[i] if i < len(compras) else [None] * 4
            filas.append(v + [None] + c)
        return filas

    return _libro({
        "ENERO": mes([[date(2026, 1, 3), "A", "PAN", 10000], [date(2026, 1, 4), "B", "PAN", 15000]],
                     [[date(2026, 1, 2), "MOLINO", "HARINA", 8000]]),
        "FEBRERO": mes([[date(2026, 2, 3), "C", "PAN", 20000]],
                       [[date(2026, 2, 1), "MOLINO", "HARINA", 5000], [date(2026, 2, 2), "INGENIO", "AZUCAR", 3000]]),
    })


# 18 · listas que no dicen qué son
def listas_sin_tipo() -> bytes:
    return _libro({
        "Hoja1": [
            ["FECHA", "DETALLE", "VALOR"],
            [date(2026, 5, 5), "Arriendo", 500000],
            [date(2026, 5, 6), "Energía", 80000],
        ],
        "Hoja2": [
            ["FECHA", "DETALLE", "VALOR"],
            [date(2026, 5, 7), "Mercancía varia", 100000],
        ],
    })


# 19 · (v2.3 · B1) hoja «CUENTAS POR PAGAR» cuya columna de nombres dice «CLIENTE»
def cxp_con_columna_cliente() -> bytes:
    return _libro({
        "COMPRAS": [
            ["COMPRAS MARZO 2026"],
            ["FECHA", "PROVEEDOR", "PRODUCTO", "TOTAL", "FORMA DE PAGO"],
            [date(2026, 3, 2), "DISTRIBUIDORA NORTE", "ARROZ", 300000, "CREDITO"],
            [date(2026, 3, 5), "DISTRIBUIDORA NORTE", "FRIJOL", 100000, "CONTADO"],
        ],
        "CUENTAS POR PAGAR": [
            ["CUENTAS POR PAGAR MARZO 2026"],
            ["CLIENTE", "VALOR", "ABONO", "SALDO"],
            ["DISTRIBUIDORA NORTE", 300000, 100000, 200000],
            ["GRANOS DEL SUR", 150000, 0, 150000],
        ],
    })


# 20 · (v2.3 · B2) un bloque titulado con un mes que todavía no llega
def bloque_de_mes_futuro() -> bytes:
    return _libro({
        "VENTAS": [
            ["VENTAS SEPTIEMBRE 2026"],
            ["FECHA", "CLIENTE", "VALOR"],
            [date(2026, 9, 5), "CLIENTE A", 100000],
            [date(2026, 9, 20), "CLIENTE B", 50000],
            [],
            ["VENTAS NOVIEMBRE 2026"],
            ["FECHA", "CLIENTE", "VALOR"],
            [date(2026, 11, 3), "CLIENTE C", 70000],
            [date(2026, 11, 15), "CLIENTE D", 30000],
        ],
        "GASTOS": [
            ["GASTOS SEPTIEMBRE 2026"],
            ["FECHA", "CONCEPTO", "VALOR"],
            [date(2026, 9, 30), "ARRIENDO", 40000],
        ],
    })


# 21 · (v2.3 · B3) once bloques mensuales con la misma rareza: la forma de pago solo
#      en la primera fila, y en cuatro de ellos una fecha copiada del mes anterior
def muchos_bloques_iguales() -> bytes:
    filas = []
    for m in range(1, 12):
        filas += [[f"VENTAS {('ENERO FEBRERO MARZO ABRIL MAYO JUNIO JULIO AGOSTO SEPTIEMBRE OCTUBRE NOVIEMBRE').split()[m - 1]} 2025"],
                  ["FECHA", "CLIENTE", "VALOR", "FORMA DE PAGO"]]
        primera = date(2025, m - 1, 28) if m in (3, 5, 7, 9) else date(2025, m, 2)
        filas += [[primera, "CLIENTE UNO", 100000, "CONTADO"],
                  [date(2025, m, 10), "CLIENTE DOS", 50000, None],
                  [date(2025, m, 20), "CLIENTE TRES", 30000, None], []]
    return _libro({"VENTAS 2025": filas})


ARCHIVOS = {
    "01_ventas_compras_bloques.xlsx": ventas_compras_bloques,
    "02_columnas_mal_rotuladas.xlsx": columnas_mal_rotuladas,
    "03_fechas_copiadas_y_futuras.xlsx": fechas_copiadas,
    "04_totales_mezclados.xlsx": totales_mezclados,
    "05_hojas_vacias.xlsx": hojas_vacias,
    "06_cartera_terceros_desconocidos.xlsx": cartera_terceros,
    "07_gastos_sin_nit.csv": gastos_csv,
    "08_balance_de_prueba.pdf": balance_pdf,
    "08b_ventas_tabla.pdf": ventas_pdf,
    "09_ventas_en_word.docx": ventas_word,
    "10_escaneado.pdf": escaneado_pdf,
    "11_libro_diario.csv": libro_diario_csv,
    "12_inventario_insuficiente.xlsx": inventario_insuficiente,
    "13_datos_repetidos.xlsx": repetidos,
    "14_titulos_mal_escritos.xlsx": titulos_mal_escritos,
    "15_inventario_inicial_y_conteo.xlsx": inventario_inicial_conteo,
    "16_balance_excel.xlsx": balance_excel,
    "17_hoja_por_mes_lado_a_lado.xlsx": hoja_por_mes_lado_a_lado,
    "18_listas_sin_tipo.xlsx": listas_sin_tipo,
    "19_cxp_con_columna_cliente.xlsx": cxp_con_columna_cliente,
    "20_bloque_de_mes_futuro.xlsx": bloque_de_mes_futuro,
    "21_muchos_bloques_iguales.xlsx": muchos_bloques_iguales,
}


def generar() -> list[Path]:
    salida = []
    for nombre, fn in ARCHIVOS.items():
        _guardar(nombre, fn())
        salida.append(AQUI / nombre)
    return salida


if __name__ == "__main__":
    for p in generar():
        print(p.name, p.stat().st_size, "bytes")
