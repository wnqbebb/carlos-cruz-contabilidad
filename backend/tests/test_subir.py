"""Una sola puerta: se sube lo que sea y el backend decide (spec v2.2 · Fase 3).

Lo que se prueba aquí es que el contador no tenga que saber qué pantalla usar
ni qué formato espera la aplicación.
"""
from __future__ import annotations

import io

from openpyxl import Workbook

from app.config import FUENTES
from app.exportar import plantilla
from app.importadores import clasificador as clas
from app.importadores import identidad as ident


# ── utilidades ──────────────────────────────────────────────────────────────
def _xlsx(filas: list[list], hoja: str = "Hoja1") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _directorio() -> bytes:
    return _xlsx([
        ["NIT", "RAZON SOCIAL", "MUNICIPIO"],
        ["[NIT]", "FARMACIA NATURISTA ANTARES S.A.S.", "Guacarí"],
        ["900123456", "DISTRIBUCIONES DEL VALLE S.A.S.", "Cali"],
        ["805004321", "FERRETERIA LA 14 LTDA", "Buga"],
    ], "CLIENTES")


def _subir(cliente_api, archivos, cliente_id=""):
    return cliente_api.post(
        f"/api/subir?cliente_id={cliente_id}",
        files=[("archivos", (n, c)) for n, c in archivos],
    )


# ── clasificación ───────────────────────────────────────────────────────────
def test_un_directorio_de_clientes_se_reconoce_como_tal(base_limpia, cliente_api):
    r = _subir(cliente_api, [("mis clientes.xlsx", _directorio())])
    assert r.status_code == 200
    assert r.json()["clase"] == clas.DIRECTORIO


def test_una_sola_fila_con_nit_no_es_un_directorio(base_limpia, cliente_api):
    """La cabecera de una contabilidad también trae NIT y nombre: no basta."""
    uno = _xlsx([["NIT", "RAZON SOCIAL"], ["[NIT]", "ANTARES SAS"]])
    assert _subir(cliente_api, [("x.xlsx", uno)]).json()["clase"] != clas.DIRECTORIO


def test_la_plantilla_oficial_es_contabilidad_y_no_pregunta_nada(base_limpia, cliente_api):
    r = _subir(cliente_api, [("ejemplo.xlsx", plantilla.construir(caso="completo"))]).json()
    assert r["clase"] == clas.CONTABILIDAD
    # La hoja EMPRESA trae nombre y NIT: no hay nada que preguntar.
    assert r["falta"] == []
    assert r["identidad"]["campos"]["razon_social"]["confianza"] == ident.SEGURO


def test_la_contabilidad_real_se_reconoce_aunque_no_traiga_identidad(base_limpia, cliente_api):
    r = _subir(cliente_api, [("CONTABILIDAD.xls", (FUENTES / "CONTABILIDAD.xls").read_bytes())]).json()
    assert r["clase"] == clas.CONTABILIDAD
    assert any("hoja de trabajo" in (h["razon"] or "").lower() for h in r["hojas"])


def test_un_documento_de_word_identifica_al_cliente(base_limpia, cliente_api):
    r = _subir(cliente_api, [
        ("estatutos.docx", (FUENTES / "CORREGIDO_acta_y_estatutos_FANANT.docx").read_bytes()),
    ]).json()
    assert r["clase"] == clas.DOCUMENTOS
    assert "ANTARES" in r["identidad"]["campos"]["razon_social"]["valor"].upper()


def test_word_mas_excel_juntos_identidad_del_word_cifras_del_excel(base_limpia, cliente_api):
    """El caso del spec: estatutos + contabilidad en la misma subida."""
    r = _subir(cliente_api, [
        ("estatutos.docx", (FUENTES / "CORREGIDO_acta_y_estatutos_FANANT.docx").read_bytes()),
        ("CONTABILIDAD.xls", (FUENTES / "CONTABILIDAD.xls").read_bytes()),
    ]).json()
    assert r["clase"] == clas.CONTABILIDAD           # hay cifras: se va a calcular
    assert "ANTARES" in r["identidad"]["campos"]["razon_social"]["valor"].upper()
    origen = r["identidad"]["campos"]["razon_social"]["origen"]
    assert "estatutos.docx" in origen                # pero el nombre salió del Word


def test_el_nombre_del_archivo_es_el_ultimo_recurso(base_limpia, cliente_api):
    anonimo = _xlsx([["Fecha", "Cuenta", "Debito", "Credito"], ["2025-01-05", "110505", 100, 0]])
    r = _subir(cliente_api, [("CONTABILIDAD_TIENDA_JUAN_PEREZ_2026.xlsx", anonimo)]).json()
    campo = r["identidad"]["campos"].get("razon_social")
    # «CONTABILIDAD» y el año son ruido; «TIENDA» se respeta porque puede ser
    # parte del nombre del negocio. Va como sugerencia, para que él confirme.
    assert campo and campo["valor"] == "TIENDA JUAN PEREZ"
    assert campo["confianza"] == ident.SUGERIDO
    assert "nit" in r["falta"]                       # lo único que se pide


def test_un_archivo_que_no_se_puede_leer_lo_dice_claro(base_limpia, cliente_api):
    r = _subir(cliente_api, [("notas.docx", b"esto no es un docx")])
    assert r.status_code == 400
    assert "word" in r.json()["detail"]["mensaje"].lower()


def test_el_doc_viejo_pide_guardarlo_como_docx(base_limpia, cliente_api):
    r = _subir(cliente_api, [("acta.doc", b"\xd0\xcf\x11\xe0 contenido viejo")])
    assert r.status_code == 400
    assert ".docx" in r.json()["detail"]["mensaje"]


# ── confirmar: el archivo nunca se vuelve a pedir ──────────────────────────
def test_confirmar_crea_el_cliente_y_sigue_con_el_mismo_archivo(base_limpia, cliente_api):
    subida = _subir(cliente_api, [("ejemplo.xlsx", plantilla.construir(caso="completo"))]).json()
    assert subida["cliente"] is None                  # todavía no existe

    r = cliente_api.post(f"/api/subir/{subida['subida_id']}/confirmar", json={})
    assert r.status_code == 200
    datos = r.json()
    assert datos["clase"] == clas.CONTABILIDAD
    assert datos["cliente_id"]
    assert datos["sesion_id"] and datos["hojas"]      # listo para mapear y calcular

    ficha = cliente_api.get(f"/api/clientes/{datos['cliente_id']}").json()
    assert ficha["razon_social"]
    assert ficha["nit"]


def test_confirmar_sin_identidad_pide_solo_lo_que_falta(base_limpia, cliente_api):
    anonimo = _xlsx([["Fecha", "Cuenta", "Debito", "Credito"], ["2025-01-05", "110505", 100, 0]])
    subida = _subir(cliente_api, [("movimientos.xlsx", anonimo)]).json()

    r = cliente_api.post(f"/api/subir/{subida['subida_id']}/confirmar", json={})
    assert r.status_code == 422
    assert r.json()["detail"]["codigo"] == "faltan_datos_cliente"
    assert "nit" in r.json()["detail"]["falta"]

    # Con el NIT puesto, sigue sin volver a subir el archivo.
    r2 = cliente_api.post(f"/api/subir/{subida['subida_id']}/confirmar",
                          json={"crear": {"nit": "[NIT]", "razon_social": "TIENDA DE PRUEBA"}})
    assert r2.status_code == 200
    assert r2.json()["sesion_id"]


def test_si_el_nit_ya_existe_se_trabaja_sobre_ese_cliente(base_limpia, cliente_api):
    existente = cliente_api.post("/api/clientes", json={
        "nit": "[NIT]", "razon_social": "FARMACIA NATURISTA ANTARES S.A.S."}).json()
    subida = _subir(cliente_api, [("ejemplo.xlsx", plantilla.construir(caso="completo"))]).json()
    # La plantilla de ejemplo trae otro NIT, así que se fuerza el del cliente.
    r = cliente_api.post(f"/api/subir/{subida['subida_id']}/confirmar",
                         json={"cliente_id": existente["id"]}).json()
    assert r["cliente_id"] == existente["id"]
    assert len(cliente_api.get("/api/clientes?estado=").json()["clientes"]) == 1


def test_un_directorio_se_puede_importar_desde_la_misma_puerta(base_limpia, cliente_api):
    subida = _subir(cliente_api, [("clientes.xlsx", _directorio())]).json()
    r = cliente_api.post(f"/api/subir/{subida['subida_id']}/confirmar", json={}).json()
    assert r["clase"] == clas.DIRECTORIO
    assert r["informe"]["insertados"] == 3
    assert cliente_api.get("/api/clientes?estado=").json()["total"] == 3


def test_la_subida_caducada_lo_dice_sin_romperse(base_limpia, cliente_api):
    r = cliente_api.post("/api/subir/noexiste/confirmar", json={})
    assert r.status_code == 404
    assert r.json()["detail"]["codigo"] == "subida_expirada"


# ── identidad suelta para el formulario de cliente nuevo ───────────────────
def test_la_ficha_se_puede_llenar_desde_documentos(base_limpia, cliente_api):
    r = cliente_api.post("/api/identidad", files=[
        ("archivos", ("cuentas.docx", (FUENTES / "cuentas_de_cobro_Word.docx").read_bytes())),
    ]).json()
    campos = r["identidad"]["campos"]
    assert campos["nit"]["valor"] == "[NIT]"
    assert "ANTARES" in campos["razon_social"]["valor"].upper()


# ── ya no se asume FANANT ───────────────────────────────────────────────────
def test_sin_cliente_la_empresa_ya_no_es_fanant(base_limpia, cliente_api):
    """Antes, calcular sin cliente usaba el NIT y la razón social de FANANT."""
    r = cliente_api.post("/api/importar/demo?caso=basico").json()
    assert r["empresa"]["nit"] in ("", None) or "ANTARES" not in (r["empresa"]["razon_social"] or "").upper()


def test_la_plantilla_en_blanco_no_trae_datos_de_nadie():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(plantilla.construir()), read_only=True)
    valores = [str(c.value or "") for fila in wb["EMPRESA"].iter_rows() for c in fila]
    texto = " ".join(valores).upper()
    assert "ANTARES" not in texto
    assert "[NIT]" not in texto
