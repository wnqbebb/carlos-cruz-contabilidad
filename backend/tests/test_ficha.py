"""La ficha del cliente se llena sola con sus documentos (spec v2.2 · Fase 6).

La prueba obligatoria usa los documentos reales de FANANT (estatutos y cartas):
deben llenar 4 socios, capital, representante legal y suplente, contador con
T.P. y CIIU. Se compara contra la ficha de referencia EN MODO LECTURA: no se
crea ni se modifica ningún cliente real.
"""
from __future__ import annotations

import io
import json
from decimal import Decimal

import pytest

from app.config import EMPRESA_PRIVADA, FUENTES
from app.importadores.ficha import extraer, tipo_documento
from app.importadores.lector import leer_archivo
from app.utils.numeros import normalizar

DOCS_FANANT = ["CORREGIDO_acta_y_estatutos_FANANT.docx", "CARTAS_VARIAS.docx"]
REFERENCIA = EMPRESA_PRIVADA


def _hay_fuentes() -> bool:
    return all((FUENTES / n).exists() for n in DOCS_FANANT) and REFERENCIA.exists()


requiere_fanant = pytest.mark.skipif(
    not _hay_fuentes(),
    reason="Los documentos reales de FANANT no están en esta copia (viven en privado/, fuera de git).",
)


def _ficha(nombres):
    return extraer({n: leer_archivo((FUENTES / n).read_bytes(), n) for n in nombres})


@requiere_fanant
def test_los_estatutos_y_las_cartas_llenan_la_ficha_de_fanant():
    ref = json.loads(REFERENCIA.read_text(encoding="utf-8"))
    f = _ficha(DOCS_FANANT)
    v = f.valores()
    # Socios: los cuatro, con su cédula y sus acciones.
    assert len(f.socios) == 4
    assert {s["cedula"] for s in f.socios} == {a["cc"].replace(".", "") for a in ref["accionistas"]}
    assert all(Decimal(s["acciones"]) == 1000 for s in f.socios)
    # Capital y acciones.
    assert Decimal(v["capital_suscrito"]) == Decimal(ref["capital_suscrito"])
    assert Decimal(v["valor_nominal_accion"]) == Decimal(ref["valor_nominal_accion"])
    assert Decimal(v["numero_acciones"]) == 4000
    # Representante legal y suplente, con cédula.
    assert normalizar(v["rep_legal"]) == normalizar(ref["rep_legal"])
    assert v["rep_legal_cc"] == ref["rep_legal_cc"].replace(".", "")
    assert normalizar(v["rep_legal_suplente"]) == normalizar(ref["rep_legal_suplente"])
    cc_suplente = next(a["cc"] for a in ref["accionistas"] if normalizar(a["nombre"]) == normalizar(ref["rep_legal_suplente"]))
    assert v["rep_legal_suplente_cc"] == cc_suplente.replace(".", "")
    # Contador con su tarjeta profesional (de las cartas).
    assert normalizar(v["contador"]) == normalizar(ref["contador"])
    assert v["contador_tp"] == ref["contador_tp"]
    assert v["contador_cc"] == ref["contador_cc"].replace(".", "")
    # CIIU principal y secundarios.
    assert ", ".join([v["ciiu"], v["ciiu_secundarios"]]) == ref["ciiu"]
    # Lo demás que trae el documento.
    assert v["sigla"] == ref["sigla"] and v["tipo_sociedad"] == "S.A.S."
    assert v["municipio"] == "Guacarí" and v["departamento"] == "Valle del Cauca"
    assert v["fecha_constitucion"] == "2024-11-26"
    assert "0002" in v["documento_constitucion"]


@requiere_fanant
def test_cada_dato_dice_de_donde_salio_y_con_que_confianza():
    j = _ficha(DOCS_FANANT).a_json()
    for campo, dato in j["campos"].items():
        assert dato["origen"], campo
        assert dato["confianza"] in ("seguro", "por confirmar"), campo
    assert j["campos"]["capital_suscrito"]["confianza"] == "seguro"
    assert j["campos"]["contador_tp"]["origen"].startswith("CARTAS_VARIAS.docx")
    assert {d["tipo"] for d in j["documentos"]} == {"estatutos", "carta"}


@requiere_fanant
def test_los_documentos_que_no_coinciden_quedan_en_conflicto():
    # La carta trae «FARMACIA NATURITA» (errata del original): se marca, no se escoge en silencio.
    j = _ficha(DOCS_FANANT).a_json()
    valores = {c["valor"] for c in j["conflictos"]["razon_social"]}
    assert "FARMACIA NATURISTA ANTARES S.A.S." in valores and "FARMACIA NATURITA ANTARES S.A.S." in valores
    assert j["campos"]["razon_social"]["valor"] == "FARMACIA NATURISTA ANTARES S.A.S."     # gana el estatuto


def _docx(parrafos: list[str], tabla: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    for p in parrafos:
        d.add_paragraph(p)
    if tabla:
        t = d.add_table(rows=len(tabla), cols=len(tabla[0]))
        for i, fila in enumerate(tabla):
            for j, v in enumerate(fila):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(lineas: list[str]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    y = 750
    for t in lineas:
        c.drawString(50, y, t)
        y -= 16
    c.save()
    return buf.getvalue()


def test_un_rut_ficticio_llena_los_datos_tributarios():
    rut = _pdf([
        "FORMULARIO DEL REGISTRO ÚNICO TRIBUTARIO - DIAN",
        "Número de Identificación Tributaria (NIT): 900.123.456 - 8",
        "Tipo de contribuyente: Persona jurídica",
        "Razón social: TALLER DE PRUEBA S.A.S.",
        "Dirección principal: Carrera 5 No. 10-20",
        "Municipio/Ciudad: Buga",
        "Departamento: Valle del Cauca",
        "Correo electrónico: taller@ejemplo.co",
        "Teléfono 1: 3001234567",
        "Actividad principal  Código: 4520",
        "Actividad secundaria  Código: 4530",
        "Responsabilidades",
        "05 - Impuesto renta y complementario régimen ordinario",
        "48 - Impuesto sobre las ventas - IVA",
        "52 - Facturador electrónico",
    ])
    f = extraer({"RUT_taller.pdf": leer_archivo(rut, "RUT_taller.pdf")})
    v = f.valores()
    assert f.documentos[0]["tipo"] == "rut"
    assert v["nit"] == "900123456" and v["razon_social"] == "TALLER DE PRUEBA S.A.S."
    assert v["municipio"] == "Buga" and v["departamento"] == "Valle del Cauca"
    assert v["email"] == "taller@ejemplo.co" and v["telefono"] == "3001234567"
    assert v["ciiu"] == "4520" and v["ciiu_secundarios"] == "4530"
    assert v["responsabilidades"] == "05, 48, 52" and v["responsable_iva"] == "true"
    assert v["tipo_persona"] == "juridica"


def test_un_certificado_de_camara_trae_matricula_y_renovacion():
    camara = _pdf([
        "CÁMARA DE COMERCIO DE BUGA",
        "CERTIFICADO DE EXISTENCIA Y REPRESENTACIÓN LEGAL",
        "Razón social: TALLER DE PRUEBA S.A.S.",
        "Matrícula No. 123456",
        "Fecha de renovación: 15 de marzo de 2026",
        "Representante legal: PEDRO PABLO PEREZ GOMEZ",
        "Capital suscrito: $ 20.000.000",
    ])
    v = extraer({"camara.pdf": leer_archivo(camara, "camara.pdf")}).valores()
    assert v["matricula_mercantil"] == "123456" and v["fecha_renovacion"] == "2026-03-15"
    assert v["rep_legal"] == "Pedro Pablo Perez Gomez" and v["capital_suscrito"] == "20000000"


def test_lo_que_no_aparece_queda_vacio():
    carta = _docx(["Señores BANCO X", "Solicitamos abrir una cuenta.", "Atentamente,"])
    v = extraer({"carta.docx": leer_archivo(carta, "carta.docx")}).valores()
    assert "capital_suscrito" not in v and "rep_legal" not in v and "contador_tp" not in v


@pytest.mark.parametrize("texto, archivo, tipo", [
    ("REGISTRO ÚNICO TRIBUTARIO", "x.pdf", "rut"),
    ("CÁMARA DE COMERCIO ... CERTIFICA", "x.pdf", "camara"),
    ("ESTATUTOS DE SOCIEDAD", "x.docx", "estatutos"),
    ("EL SUSCRITO CONTADOR HACE CONSTAR", "x.docx", "carta"),
])
def test_tipo_de_documento(texto, archivo, tipo):
    assert tipo_documento(texto, archivo) == tipo


def test_la_api_de_identidad_devuelve_la_ficha(cliente_api):
    estatutos = _docx([
        "ACTA DE CONSTITUCIÓN",
        "Se constituye una sociedad por acciones simplificada denominada PANES DEL SUR S.A.S. distinguida con la sigla “PDS”.",
        "El capital suscrito y pagado es de $ 10.000.000, dividido en mil (1.000) acciones.",
        "El representante legal principal será el señor JUAN CAMILO RIOS VEGA y su suplente será la señora ANA LUCIA VEGA.",
    ], [["NOMBRE", "CEDULA", "ACCIONES", "%"], ["Juan Camilo Rios Vega", "10.111.222", "600", "60"],
        ["Ana Lucia Vega", "20.333.444", "400", "40"]])
    r = cliente_api.post("/api/identidad", files=[("archivos", ("acta.docx", estatutos))]).json()
    campos = r["ficha"]["campos"]
    assert campos["razon_social"]["valor"] == "PANES DEL SUR S.A.S." and campos["sigla"]["valor"] == "PDS"
    assert campos["capital_suscrito"]["valor"] == "10000000"
    assert campos["rep_legal"]["valor"] == "Juan Camilo Rios Vega" and campos["rep_legal_cc"]["valor"] == "10111222"
    assert len(r["ficha"]["socios"]) == 2 and r["ficha"]["socios"][0]["participacion"] == "0.6"


def test_ciiu_con_y_entre_codigos():
    acta = _docx(["ACTA DE CONSTITUCIÓN", "Actividad identificada en el CIIU con los códigos 1081 y 4721."])
    v = extraer({"acta.docx": leer_archivo(acta, "acta.docx")}).valores()
    assert v["ciiu"] == "1081" and v["ciiu_secundarios"] == "4721"
