"""Libros oficiales y todos los entregables desde cualquier archivo (spec v2.2 · Fase 5, H05, H06, H14).

«Subo un Excel (o un PDF o un Word) y la aplicación calcula el balance de
prueba, el definitivo, los estados financieros, los saldos de inventario, el
libro diario y el libro mayor.» Aquí se comprueba esa frase, formato por formato.
"""
from __future__ import annotations

import io
import zipfile
from collections import defaultdict
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests.test_archivos_variados import BANCO, _calcular, _subir

D = Decimal

ENTREGABLES = ("libro_diario", "mayor_balances", "balance_prueba", "balance_definitivo", "situacion_financiera",
               "estado_resultados", "cambios_patrimonio", "flujo_efectivo", "indicadores", "inventario_saldos")


@pytest.mark.parametrize("archivo", [
    "01_ventas_compras_bloques.xlsx",       # Excel de registros auxiliares
    "11_libro_diario.csv",                   # CSV en partida doble
    "08_balance_de_prueba.pdf",              # PDF con texto
    "08b_ventas_tabla.pdf",                  # tabla en PDF
    "09_ventas_en_word.docx",                # tabla en Word
    "16_balance_excel.xlsx",                 # balance en Excel
])
def test_todos_los_entregables_desde_cualquier_formato(cliente_api, archivo):
    imp = _subir(cliente_api, archivo)
    res = _calcular(cliente_api, imp)
    faltan = [k for k in ENTREGABLES if k not in res["reportes"]]
    assert not faltan, f"{archivo}: faltan {faltan}"
    assert res["notas"], "las notas a los estados financieros siempre van"
    # Sin kardex, el saldo de inventario aparece igual, con la explicación y sin ceros inventados.
    inv = res["reportes"]["inventario_saldos"]
    if not res["inventario"]["productos"]:
        assert inv["filas"][0]["tipo"] == "nota" and "no trae movimientos de inventario" in inv["filas"][0]["valores"]["descripcion"]
    assert res["reportes"]["libro_diario"]["verificacion"]["cuadra"]
    assert res["reportes"]["mayor_balances"]["verificacion"]["cuadra"]


def test_libro_diario_oficial(cliente_api):
    imp = _subir(cliente_api, "01_ventas_compras_bloques.xlsx")
    res = _calcular(cliente_api, imp)
    ld = res["reportes"]["libro_diario"]
    claves = [c["clave"] for c in ld["columnas"]]
    assert claves[:9] == ["fecha", "comprobante", "tipo", "codigo", "cuenta", "tercero", "descripcion", "debito", "credito"]
    secciones = [f["valores"]["descripcion"] for f in ld["filas"] if f["tipo"] == "seccion"]
    # Orden cronológico: compras del 2 y 3 de enero, ventas del 10 y 20, febrero al final.
    assert [s.split(" · ")[0] for s in secciones] == ["CMP-0001", "CMP-0002", "VTA-0001", "VTA-0002", "VTA-0003"]
    assert "Factura de venta" in secciones[2] and "2026-01-10" in secciones[2]
    # Total por comprobante y del periodo.
    sub = [f for f in ld["filas"] if f["tipo"] == "subtotal"]
    assert all(D(f["valores"]["debito"]) == D(f["valores"]["credito"]) for f in sub)
    total = next(f for f in ld["filas"] if f["tipo"] == "total")["valores"]
    # Ventas 115.000 + costo 75.000 + compras 130.000 = 320.000 a cada lado
    assert D(total["debito"]) == D(total["credito"]) == D("320000")
    assert ld["verificacion"]["comprobantes"] == 5
    linea = next(f for f in ld["filas"] if f["tipo"] == "linea")
    assert linea["valores"]["origen"].endswith("› COMPRAS › fila 3")


def test_libro_diario_incluye_los_ajustes_aceptados(cliente_api):
    imp = _subir(cliente_api, "15_inventario_inicial_y_conteo.xlsx")
    res = _calcular(cliente_api, imp)
    secciones = [f["valores"]["descripcion"] for f in res["reportes"]["libro_diario"]["filas"] if f["tipo"] == "seccion"]
    assert any(s.startswith("AJ-FIS") and "ajuste" in s for s in secciones)   # faltante del conteo físico


def test_mayor_y_balances_con_subtotales_por_grupo_y_clase(cliente_api):
    imp = _subir(cliente_api, "06_cartera_terceros_desconocidos.xlsx")
    res = _calcular(cliente_api, imp)
    mb = res["reportes"]["mayor_balances"]
    assert [c["clave"] for c in mb["columnas"]] == ["codigo", "cuenta", "ini_d", "ini_c", "mov_d", "mov_c", "fin_d", "fin_c"]
    subtotales = {f["valores"]["cuenta"]: f["valores"] for f in mb["filas"] if f["tipo"] == "subtotal"}
    # Clase 1: caja 60.000 + clientes 150.000 = 210.000 de nuevo saldo débito
    assert D(subtotales["Total clase 1"]["fin_d"]) == D("210000")
    assert D(subtotales["Total grupo 11"]["fin_d"]) == D("60000")
    assert D(subtotales["Total grupo 13"]["fin_d"]) == D("150000")
    # Saldo anterior: Doña Rosa 80.000 en clientes y Postobón 200.000 en proveedores
    assert D(subtotales["Total grupo 13"]["ini_d"]) == D("80000")
    assert D(subtotales["Total grupo 22"]["ini_c"]) == D("200000")
    total = next(f for f in mb["filas"] if f["tipo"] == "total")["valores"]
    assert D(total["fin_d"]) == D(total["fin_c"])


def test_libros_en_excel_y_pdf(cliente_api):
    imp = _subir(cliente_api, "11_libro_diario.csv")
    res = _calcular(cliente_api, imp)
    sid, periodo = imp["sesion_id"], res["periodo"]["id"]
    wb = load_workbook(io.BytesIO(cliente_api.get(f"/api/exportar/{sid}/excel").content))
    assert {"Libro diario", "Mayor y balances"} <= set(wb.sheetnames)
    pdf = cliente_api.get(f"/api/exportar/{sid}/pdf").content
    from pypdf import PdfReader

    texto = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "LIBRO DIARIO" in texto and "LIBRO MAYOR Y BALANCES" in texto
    for libro in ("libro-diario", "mayor-balances"):
        for formato in ("excel", "pdf"):
            r = cliente_api.get(f"/api/periodos/{periodo}/{libro}/{formato}")
            assert r.status_code == 200 and len(r.content) > 1000, (libro, formato)


def test_el_pdf_del_libro_empieza_en_la_primera_pagina_y_con_miles(cliente_api):
    """El primer informe horizontal dejaba una página en blanco, y los importes guardados (texto)
    salían sin separador de miles."""
    from pypdf import PdfReader

    imp = _subir(cliente_api, "11_libro_diario.csv")
    res = _calcular(cliente_api, imp)
    datos = cliente_api.get(f"/api/periodos/{res['periodo']['id']}/libro-diario/pdf").content
    primera = PdfReader(io.BytesIO(datos)).pages[0].extract_text()
    assert "LIBRO DIARIO" in primera
    assert "5.000.000" in primera and "1.190.000" in primera
    assert "NIT " in primera and "-" in primera.split("NIT ", 1)[1].split()[0]     # con dígito de verificación


def test_un_periodo_guardado_antes_de_los_libros_los_arma_con_lo_guardado(cliente_api):
    """Periodos de la v2.1 (como FANANT enero 2025): sin libros en el resultado guardado."""
    from app.repositorio import periodos as repo

    imp = _subir(cliente_api, "11_libro_diario.csv")
    res = _calcular(cliente_api, imp)
    pid = res["periodo"]["id"]
    guardado = repo.resultado(pid)
    viejo = {**guardado["resultado"], "reportes": {k: v for k, v in guardado["resultado"]["reportes"].items()
                                                   if k not in ("libro_diario", "mayor_balances")}}
    from app.db import conexion
    from app.esquema import resultados as TR
    from sqlalchemy import update

    with conexion() as cn:
        cn.execute(update(TR).where(TR.c.periodo_id == pid).values(payload=viejo))
    datos = cliente_api.get(f"/api/periodos/{pid}/resultado").json()["resultado"]
    ld = datos["reportes"]["libro_diario"]
    assert ld["verificacion"]["cuadra"] and ld["verificacion"]["comprobantes"] == 3
    assert datos["reportes"]["mayor_balances"]["verificacion"]["cuadra"]


def test_movimientos_por_periodo_con_sumas_exactas(cliente_api):
    imp = _subir(cliente_api, "11_libro_diario.csv")
    res = _calcular(cliente_api, imp, periodizacion="por_periodo")
    enero = next(p for p in res["periodos_procesados"] if p["desde"] == "2026-01-01")
    m = cliente_api.get(f"/api/clientes/{imp['cliente_id']}/movimientos?periodo_id={enero['periodo_id']}").json()
    assert m["total"] == 5
    # Sumas como texto exacto, hechas con Decimal (no con SUM de SQLite, que da float)
    assert m["suma_debito"] == m["suma_credito"] and D(m["suma_debito"]) == D("6190000")


def test_h14_el_excel_no_usa_indigo_ni_verde(cliente_api):
    imp = _subir(cliente_api, "01_ventas_compras_bloques.xlsx")
    _calcular(cliente_api, imp)
    datos = cliente_api.get(f"/api/exportar/{imp['sesion_id']}/excel").content
    with zipfile.ZipFile(io.BytesIO(datos)) as z:
        xml = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml")).upper()
    for prohibido in ("4338CA", "EEF2FF", "059669", "ECFDF5"):
        assert prohibido not in xml, prohibido


def test_cada_comprobante_del_libro_cuadra_en_todo_el_banco(cliente_api):
    """Partida doble comprobante por comprobante, en todos los archivos del banco que calculan."""
    for archivo in ("02_columnas_mal_rotuladas.xlsx", "04_totales_mezclados.xlsx", "07_gastos_sin_nit.csv",
                    "12_inventario_insuficiente.xlsx", "13_datos_repetidos.xlsx", "17_hoja_por_mes_lado_a_lado.xlsx"):
        imp = _subir(cliente_api, archivo)
        res = _calcular(cliente_api, imp)
        suma: dict[str, list[Decimal]] = defaultdict(lambda: [D(0), D(0)])
        for f in res["reportes"]["libro_diario"]["filas"]:
            if f["tipo"] == "linea":
                suma[f["valores"]["comprobante"]][0] += D(f["valores"]["debito"] or 0)
                suma[f["valores"]["comprobante"]][1] += D(f["valores"]["credito"] or 0)
        assert suma and all(d == c for d, c in suma.values()), archivo
    assert (BANCO / "01_ventas_compras_bloques.xlsx").exists()
