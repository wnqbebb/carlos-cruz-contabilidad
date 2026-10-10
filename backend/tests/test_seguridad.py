"""Pruebas de seguridad (v2.3 · Fase 6): una por control de docs/SEGURIDAD.md."""
from __future__ import annotations

import io
import logging
import os
import time
import zipfile
from datetime import datetime, timedelta, timezone

import pyotp
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, select, update

from app import sesion
from app.db import conexion, lectura
from app.esquema import codigos_recuperacion, sesiones_acceso, usuarios
from app.main import app
from app.seguridad import aislado, archivos, cifrado, claves, incidentes, respaldo, secretos
from app.seguridad import http as seg_http
from tests.conftest import CLAVE_PRUEBA, USUARIO_PRUEBA, entrar

CLAVE_NUEVA = "Una-Clave-Nueva-Larga-2026"


def _vencer_reautenticacion(c: TestClient) -> None:
    with conexion() as cn:
        cn.execute(update(sesiones_acceso).where(sesiones_acceso.c.id == sesion._hash(c.cookies.get(sesion.COOKIE)))
                   .values(reautenticada=datetime.now(timezone.utc) - timedelta(minutes=10)))
    sesion.olvidar_fallos()


# ── C15, C17: nada interno a la vista ───────────────────────────────────
def test_documentacion_de_la_api_desactivada(cliente_sin_sesion):
    for ruta in ("/docs", "/redoc", "/openapi.json"):
        assert cliente_sin_sesion.get(ruta).status_code == 404


def test_salud_sin_sesion_solo_ok_y_version(cliente_sin_sesion):
    assert set(cliente_sin_sesion.get("/api/salud").json()) == {"ok", "version"}


def test_las_rutas_publicas_son_exactamente_las_previstas():
    assert sesion.PUBLICAS == {"/api/salud", "/api/sesion", "/api/acceso/primer-uso", "/api/acceso/recuperar",
                               "/api/acceso/politica"}


# ── C11, C10: host, origen y CSRF ───────────────────────────────────────
def test_host_ajeno_rechazado(cliente_sin_sesion):
    r = cliente_sin_sesion.get("/api/salud", headers={"host": "atacante.example"})
    assert r.status_code == 400
    assert cliente_sin_sesion.get("/api/salud", headers={"host": "localhost:8000"}).status_code == 200
    assert cliente_sin_sesion.get("/api/salud", headers={"host": "127.0.0.1:8000"}).status_code == 200


def test_origen_ajeno_rechazado(cliente_api):
    r = cliente_api.post("/api/clientes", json={"nit": "900000001", "razon_social": "X"},
                         headers={"origin": "https://atacante.example"})
    assert r.status_code == 403 and r.json()["detail"]["codigo"] == "origen"
    r = cliente_api.post("/api/sesion", json={"usuario": "x", "clave": "y"}, headers={"origin": "https://atacante.example"})
    assert r.status_code == 403


def test_sin_token_csrf_no_se_modifica_nada(cliente_api):
    sin = dict(cliente_api.headers)
    del cliente_api.headers[sesion.ENCABEZADO_CSRF]
    r = cliente_api.post("/api/clientes", json={"nit": "900000002", "razon_social": "SIN TOKEN"})
    assert r.status_code == 403 and r.json()["detail"]["codigo"] == "csrf"
    cliente_api.headers[sesion.ENCABEZADO_CSRF] = "otro-token"
    assert cliente_api.post("/api/clientes", json={"nit": "900000002", "razon_social": "X"}).status_code == 403
    cliente_api.headers.update(sin)
    assert cliente_api.get("/api/clientes").status_code == 200      # las lecturas no lo piden


# ── C13: encabezados ────────────────────────────────────────────────────
def test_encabezados_de_seguridad(cliente_api):
    for ruta in ("/api/tablero", "/"):
        h = cliente_api.get(ruta).headers
        assert "script-src 'self'" in h["content-security-policy"] and "frame-ancestors 'none'" in h["content-security-policy"]
        assert "unsafe-eval" not in h["content-security-policy"]
        assert h["x-content-type-options"] == "nosniff" and h["referrer-policy"] == "no-referrer"
        assert "camera=()" in h["permissions-policy"] and h["x-frame-options"] == "DENY"
    assert cliente_api.get("/api/tablero").headers["cache-control"] == "no-store"
    assert "strict-transport-security" in cliente_api.get("/api/tablero", headers={"x-forwarded-proto": "https"}).headers


# ── C12: límite de solicitudes ──────────────────────────────────────────
def test_limite_de_solicitudes(cliente_api, monkeypatch):
    seg_http.olvidar_limites()
    monkeypatch.setenv("CC_LIMITE_GENERAL", "5")
    codigos = [cliente_api.get("/api/clientes").status_code for _ in range(7)]
    assert codigos[:5] == [200] * 5 and codigos[-1] == 429
    monkeypatch.setenv("CC_LIMITE_GENERAL", "1000000")
    seg_http.olvidar_limites()


# ── C1, C2, C3: primer uso, recuperación y contraseñas ─────────────────
@pytest.fixture
def sin_usuario(base_limpia, monkeypatch):
    monkeypatch.delenv("CC_USUARIO", raising=False)
    monkeypatch.delenv("CC_CLAVE_HASH", raising=False)
    sesion.olvidar_fallos()
    with TestClient(app) as c:
        yield c


def test_primer_uso_solo_desde_el_mismo_equipo(sin_usuario, monkeypatch):
    monkeypatch.setenv("CC_HOSTS_LOCALES", "")
    assert sin_usuario.get("/api/sesion").json()["puede_crear"] is False
    r = sin_usuario.post("/api/acceso/primer-uso", json={"usuario": "carlos", "clave": CLAVE_NUEVA})
    assert r.status_code == 403 and r.json()["detail"]["codigo"] == "solo_local"


def test_primer_uso_crea_acceso_con_10_codigos_y_solo_su_hash(sin_usuario):
    assert sin_usuario.get("/api/sesion").json()["puede_crear"] is True
    debil = sin_usuario.post("/api/acceso/primer-uso", json={"usuario": "carlos", "clave": "password1234"})
    assert debil.status_code == 422 and debil.json()["detail"]["codigo"] == "clave_debil"
    r = sin_usuario.post("/api/acceso/primer-uso", json={"usuario": "carlos", "clave": CLAVE_NUEVA})
    assert r.status_code == 200
    codigos = r.json()["codigos"]
    assert len(codigos) == 10 and len(set(codigos)) == 10
    with lectura() as cn:
        hashes = [f.hash for f in cn.execute(select(codigos_recuperacion.c.hash)).all()]
        guardado = cn.execute(select(usuarios.c.hash)).scalar()
    assert all(h.startswith("$argon2id$") for h in hashes) and not any(c in "".join(hashes) for c in codigos)
    assert guardado.startswith("$argon2id$") and CLAVE_NUEVA not in guardado
    # Ya hay usuario: no se puede volver a crear.
    assert sin_usuario.post("/api/acceso/primer-uso", json={"usuario": "otro", "clave": CLAVE_NUEVA}).status_code == 409
    # ¿Olvidó la contraseña? Con un código: una vez sí, la segunda no.
    nueva = "Otra-Clave-Muy-Segura-77"
    r = sin_usuario.post("/api/acceso/recuperar", json={"usuario": "carlos", "codigo": codigos[0], "nueva": nueva})
    assert r.status_code == 200 and r.json()["codigos_restantes"] == 9
    assert sin_usuario.post("/api/acceso/recuperar", json={"usuario": "carlos", "codigo": codigos[0],
                                                           "nueva": nueva + "x"}).status_code == 401
    sesion.olvidar_fallos()
    assert sin_usuario.post("/api/sesion", json={"usuario": "carlos", "clave": nueva}).status_code == 200


def test_politica_de_contrasenas():
    assert not claves.evaluar("corta")["valida"]
    assert not claves.evaluar("password12345678")["valida"]             # común
    assert not claves.evaluar("carlos-cruz-2026-x", "carlos")["valida"]  # contiene el usuario
    ev = claves.evaluar("Mi-Contabilidad-Cuadra-2026!")
    assert ev["valida"] and ev["puntaje"] >= 3


def test_hash_bcrypt_anterior_se_migra_a_argon2id(cliente_sin_sesion):
    entrar(cliente_sin_sesion)                                           # el usuario de prueba viene en bcrypt
    with lectura() as cn:
        assert cn.execute(select(usuarios.c.hash)).scalar().startswith("$argon2id$")


# ── C4: verificación en dos pasos ──────────────────────────────────────
def test_totp_activar_ingresar_y_desactivar(cliente_api, monkeypatch):
    r = cliente_api.post("/api/acceso/totp/iniciar")
    assert r.status_code == 200 and r.json()["qr"].startswith("data:image/svg+xml")
    secreto = r.json()["secreto"]
    assert cliente_api.post("/api/acceso/totp/confirmar", json={"codigo": "000000"}).status_code == 422
    r = cliente_api.post("/api/acceso/totp/confirmar", json={"codigo": pyotp.TOTP(secreto).now()})
    assert r.status_code == 200 and len(r.json()["codigos"]) == 10
    codigos = r.json()["codigos"]
    with lectura() as cn:
        assert secreto not in (cn.execute(select(usuarios.c.totp_secreto)).scalar() or "")    # cifrado
    with TestClient(app) as otro:
        sesion.olvidar_fallos()
        r = otro.post("/api/sesion", json={"usuario": USUARIO_PRUEBA, "clave": CLAVE_PRUEBA})
        assert r.status_code == 401 and r.json()["detail"]["codigo"] == "requiere_totp"
        r = otro.post("/api/sesion", json={"usuario": USUARIO_PRUEBA, "clave": CLAVE_PRUEBA,
                                           "codigo": pyotp.TOTP(secreto).now()})
        assert r.status_code == 200
    # La base copiada a otro equipo (otra clave de datos): el código TOTP ya no se puede comprobar,
    # pero no hay error 500 y un código de recuperación sigue sirviendo.
    with monkeypatch.context() as m, TestClient(app) as otro:
        m.setenv("CC_CLAVE_DATOS", "ab" * 32)
        sesion.olvidar_fallos()
        r = otro.post("/api/sesion", json={"usuario": USUARIO_PRUEBA, "clave": CLAVE_PRUEBA,
                                           "codigo": pyotp.TOTP(secreto).now()})
        assert r.status_code == 401
        r = otro.post("/api/sesion", json={"usuario": USUARIO_PRUEBA, "clave": CLAVE_PRUEBA, "codigo": codigos[0]})
        assert r.status_code == 200
    _vencer_reautenticacion(cliente_api)
    assert cliente_api.post("/api/acceso/totp/desactivar").status_code == 403
    cliente_api.post("/api/sesion/reautenticar", json={"clave": CLAVE_PRUEBA, "codigo": pyotp.TOTP(secreto).now()})
    assert cliente_api.post("/api/acceso/totp/desactivar").status_code == 200


# ── C8: reautenticación ─────────────────────────────────────────────────
def test_acciones_delicadas_piden_la_contrasena(cliente_api):
    cid = cliente_api.post("/api/clientes", json={"nit": "900000003", "razon_social": "BORRABLE"}).json()["id"]
    _vencer_reautenticacion(cliente_api)
    for metodo, ruta in (("DELETE", f"/api/clientes/{cid}?definitivo=true"), ("POST", "/api/clientes/demostracion/eliminar"),
                         ("POST", "/api/versiones/1/restaurar"), ("POST", "/api/acceso/codigos")):
        r = cliente_api.request(metodo, ruta)
        assert r.status_code == 403 and r.json()["detail"]["codigo"] == "reautenticar", ruta
    assert cliente_api.post("/api/sesion/reautenticar", json={"clave": "mala"}).status_code == 401
    assert cliente_api.post("/api/sesion/reautenticar", json={"clave": CLAVE_PRUEBA}).status_code == 200
    assert cliente_api.delete(f"/api/clientes/{cid}?definitivo=true").status_code == 200
    # Archivar no es destructivo: no pide la contraseña.
    otro = cliente_api.post("/api/clientes", json={"nit": "900000004", "razon_social": "ARCHIVABLE"}).json()["id"]
    _vencer_reautenticacion(cliente_api)
    assert cliente_api.delete(f"/api/clientes/{otro}").status_code == 200


def test_cambiar_contrasena_cierra_las_otras_sesiones(cliente_api):
    with TestClient(app) as otro:
        entrar(otro)
        assert cliente_api.post("/api/acceso/clave", json={"actual": "mala", "nueva": CLAVE_NUEVA}).status_code == 401
        r = cliente_api.post("/api/acceso/clave", json={"actual": CLAVE_PRUEBA, "nueva": CLAVE_NUEVA})
        assert r.status_code == 200 and r.json()["sesiones_cerradas"] >= 1
        assert otro.get("/api/tablero").status_code == 401


# ── C20–C24: archivos subidos ───────────────────────────────────────────
def _bomba_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/relleno.bin", b"\x00" * (400 * 1024 * 1024))
    return buf.getvalue()


def _xlsx_con_entidad() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><!DOCTYPE t [<!ENTITY x SYSTEM "file:///c:/windows/win.ini">]><Types>&x;</Types>')
    return buf.getvalue()


def test_bomba_zip_rechazada_rapido(cliente_api):
    inicio = time.time()
    r = cliente_api.post("/api/subir", files=[("archivos", ("bomba.xlsx", _bomba_zip()))])
    assert r.status_code == 400 and r.json()["detail"]["codigo"] == "archivo_sospechoso"
    assert time.time() - inicio < 10


def test_xml_con_entidad_externa_rechazado(cliente_api):
    r = cliente_api.post("/api/subir", files=[("archivos", ("xxe.xlsx", _xlsx_con_entidad()))])
    assert r.status_code == 400 and r.json()["detail"]["codigo"] == "archivo_sospechoso"
    import openpyxl.xml
    assert openpyxl.xml.DEFUSEDXML


def test_imagen_gigante_rechazada(cliente_api):
    from PIL import Image

    buf = io.BytesIO()
    Image.new("1", (12000, 6000)).save(buf, format="PNG")            # 72 megapíxeles, pocos bytes
    cid = cliente_api.post("/api/clientes", json={"nit": "10000099", "razon_social": "P", "tipo_persona": "natural"}).json()["id"]
    r = cliente_api.post(f"/api/renta/{cid}/2025/documentos", files=[("archivos", ("foto.png", buf.getvalue(), "image/png"))])
    assert r.status_code == 400 and r.json()["detail"]["codigo"] == "imagen_gigante"


def test_tipo_por_contenido_no_por_extension(cliente_api):
    r = cliente_api.post("/api/subir", files=[("archivos", ("balance.xlsx", b"%PDF-1.4 esto es un pdf"))])
    assert r.status_code == 400 and r.json()["detail"]["codigo"] == "contenido_no_corresponde"
    assert archivos.tipo_por_contenido(b"MZ\x90\x00\x03") == "desconocido"


def test_macros_avisadas_nunca_ejecutadas():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/vbaProject.bin", b"\x00" * 100)
    assert any("macros" in a for a in archivos.validar("libro.xlsm", buf.getvalue(), (".xlsm",)))


def test_nombre_con_rutas_es_inofensivo(cliente_api):
    from app.repositorio import sesiones, subidas

    r = cliente_api.post("/api/subir", files=[("archivos", ("../../../windows/evil.csv", b"fecha;valor\n2025-01-01;10\n"))])
    assert r.status_code in (200, 400)
    if r.status_code == 200:
        [ref] = sesiones.obtener(r.json()["subida_id"])["archivos"]
        from pathlib import Path
        assert Path(ref["ruta"]).resolve().parent.parent == subidas.carpeta().resolve()
        assert ".." not in ref["nombre"] and "/" not in ref["nombre"]


def _cuelga():  # usado por el proceso aislado
    while True:
        time.sleep(0.1)


def _come_memoria():
    bloque = []
    while True:
        bloque.append(bytearray(50 * 1024 * 1024))
        time.sleep(0.05)


def test_lectura_aislada_con_tiempo_y_memoria_limitados(monkeypatch):
    monkeypatch.setenv("CC_AISLAR", "1")
    inicio = time.time()
    with pytest.raises(aislado.ArchivoNoProcesable) as ex:
        aislado.ejecutar("tests.test_seguridad._cuelga", tiempo=3)
    assert ex.value.codigo == "tiempo_agotado" and time.time() - inicio < 30
    with pytest.raises(aislado.ArchivoNoProcesable) as ex:
        aislado.ejecutar("tests.test_seguridad._come_memoria", tiempo=60, memoria_mb=300)
    assert ex.value.codigo == "memoria_agotada"
    assert aislado.ejecutar("app.seguridad.archivos.tipo_por_contenido", b"%PDF-1.7") == "pdf"


def test_pdf_malicioso_no_tumba_el_servidor(cliente_api, monkeypatch):
    monkeypatch.setenv("CC_AISLAR", "1")
    # Árbol de páginas que se apunta a sí mismo.
    pdf = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[2 0 R]/Count 1>>endobj\n"
           b"trailer<</Root 1 0 R>>\n%%EOF")
    r = cliente_api.post("/api/subir", files=[("archivos", ("raro.pdf", pdf))])
    assert r.status_code in (200, 400, 422)
    assert cliente_api.get("/api/salud").status_code == 200


# ── C28: inyección de fórmulas ──────────────────────────────────────────
def test_formula_del_usuario_se_exporta_como_texto():
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = '=HYPERLINK("http://atacante.example","clic")'
    ws["A2"] = "=SUM(B1:B9)"
    ws["A3"] = "@cmd"
    archivos.blindar_libro(wb)
    buf = io.BytesIO()
    wb.save(buf)
    leido = load_workbook(io.BytesIO(buf.getvalue())).active
    assert leido["A1"].data_type == "s" and leido["A2"].data_type == "f" and leido["A3"].data_type == "s"
    assert archivos.texto_seguro("=1+1") == "'=1+1" and archivos.texto_seguro(5) == 5


# ── C16, C18: incidentes y registros ────────────────────────────────────
def test_error_inesperado_da_codigo_de_incidente_sin_detalle(cliente_api, monkeypatch):
    from app.inteligencia import tablero

    def falla(*a, **k):
        raise RuntimeError(r"C:\Users\Admin\secreto postgres://usuario:clave@servidor/base")

    monkeypatch.setattr(tablero, "armar", falla)
    with TestClient(app, raise_server_exceptions=False) as c:
        entrar(c)
        r = c.get("/api/tablero")
    assert r.status_code == 500
    cuerpo = r.text
    assert "INC-" in cuerpo and "Users" not in cuerpo and "postgres" not in cuerpo and "RuntimeError" not in cuerpo


def test_registros_sin_datos_sensibles(caplog):
    texto = incidentes.limpiar("clave=Secreta123 cc_sesion=abcdefghijklmnop postgres://u:p@h/db cédula 1234567890 "
                               "NIT 900.123.456-7 total $ 1.234.567")
    assert "Secreta123" not in texto and "abcdefghijklmnop" not in texto and "u:p@" not in texto
    assert "1234567890" not in texto and "****7890" in texto and "900.123.456" not in texto and "****3456" in texto
    registro = logging.getLogger("carloscruz.prueba")
    registro.addFilter(incidentes.FiltroSensible())
    with caplog.at_level(logging.INFO):
        registro.info("ingreso de %s con password=%s", "1023456789", "otra-clave")
    assert "otra-clave" not in caplog.text and "1023456789" not in caplog.text


# ── C25, C30, C32: cifrado, secretos y copias ───────────────────────────
def test_cifrado_detecta_alteraciones():
    sellado = cifrado.cifrar(b"balance de prueba")
    assert cifrado.descifrar(sellado) == b"balance de prueba"
    alterado = sellado[:-1] + bytes([sellado[-1] ^ 1])
    with pytest.raises(cifrado.CifradoInvalido):
        cifrado.descifrar(alterado)


def test_secretos_salen_del_archivo_de_texto(tmp_path, monkeypatch):
    monkeypatch.setenv("CC_MIGRAR_SECRETOS", "1")
    env = tmp_path / ".env"
    env.write_text("ALMACENAMIENTO=supabase\nDATABASE_URL=postgresql://u:secreto@h/db\nCLAVE_SESION=abc\n", encoding="utf-8")
    movidos = secretos.migrar_archivo(env)
    texto = env.read_text(encoding="utf-8")
    assert set(movidos) == {"DATABASE_URL", "CLAVE_SESION"}
    assert "secreto" not in texto and "ALMACENAMIENTO=supabase" in texto
    assert secretos.leer("DATABASE_URL") == "postgresql://u:secreto@h/db"
    secretos.borrar("DATABASE_URL")
    secretos.borrar("CLAVE_SESION")


def test_cada_instalacion_tiene_su_propio_espacio_de_secretos(tmp_path, monkeypatch):
    a = secretos.servicio()
    otro = tmp_path / "otra.env"
    otro.write_text("ALMACENAMIENTO=local\n", encoding="utf-8")
    monkeypatch.setenv("CC_ENV", str(otro))
    assert secretos.servicio() != a


def test_el_espacio_de_secretos_va_con_el_archivo_de_configuracion(tmp_path, monkeypatch):
    """Una instancia con otra base pero el mismo archivo de configuración es la misma instalación:
    nunca se lleva los secretos de ese archivo a otro espacio del almacén."""
    a = secretos.servicio()
    monkeypatch.setenv("CC_SQLITE", str(tmp_path / "otra-base.db"))
    assert secretos.servicio() == a


def test_con_cc_env_solo_se_lee_ese_archivo(tmp_path, monkeypatch):
    from app import config

    propio = tmp_path / "propio.env"
    propio.write_text("X=1\n", encoding="utf-8")
    monkeypatch.setenv("CC_ENV", str(propio))
    assert config._candidatos_env() == [propio] and config.archivo_env() == propio


def test_copia_de_seguridad_cifrada_que_se_restaura(cliente_api, tmp_path, monkeypatch):
    monkeypatch.setenv("CC_RESPALDOS", str(tmp_path))
    cliente_api.post("/api/clientes", json={"nit": "901234567", "razon_social": "RESPALDADA S.A.S."})
    ruta = respaldo.crear()
    contenido = ruta.read_bytes()
    assert contenido.startswith(b"CC1") and b"RESPALDADA" not in contenido and b"901234567" not in contenido
    assert respaldo.probar(ruta)
    datos = respaldo.leer(ruta)
    assert any(c["razon_social"] == "RESPALDADA S.A.S." for c in datos["tablas"]["clientes"])
    ruta.write_bytes(contenido[:-3] + b"xyz")
    with pytest.raises(cifrado.CifradoInvalido):
        respaldo.leer(ruta)


def test_eliminar_cliente_borra_su_rastro(cliente_api):
    from app.esquema import bitacora as TB

    cid = cliente_api.post("/api/clientes", json={"nit": "10000077", "razon_social": "PERSONA X",
                                                 "tipo_persona": "natural"}).json()["id"]
    cliente_api.get(f"/api/renta/{cid}/2025")
    assert cliente_api.delete(f"/api/clientes/{cid}?definitivo=true").status_code == 200
    with lectura() as cn:
        filas = cn.execute(select(TB.c.detalle, TB.c.cliente_id)).all()
    assert not any(f.cliente_id == cid for f in filas)
    assert not any("PERSONA X" in str(f.detalle) or "10000077" in str(f.detalle) for f in filas)


def test_bitacora_de_seguridad_con_ip(cliente_sin_sesion):
    sesion.olvidar_fallos()
    cliente_sin_sesion.post("/api/sesion", json={"usuario": USUARIO_PRUEBA, "clave": "mala"})
    entrar(cliente_sin_sesion)
    estado = cliente_sin_sesion.get("/api/acceso/estado").json()
    assert estado["fallidos_24h"] >= 1 and estado["codigos_restantes"] == 0


def test_recortes_de_fotos_sin_metadatos():
    """C27: de una foto con EXIF y GPS solo se guardan recortes PNG re-codificados, sin metadatos."""
    np = pytest.importorskip("numpy")
    from PIL import Image

    from app.renta import ocr

    exif = Image.Exif()
    exif[0x010F] = "Telefono de prueba"            # fabricante
    exif[0x8825] = {1: "N", 2: (4.0, 36.0, 0.0)}   # GPS
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, "JPEG", exif=exif)
    assert b"Exif" in buf.getvalue()
    foto = np.array(Image.open(io.BytesIO(buf.getvalue())).convert("L"))
    png = ocr.recorte_png(ocr.LecturaImagen(tabla=foto), (10, 10, 200, 40))
    assert png.startswith(b"\x89PNG")
    assert not any(m in png for m in (b"Exif", b"eXIf", b"GPS", b"tEXt", b"iTXt", b"Telefono"))


def test_ninguna_ruta_frena_al_servidor():
    """Las rutas son `def`: FastAPI las corre en hilos y un OCR o un Excel grande no congelan a las demás."""
    import inspect

    asincronas = [r.path for r in app.routes
                  if getattr(r, "endpoint", None) and inspect.iscoroutinefunction(r.endpoint)]
    assert asincronas == []


def test_escribir_un_registro_nunca_frena_al_servidor(monkeypatch):
    """Una consola o un disco lentos (ventana con texto seleccionado, antivirus) no congelan el servidor."""
    import threading

    monkeypatch.delenv("CC_REGISTRO_DIRECTO", raising=False)
    escrito = threading.Event()

    class Lento(logging.Handler):
        def emit(self, record):
            time.sleep(1.0)
            escrito.set()

    lg = logging.getLogger("carloscruz.prueba-lento")
    lg.addHandler(Lento())
    lg.setLevel(logging.INFO)
    try:
        incidentes.desacoplar(("carloscruz.prueba-lento",))
        t = time.perf_counter()
        lg.info("acceso de prueba")
        assert time.perf_counter() - t < 0.2       # quien registra no espera
        assert escrito.wait(5)                     # y el registro sí se escribe
    finally:
        incidentes.detener_escritores()
        for h in list(lg.handlers):
            lg.removeHandler(h)


def test_clave_de_datos_de_la_nube_es_estable(monkeypatch):
    """Rescate H1: en Render la clave viene del entorno (texto libre generado por Render) y no cambia."""
    from app.seguridad import cifrado, secretos

    monkeypatch.setenv("CC_CLAVE_DATOS", "Zx9+generado/por=render")
    k1 = secretos.clave_datos()
    assert len(k1) == 32 and secretos.clave_datos() == k1
    caja = cifrado.cifrar(b"hola", contexto=b"prueba")
    assert cifrado.descifrar(caja, contexto=b"prueba") == b"hola"
    monkeypatch.setenv("CC_CLAVE_DATOS", "ab" * 32)
    assert secretos.clave_datos() == bytes.fromhex("ab" * 32)
