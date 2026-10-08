"""Lectura de PDF y los tres casos de ejemplo."""
from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.exportar import casos, plantilla
from app.importadores.lector import leer_archivo
from app.importadores.documentos import es_hoja_de_texto
from app.importadores.pdf_lector import PdfSinTexto, _partir_linea, leer_pdf


# ── el parser de líneas de PDF ──────────────────────────────────────────────
@pytest.mark.parametrize(
    "linea",
    [
        "1105 Caja general 37.144.505 30.000.000",
        "110505 Caja 1.500.000",
        "2510 Cesantias consolidadas 135.237,55",
        "4135 Comercio al por mayor y al por menor 22.641",
    ],
)
def test_reconoce_una_fila_de_balance(linea):
    fila = _partir_linea(linea)
    assert fila is not None, f"no reconoció «{linea}»"
    assert fila[0].isdigit() and len(fila[0]) >= 4   # código PUC
    assert any(c.isalpha() for c in fila[1])         # nombre de cuenta
    assert isinstance(fila[2], Decimal)              # importe exacto, nunca float


@pytest.mark.parametrize(
    "linea",
    [
        "FARMACIA NATURISTA ANTARES S.A.S. - NIT [NIT]-9",  # el NIT no es plata
        "Pagina 1 de 9",
        "Representante legal - C.C. [CEDULA]",
        "Contador publico - T.P. 103028-T",
        "TOTAL ACTIVO 37.144.505",            # sin código: es un total, no una cuenta
        "Generado el 06/10/2026",
        "12 Inversiones 0",                   # código de 2 dígitos e importe suelto
        "Del 1 de enero de 2025 al 31 de enero de 2025",
    ],
)
def test_descarta_lo_que_no_es_una_cuenta(linea):
    """Más vale una fila que no se leyó que una fila inventada.

    Una fila falsa ensucia los estados financieros en silencio; una que falta se
    nota de inmediato.
    """
    assert _partir_linea(linea) is None, f"aceptó por error «{linea}»"


def test_un_pdf_sin_texto_lo_dice_claro():
    """Un PDF escaneado no se puede leer: hay que avisar, no devolver vacío."""
    minimo = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
              b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
              b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]>>endobj\n"
              b"trailer<</Root 1 0 R>>")
    with pytest.raises((PdfSinTexto, ValueError)) as ex:
        leer_pdf(minimo, "escaneo.pdf")
    assert "escane" in str(ex.value).lower() or "no se pudo abrir" in str(ex.value).lower()


def test_el_lector_general_acepta_pdf(empresa):
    """Un PDF generado por la propia aplicación se vuelve a leer sin basura."""
    from app import motor
    from app.contabilidad.puc import Mapeador
    from app.exportar import pdf as gen_pdf
    from app.importadores.detector import detectar_archivos

    dets = detectar_archivos([("ej.xlsx", plantilla.construir(caso="completo"))], Mapeador(), empresa)
    items = motor.items_mapeo(dets, Mapeador())
    mapeo = {i["normalizado"]: i["codigo"] for i in items if i["codigo"]}
    paq, al, an, ae = motor.preparar_paquete(dets, {}, mapeo)
    res = motor.calcular(paq, empresa, motor.Config(), {}, al, an, ae)

    hojas = leer_archivo(gen_pdf.generar(res, empresa), "estados.pdf")

    # Desde la v2.2 el PDF entrega sus TABLAS (una hoja por tabla) más una hoja
    # «(texto)» con los párrafos, que es la que usa la identidad del cliente.
    texto = [h for h in hojas if es_hoja_de_texto(h)]
    tablas = [h for h in hojas if not es_hoja_de_texto(h)]
    assert len(texto) == 1, "falta la hoja de texto del PDF"
    assert tablas, "no se reconoció ninguna tabla del PDF"
    assert "ANTARES" in texto[0].texto_plano.upper()

    # Al menos una tabla tiene que traer cuentas con código PUC y nombre.
    def es_cuenta(fila):
        codigo, nombre = str(fila[0] or "").strip(), str(fila[1] or "")
        return codigo.isdigit() and len(codigo) >= 4 and any(c.isalpha() for c in nombre)

    assert any(any(es_cuenta(f) for f in t.valores if len(f) >= 2) for t in tablas),         "ninguna tabla del PDF trae cuentas reconocibles"


# ── los tres casos de ejemplo ───────────────────────────────────────────────
def test_hay_tres_casos_y_se_distinguen():
    lista = casos.listar()
    assert {c["id"] for c in lista} == {"completo", "mediocre", "basico"}
    por_id = {c["id"]: c for c in lista}
    assert por_id["completo"]["con_inventario"] and por_id["completo"]["con_nomina"]
    assert not por_id["basico"]["con_inventario"] and not por_id["basico"]["con_nomina"]
    assert por_id["completo"]["movimientos"] > por_id["basico"]["movimientos"]
    for c in lista:
        assert c["nombre"] and c["descripcion"], f"{c['id']} sin descripción"


@pytest.mark.parametrize("caso", ["completo", "mediocre", "basico"])
def test_cada_caso_genera_un_excel_cargable(caso):
    datos = plantilla.construir(caso=caso)
    wb = load_workbook(io.BytesIO(datos))
    assert "EMPRESA" in wb.sheetnames and "MOVIMIENTOS" in wb.sheetnames
    # La hoja de movimientos trae datos, no solo el encabezado.
    assert wb["MOVIMIENTOS"].max_row > 1


def test_el_caso_completo_cuadra(empresa, detectar):
    """El ejemplo «completo» tiene que cuadrar: es la carta de presentación."""
    from app import motor
    from app.contabilidad.puc import Mapeador
    from app.importadores.detector import detectar_archivos

    dets = detectar_archivos([("ej.xlsx", plantilla.construir(caso="completo"))], Mapeador(), empresa)
    items = motor.items_mapeo(dets, Mapeador())
    mapeo = {i["normalizado"]: i["codigo"] for i in items if i["codigo"]}
    paq, al, an, ae = motor.preparar_paquete(dets, {}, mapeo)
    res = motor.calcular(paq, empresa, motor.Config(), {}, al, an, ae)
    assert res["resumen"]["esf_cuadra"], "el caso completo no cuadra"
    assert res["resumen"]["bp_cuadra"]


def test_el_caso_mediocre_trae_los_problemas_a_proposito(empresa):
    """El ejemplo «mediocre» existe para enseñar qué detecta el sistema.

    Si algún día empieza a cuadrar y a no dar alertas, dejó de servir.
    """
    from app import motor
    from app.contabilidad.puc import Mapeador
    from app.importadores.detector import detectar_archivos

    dets = detectar_archivos([("ej.xlsx", plantilla.construir(caso="mediocre"))], Mapeador(), empresa)
    items = motor.items_mapeo(dets, Mapeador())
    sin_codigo = [i for i in items if not i["codigo"]]
    assert sin_codigo, "el caso mediocre debe traer cuentas sin código PUC para mapear"

    mapeo = {i["normalizado"]: i["codigo"] for i in items if i["codigo"]}
    paq, al, an, ae = motor.preparar_paquete(dets, {}, mapeo)
    res = motor.calcular(paq, empresa, motor.Config(), {}, al, an, ae)
    errores = res["resumen"]["alertas"].get("error", 0)
    assert errores > 0, "el caso mediocre debería producir errores que el contador vea"
