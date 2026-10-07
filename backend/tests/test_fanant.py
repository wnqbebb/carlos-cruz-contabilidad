"""Pruebas obligatorias (sección 11 del PROMPT) y de regresión."""
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from app import motor
from app.contabilidad.estados import resultados
from app.contabilidad.mayor import construir_mayor
from app.contabilidad.puc import Mapeador
from app.exportar import excel, pdf
from app.exportar.plantilla import construir
from app.importadores.detector import detectar_archivos
from app.inventario import kardex
from app.modelos import Empleado, MovInventario
from app.nomina.calculo import liquidar

D = Decimal


def _calcular(dets, empresa, incluir=None, decisiones=None, config=None):
    items = motor.items_mapeo(dets, Mapeador())
    mapeo = {i["normalizado"]: i["codigo"] for i in items if i["codigo"]}
    paq, al, an, ae = motor.preparar_paquete(dets, incluir or {}, mapeo)
    return paq, motor.calcular(paq, empresa, config or motor.Config(), decisiones or {}, al, an, ae)


def _solo(dets, hoja):
    return {d.id: d.hoja == hoja for d in dets}


# 1 ─ cuenta T
def test_1_cuenta_t(detectar, empresa):
    dets = detectar("CONTABILIDAD.xls")
    paq, _ = _calcular(dets, empresa, _solo(dets, "cuenta T"))
    mayor = construir_mayor(paq.saldos_iniciales, paq.movimientos)
    caja = mayor["110505"]
    assert (caja.mov_d, caja.mov_c, caja.saldo) == (D("7822641"), D("678136"), D("7144505"))
    assert sum(c.mov_d for c in mayor.values()) == sum(c.mov_c for c in mayor.values()) == D("8500777")
    assert resultados(mayor)["utilidad_neta"] == D("-653215")
    assert mayor["3105"].mov_c == D("7800000")
    assert mayor["240805"].mov_d == D("2280")


# 2 ─ hoja de trabajo
def test_2_hoja_trabajo(detectar, empresa):
    dets = detectar("CONTABILIDAD.xls")
    paq, res = _calcular(dets, empresa, _solo(dets, "HojaTRABAJO 01--02--03"))
    assert sum(s.debito for s in paq.saldos_iniciales) == sum(s.credito for s in paq.saldos_iniciales) == D("30000000")
    bp = res["reportes"]["balance_prueba"]["verificacion"]
    assert bp["cuadra"] and bp["saldos_finales"]["debito"] == D("37822641")  # caja 37.144.505 + IVA 2.280 + gastos 675.856
    assert res["resumen"]["utilidad_neta"] == D("-653215")
    assert res["resumen"]["esf_cuadra"] and res["resumen"]["total_activo"] == D("37144505")


# 3 ─ hoja vieja con descuadre
def test_3_descuadre_plantilla_antigua(detectar, empresa):
    dets = detectar("CONTABILIDAD.xls")
    d = next(x for x in dets if x.hoja == "HOJA DE TRABAJO ABRIL MAYO Y JUNIO")
    assert d.incluir is False and d.resumen["descuadre_archivo"] == D("598")
    assert any(a.codigo == "E8" and "598" in a.mensaje for a in d.paquete.alertas)
    assert d.resumen["años_titulo"] == [2012]
    _, res = _calcular(dets, empresa)
    assert any(a.codigo == "E19" for a in res["alertas"])  # título 2012 vs periodo 2025


def test_duplicado_cuenta_t_excluido(detectar):
    dets = detectar("CONTABILIDAD.xls")
    ct = next(x for x in dets if x.hoja == "cuenta T")
    assert ct.incluir is False and "Duplica" in ct.motivo


# 4 ─ estados financieros del contador
def test_4_auditoria_ef(detectar, empresa):
    dets = detectar("ESTADOS_FINANCIEROS.xlsx")
    h2 = next(x for x in dets if x.hoja == "Hoja2")
    aud = h2.paquete.auditoria_ef[0]
    assert aud["utilidad_corregida"] == D("4677120")
    codigos = {x["codigo"] for x in aud["hallazgos"]}
    assert {"E1", "E2", "E3", "E4", "E19", "E20", "E21"} <= codigos
    comp = {c["concepto"]: c for c in aud["comparativo"]}
    assert comp["Utilidad del ejercicio"]["archivo"] == D("4977120")
    assert comp["Capital social"]["archivo"] == D("1920880") and comp["Capital social"]["correcto"] == D("30000000")
    assert comp["Activo − (Pasivo + Patrimonio)"]["correcto"] == D("-27779120")
    # reconstrucción con el motor
    _, res = _calcular(dets, empresa, _solo(dets, "Hoja2"))
    assert res["resumen"]["utilidad_neta"] == D("4677120")
    assert res["resumen"]["esf_cuadra"] is False
    assert {"E1", "E2", "E3"} <= {a.codigo for a in res["alertas"]}


# 5 ─ nómina #002
def test_5_nomina_002(detectar):
    e = Empleado("ADRIANA DURAN JARAMILLO", salario_basico=D("1423500"), aux_transporte="si", mes=1, año=2025)
    l = liquidar(e, 2025)
    assert (l.salud_emp, l.pension_emp) == (D("56940"), D("56940"))
    assert l.neto == D("1509620")
    assert (l.cesantias, l.intereses, l.prima, l.vacaciones) == (D("135237.55"), D("16235"), D("135237.55"), D("59359.95"))
    assert (l.pension_empr, l.arl, l.caja) == (D("170820"), D("7430.67"), D("56940"))
    assert l.salud_empr == l.sena == l.icbf == 0  # exoneración art. 114-1 activa
    dets = detectar("NOMINA__enero__2025.xlsx")
    d = next(x for x in dets if x.hoja == "#002")
    assert d.incluir and d.resumen["mes"] == 1 and d.resumen["año"] == 2025
    assert {"E9", "E10", "E11", "E12", "E15", "E16"} <= {x["codigo"] for x in d.paquete.auditoria_nomina}
    assert all(not x.incluir for x in dets if x.hoja in ("#001", "#003"))


def test_nomina_sin_exoneracion():
    l = liquidar(Empleado("X", salario_basico=D("1423500"), aux_transporte="si", mes=1), 2025, exonerado=False)
    assert (l.salud_empr, l.sena, l.icbf) == (D("120997.50"), D("28470"), D("42705"))


# 6 ─ kardex
def _movs():
    return [MovInventario("P1", "compra", D(10), date(2025, 1, 1), costo_unitario=D(1000)),
            MovInventario("P1", "compra", D(10), date(2025, 1, 2), costo_unitario=D(1200)),
            MovInventario("P1", "venta", D(15), date(2025, 1, 3))]


def test_6_kardex_promedio():
    p = kardex.calcular(_movs(), "promedio")[0]["P1"]
    assert p.costo_ventas == D("16500") and p.saldo_cant == 5 and p.saldo_total == D("5500") and p.costo_promedio == D("1100")


def test_6_kardex_peps():
    p = kardex.calcular(_movs(), "peps")[0]["P1"]
    assert p.costo_ventas == D("16000") and p.saldo_total == D("6000")


def test_inventario_negativo():
    _, alertas = kardex.calcular([MovInventario("P9", "venta", D(3))])
    assert any(a.codigo == "INV-NEG" for a in alertas)


# 7 y 8 ─ demo completa: cierre, ecuación patrimonial, sumas iguales
def _demo(empresa):
    dets = detectar_archivos([("demo.xlsx", construir(demo=True))], Mapeador(), empresa)
    return _calcular(dets, empresa)


def test_7_cierre(empresa):
    _, res = _demo(empresa)
    mayor = res["mayor_ajustado"]
    from app.contabilidad.cierre import asiento_cierre, mayor_despues_cierre
    lineas, utilidad = asiento_cierre(mayor)
    despues = mayor_despues_cierre(mayor, lineas)
    assert all(c.neto == 0 for c in despues.values() if c.clase in "4567" and not c.codigo.startswith("5905"))
    cuenta = "3605" if utilidad > 0 else "3610"
    assert abs(despues[cuenta].neto) == abs(utilidad) == abs(res["resumen"]["utilidad_neta"])
    assert sum(c.fin_d for c in despues.values()) == sum(c.fin_c for c in despues.values())


def test_8_ecuacion_y_sumas(empresa, detectar):
    for _, res in (_demo(empresa), _calcular(detectar("CONTABILIDAD.xls", "NOMINA__enero__2025.xlsx"), empresa)):
        r = res["resumen"]
        assert r["bp_cuadra"] and r["ajustado_cuadra"] and r["esf_cuadra"] and r["hoja_trabajo_cuadra"]
        assert res["reportes"]["flujo_efectivo"]["verificacion"]["cuadra"]
        assert res["reportes"]["balance_definitivo"]["verificacion"]["cuadra"]
        assert r["total_activo"] == r["total_pasivo"] + r["total_patrimonio"]


def test_demo_inventario_y_ajustes(empresa):
    _, res = _demo(empresa)
    inv = res["inventario"]
    assert len(inv["productos"]) == 5
    assert any(v["estado"] == "Vencido" for v in inv["vencimientos"])
    assert {f["codigo"]: f["estado"] for f in inv["fisico"]}["P002"] == "Faltante"
    ids = {a["id"]: a for a in res["ajustes"]}
    assert ids["inventario_costo"]["aceptado"] and ids["depreciacion"]["aceptado"] and ids["nomina_causacion"]["aceptado"]
    assert inv["conciliacion"]["diferencia"] == 0
    assert D(str(ids["depreciacion"]["total"])) == D("16666.67")
    assert res["resumen"]["utilidad_neta"] > 0
    assert not [a for a in res["alertas"] if a.severidad == "error"]


# 9 ─ importador tolerante
def test_9_encabezados_desordenados(empresa):
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(["Crédito", "Débito", "Descripción", "Nombre de la cuenta", "Código", "Fecha", "Comprobante"])
    ws.append(["", "$ 1.423.500,00", "Aporte", "Caja", "110505", "15/01/2025", "C1"])
    ws.append(["1.423.500,00", "", "Aporte", "Capital social", "", "2025-01-15", "C1"])
    ws.append(["", "$ 300.000,oo", "Honorarios", "Gto honorarios", "", "2025-01-20", "C2"])
    ws.append(["300.000", "", "Pago", "Cajá", "", "2025-01-20", "C2"])
    buf = io.BytesIO()
    wb.save(buf)
    dets = detectar_archivos([("x.xlsx", buf.getvalue())], Mapeador(), empresa)
    assert dets[0].formato == "plantilla" and dets[0].resumen["movimientos"] == 4
    paq, res = _calcular(dets, empresa)
    assert {m.cuenta for m in paq.movimientos} == {"110505", "3105", "5110"}
    assert res["resumen"]["bp_cuadra"]
    assert paq.movimientos[0].fecha == date(2025, 1, 15)


def test_exportaciones(empresa):
    _, res = _demo(empresa)
    contenido = excel.libro_completo(res, empresa)
    wb = load_workbook(io.BytesIO(contenido))
    hojas = set(wb.sheetnames)

    # El libro va agrupado: los cuatro estados financieros comparten una sola
    # hoja, igual que los balances y el inventario. Antes eran 18 hojas sueltas.
    assert {"Portada", "Estados financieros", "Balances", "Hoja de trabajo",
            "Libro mayor", "Inventario", "Ajustes", "Alertas",
            "EF formato contador", "Saldos siguiente periodo"} <= hojas
    assert len(hojas) <= 14, f"demasiadas hojas: {sorted(hojas)}"

    # Y dentro de esa hoja están de verdad los cuatro estados.
    celdas = [str(c.value) for fila in wb["Estados financieros"].iter_rows() for c in fila if c.value]
    ef = " | ".join(celdas)
    for titulo in ("SITUACIÓN FINANCIERA", "RESULTADO", "PATRIMONIO", "EFECTIVO"):
        assert titulo in ef.upper(), f"falta la sección {titulo} en la hoja de estados financieros"

    # La portada lleva al cliente identificado y enlaces a las demás hojas.
    portada = wb["Portada"]
    textos = [str(c.value) for fila in portada.iter_rows() for c in fila if c.value]
    assert empresa.razon_social in textos
    enlaces = [c.hyperlink.target or c.hyperlink.location
               for fila in portada.iter_rows() for c in fila if c.hyperlink]
    assert len(enlaces) >= 6, "la portada debe tener índice navegable"

    assert pdf.generar(res, empresa)[:4] == b"%PDF"


def test_excel_desde_un_periodo_guardado(empresa):
    """Exportar un periodo guardado debe dar el mismo libro que el cálculo en vivo.

    En la base los importes se guardan como CADENA decimal (ver exactitud.py),
    así que el exportador tiene que aceptar las dos formas. Si no, el Excel de
    un periodo viejo saldría con las celdas vacías o como texto.
    """
    from app.exactitud import a_json

    _, res = _demo(empresa)
    guardado = a_json({k: v for k, v in res.items() if k != "mayor_ajustado"})

    libro_vivo = load_workbook(io.BytesIO(excel.libro_completo(res, empresa)))
    libro_base = load_workbook(io.BytesIO(excel.libro_completo(guardado, empresa)))

    # La hoja del formato del contador necesita el mayor en memoria; el resto
    # de hojas tienen que estar en los dos libros.
    assert set(libro_base.sheetnames) | {"EF formato contador"} == set(libro_vivo.sheetnames)

    vivo = libro_vivo["Estados financieros"]
    base = libro_base["Estados financieros"]
    numeros_vivo = [c.value for fila in vivo.iter_rows() for c in fila if isinstance(c.value, float)]
    numeros_base = [c.value for fila in base.iter_rows() for c in fila if isinstance(c.value, float)]
    assert numeros_vivo == numeros_base, "los importes no coinciden entre el cálculo en vivo y el guardado"
    assert numeros_base, "la hoja guardada no trae ningún importe"


def test_api(tmp_path):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    r = c.post("/api/importar/ejemplo")
    assert r.status_code == 200
    datos = r.json()
    assert datos["periodo_sugerido"]["desde"] == "2025-01-01"
    mapeo = {i["normalizado"]: i["codigo"] for i in datos["mapeo"] if i["codigo"]}
    r = c.post("/api/calcular", json={"sesion_id": datos["sesion_id"], "mapeo": mapeo, "recordar_alias": False})
    assert r.status_code == 200, r.text
    assert r.json()["resumen"]["bp_cuadra"]
    assert c.get(f"/api/exportar/{datos['sesion_id']}/excel").status_code == 200
    assert c.get(f"/api/exportar/{datos['sesion_id']}/pdf").content[:4] == b"%PDF"
    r = c.post("/api/importar/demo")
    assert r.status_code == 200
    assert c.get("/api/plantilla").status_code == 200
