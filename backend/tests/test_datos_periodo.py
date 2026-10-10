"""«Datos del periodo»: editar dentro de la aplicación (rescate H7 · recorrido J5 por la API)."""
from __future__ import annotations

from decimal import Decimal

from tests.test_entrada import _caso

D = Decimal


def _datos(api, pid):
    r = api.get(f"/api/periodos/{pid}/datos")
    assert r.status_code == 200, r.text
    return r.json()


def _guardar(api, pid, tablas, esperado=200):
    r = api.put(f"/api/periodos/{pid}/datos", json={"tablas": tablas})
    assert r.status_code == esperado, r.text
    return r.json()


def test_editar_valor_cuenta_agregar_borrar_y_saldo_inicial(cliente_api):
    c, calculado = _caso(cliente_api, "completo", "900100158")
    pid = calculado["periodo"]["id"]
    vista = _datos(cliente_api, pid)
    assert vista["editable"] is True and vista["validacion"]["cuadra"] is True
    movs = vista["tablas"]["movimientos"]
    saldos = vista["tablas"]["saldos_iniciales"]
    n = len(movs)

    # 1. Cambiar el valor de un comprobante (las dos líneas: sigue cuadrando) y la cuenta de una línea.
    papeleria = [i for i, m in enumerate(movs) if m["comprobante"] == "CE-003"]
    assert len(papeleria) == 2
    for i in papeleria:
        if D(movs[i]["debito"]):
            movs[i]["debito"] = "95000.00"
            movs[i]["cuenta"] = "519595"            # de papelería a otros gastos diversos
        else:
            movs[i]["credito"] = "95000.00"
    # 2. Agregar un comprobante nuevo (dos filas).
    movs.append({"fecha": "2025-01-30", "comprobante": "CE-099", "cuenta": "513550", "nombre_cuenta": "Transporte",
                 "debito": "40000", "credito": "0", "descripcion": "Flete digitado en la aplicación"})
    movs.append({"fecha": "2025-01-30", "comprobante": "CE-099", "cuenta": "110505", "nombre_cuenta": "Caja",
                 "debito": "0", "credito": "40000", "descripcion": "Pago flete"})
    # 3. Borrar un comprobante entero (comisiones bancarias).
    movs = [m for m in movs if m["comprobante"] != "NC-002"]
    # 4. Cambiar un saldo inicial (caja +100.000 contra capital, para que siga cuadrando).
    for s in saldos:
        if s["cuenta"] == "110505":
            s["debito"] = str(D(s["debito"]) + 100000)
        if s["cuenta"] == "3105":
            s["credito"] = str(D(s["credito"]) + 100000)

    nuevo = _guardar(cliente_api, pid, {"movimientos": movs, "saldos_iniciales": saldos})
    assert nuevo["guardado"] is True and nuevo["validacion"]["cuadra"] is True
    assert len(nuevo["tablas"]["movimientos"]) == n + 2 - 2
    assert nuevo["versiones"] == 1

    # Los estados se recalcularon solos: utilidad = antes − 10.000 (papelería) − 40.000 (flete) + 15.000 (comisión)
    antes = D(calculado["resumen"]["utilidad_neta"])
    res = cliente_api.get(f"/api/periodos/{pid}/resultado").json()["resultado"]
    assert D(res["resumen"]["utilidad_neta"]) == antes - 10000 - 40000 + 15000
    assert D(res["resumen"]["descuadre_esf"]) == 0
    # El libro diario guardado refleja el cambio.
    diario = cliente_api.get(f"/api/clientes/{c['id']}/movimientos?periodo_id={pid}&por_pagina=2000").json()
    assert any(m["comprobante"] == "CE-099" for m in diario["movimientos"])
    assert not any(m["comprobante"] == "NC-002" for m in diario["movimientos"])

    # 5. Deshacer: vuelve exactamente a como estaba.
    r = cliente_api.post(f"/api/periodos/{pid}/deshacer")
    assert r.status_code == 200, r.text
    res = cliente_api.get(f"/api/periodos/{pid}/resultado").json()["resultado"]
    assert D(res["resumen"]["utilidad_neta"]) == antes
    # La versión editada no se pierde: quedó en el historial.
    assert len(cliente_api.get(f"/api/periodos/{pid}/versiones").json()["versiones"]) >= 2


def test_descuadre_se_avisa_y_cuenta_inexistente_no_se_guarda(cliente_api):
    _, calculado = _caso(cliente_api, "completo", "900100158")
    pid = calculado["periodo"]["id"]
    movs = _datos(cliente_api, pid)["tablas"]["movimientos"]
    movs[0]["debito"] = str(D(movs[0]["debito"] or 0) + 1000) if D(movs[0]["debito"] or 0) else movs[0]["debito"]
    if not D(movs[0]["debito"] or 0):
        movs[0]["credito"] = str(D(movs[0]["credito"]) + 1000)
    v = cliente_api.post(f"/api/periodos/{pid}/datos/validar", json={"tablas": {"movimientos": movs}}).json()
    assert v["cuadra"] is False and v["descuadres"]
    # Se deja guardar (el contador decide), con el descuadre a la vista.
    nuevo = _guardar(cliente_api, pid, {"movimientos": movs})
    assert nuevo["validacion"]["cuadra"] is False
    # Una cuenta que no existe sí bloquea.
    movs[1]["cuenta"] = "999999"
    r = cliente_api.put(f"/api/periodos/{pid}/datos", json={"tablas": {"movimientos": movs}})
    assert r.status_code == 422 and "999999" in r.json()["detail"]["mensaje"]


def test_periodo_cerrado_pide_reabrir(cliente_api):
    c, calculado = _caso(cliente_api, "completo", "900100158")
    pid = calculado["periodo"]["id"]
    r = cliente_api.post(f"/api/periodos/{pid}/cerrar")
    assert r.status_code == 200, r.text
    vista = _datos(cliente_api, pid)
    assert vista["editable"] is False and "Reábralo" in vista["motivo_bloqueo"]
    r = cliente_api.put(f"/api/periodos/{pid}/datos", json={"tablas": {"movimientos": vista["tablas"]["movimientos"]}})
    assert r.status_code == 409


def test_descargar_datos_para_editar_y_volver_a_subir(cliente_api):
    """J6: «Descargar datos para editar» → modificar en Excel → subir → el mismo periodo se actualiza."""
    import io

    from openpyxl import load_workbook

    c, calculado = _caso(cliente_api, "completo", "900100158")
    pid = calculado["periodo"]["id"]
    r = cliente_api.get(f"/api/periodos/{pid}/datos/excel")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["MOVIMIENTOS"]
    encab = [c.value for c in ws[1]]
    i_comp, i_deb, i_cre = encab.index("No. comprobante"), encab.index("Débito"), encab.index("Crédito")
    for fila in ws.iter_rows(min_row=2):
        if fila[i_comp].value == "CE-003":
            if fila[i_deb].value:
                fila[i_deb].value = 95000
            else:
                fila[i_cre].value = 95000
    buf = io.BytesIO()
    wb.save(buf)
    p = cliente_api.post("/api/subir", files=[("archivos", ("datos editados.xlsx", buf.getvalue()))],
                         params={"cliente_id": c["id"]}).json()
    assert p["clase"] == "contabilidad"
    imp = cliente_api.post(f"/api/subir/{p['subida_id']}/confirmar", json={"cliente_id": c["id"]}).json()
    peticion = {"sesion_id": imp["sesion_id"], "cliente_id": c["id"],
                "mapeo": {m["normalizado"]: m["codigo"] for m in imp["mapeo"] if m.get("codigo")}}
    r = cliente_api.post("/api/calcular", json=peticion)
    assert r.status_code == 200, r.text
    assert r.json()["periodo"]["id"] == pid            # el MISMO periodo, no uno nuevo
    assert D(r.json()["resumen"]["utilidad_neta"]) == D(calculado["resumen"]["utilidad_neta"]) - 10000


def test_archivos_del_periodo_ver_y_quitar(cliente_api):
    _, calculado = _caso(cliente_api, "completo", "900100158")
    pid = calculado["periodo"]["id"]
    archivos = _datos(cliente_api, pid)["archivos"]
    assert archivos and archivos[0]["filas"] > 0
    r = cliente_api.post(f"/api/periodos/{pid}/recalcular")
    assert r.status_code == 200 and r.json()["guardado"] is True
