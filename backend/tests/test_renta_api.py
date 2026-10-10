"""Declaración de renta de punta a punta por la API (v2.3 · Fase 5), con el caso A anonimizado."""
import io
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.renta.lectura import interpretar_filas, validar_sumas

CASO_A = json.loads((Path(__file__).parent / "renta" / "caso_a.json").read_text(encoding="utf-8"))


def excel_exogena(caso: dict) -> bytes:
    """El reporte de exógena del caso como lo descarga el portal (Excel)."""
    wb = Workbook()
    ws = wb.active
    d = caso["documento"]
    ws.append(["Tipo de documento:", d["tipo"]])
    ws.append(["Identificación:", d["numero"]])
    ws.append(["Nombres / Razón social:", d["nombre"]])
    nombres = {1: "Ingresos", 2: "Patrimonio", 3: "Consumo TC", 4: "Movimiento", 5: "Compras"}
    for n, v in caso["topes"].items():
        ws.append(["", "", f"Tope {n} - {nombres[int(n)]}", int(v), ""])
    for l in caso["lineas"]:
        ws.append([l["entidad"], l["titular"], l["detalle"], int(l["valor"]), l["uso"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cliente_natural(api, nit="10000001", nombre="CONTRIBUYENTE A"):
    r = api.post("/api/clientes", json={"nit": nit, "razon_social": nombre, "tipo_persona": "natural",
                                       "regimen": "no_responsable_iva"})
    assert r.status_code in (200, 201), r.text
    return r.json()["id"]


def _subir(api, cid, nombre="exogena.xlsx", contenido=None):
    contenido = contenido or excel_exogena(CASO_A)
    r = api.post(f"/api/renta/{cid}/2025/documentos", files=[("archivos", (nombre, contenido, "application/octet-stream"))])
    assert r.status_code == 200, r.text
    return r.json()


def test_caso_a_de_punta_a_punta(cliente_api):
    cid = _cliente_natural(cliente_api)
    v = _subir(cliente_api, cid)
    res = v["resultado"]
    assert v["estado"] == "borrador"
    assert len(res["lineas"]) == 23
    # Obligada solo por el Tope 4 (consignaciones).
    assert res["obligacion"]["obligado"]
    assert [m["tope"] for m in res["obligacion"]["motivos"] if m["supera"]] == ["consignaciones"]
    # Las sumas: todo cuadra; el Tope 4 trae 3 pesos de diferencia de la propia fuente.
    estados = {x["tope"]: x["estado"] for x in res["validacion"]}
    assert estados[2] == "cuadra" and estados[3] == "cuadra" and estados[4] == "redondeo"
    # Las 2 líneas «NO REGISTRA NO» quedan marcadas y hay pregunta para ellas.
    assert sum(1 for l in res["lineas"] if l["no_titular"]) == 2
    assert any(q["id"] == "no_titular" for q in res["preguntas"])
    assert len(res["preguntas"]) <= 5
    # Aviso: consignaciones (≈ 93 M) muy por encima de la facturación electrónica (≈ 4,8 M).
    assert any("consignaciones" in m["texto"].lower() and "facturación electrónica" in m["texto"] for m in res["marcas"])

    # Lo que agregó el contador por conocer el negocio.
    for categoria, valor in (("ingreso_no_laboral", "65400000"), ("costo_no_laboral", "32800000")):
        r = cliente_api.post(f"/api/renta/{cid}/2025/cambio",
                             json={"tipo": "agregar", "categoria": categoria, "valor": valor, "descripcion": "Negocio"})
        assert r.status_code == 200, r.text
    # La app propone el máximo legal del 1 % y deja ajustarlo: el contador declaró 48.000.
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio",
                         json={"tipo": "beneficio", "id": "compras_fe", "valor": True, "uno_por_ciento": "48000"})
    res = r.json()["resultado"]
    assert D(res["maximo_1pct"]) == D("78000")
    casillas = {c["casilla"]: D(c["optimizada"]) for c in res["casillas"]}
    for n, valor in CASO_A["esperado_210"].items():
        assert casillas[int(n)] == D(valor), f"casilla {n}: {casillas[int(n)]} ≠ {valor}"
    assert res["motivos_incompleto"] == []
    assert D(res["cifras"]["a_favor"]) == D("5000")


def test_reclasificar_y_excluir_recalculan(cliente_api):
    cid = _cliente_natural(cliente_api)
    res = _subir(cliente_api, cid)["resultado"]
    deuda = next(l for l in res["lineas"] if l["categoria"] == "deuda")
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "excluir", "linea": deuda["id"]})
    casillas = {c["casilla"]: D(c["optimizada"]) for c in r.json()["resultado"]["casillas"]}
    # 52.201.487 − 22.900 (la primera deuda) = 52.178.587 → 52.179.000
    assert deuda["valor"] == "22900" and casillas[30] == D("52179000")
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "reclasificar", "linea": deuda["id"],
                                                                 "categoria": "informativo"})
    assert r.status_code == 200
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "reclasificar", "linea": deuda["id"],
                                                                 "categoria": "no_existe"})
    assert r.status_code == 422


def test_persona_juridica_no_aplica(cliente_api):
    cid = cliente_api.post("/api/clientes", json={"nit": "900123456", "razon_social": "EMPRESA S.A.S."}).json()["id"]
    r = cliente_api.get(f"/api/renta/{cid}/2025")
    assert r.status_code == 409
    assert "110" in r.json()["detail"]["mensaje"]


def _negocio(api, cid, ingresos="65400000", costos="32800000"):
    r = api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "negocio", "ingresos": ingresos, "costos": costos})
    assert r.status_code == 200, r.text
    return r.json()["resultado"]


def test_descargas(cliente_api):
    cid = _cliente_natural(cliente_api)
    _subir(cliente_api, cid)
    _negocio(cliente_api, cid)
    _negocio(cliente_api, cid)
    pdf = cliente_api.get(f"/api/renta/{cid}/2025/descargar/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    resumen = cliente_api.get(f"/api/renta/{cid}/2025/descargar/resumen")
    assert resumen.status_code == 200 and resumen.content.startswith(b"%PDF")
    xls = cliente_api.get(f"/api/renta/{cid}/2025/descargar/excel")
    assert xls.status_code == 200
    wb = load_workbook(io.BytesIO(xls.content))
    assert wb.sheetnames == ["Formulario 210", "Líneas del reporte", "De dónde sale cada peso", "Avisos"]
    ws = wb["Formulario 210"]
    assert ws["A1"].value.endswith("BORRADOR")
    # Ninguna celda de texto empieza por «=»: no hay inyección de fórmulas.
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows(values_only=True):
            assert not any(isinstance(c, str) and c.startswith("=") for c in fila)


def test_marcar_presentada_y_cartera(cliente_api):
    cid = _cliente_natural(cliente_api)
    _subir(cliente_api, cid)
    _negocio(cliente_api, cid)
    r = cliente_api.post(f"/api/renta/{cid}/2025/estado", json={"estado": "presentada"})
    assert r.status_code == 422
    r = cliente_api.post(f"/api/renta/{cid}/2025/estado",
                         json={"estado": "presentada", "numero": "2118000000001", "fecha": "2026-09-07"})
    assert r.status_code == 200 and r.json()["estado"] == "presentada"
    c = cliente_api.get("/api/renta/cartera?anio=2025").json()
    fila = next(x for x in c["declaraciones"] if x["cliente_id"] == cid)
    assert fila["estado"] == "presentada" and fila["obligado"] is True


def test_tarea_en_el_tablero_si_vence(cliente_api):
    from app.inteligencia import tablero

    cid = _cliente_natural(cliente_api, nit="10000020", nombre="CONTRIBUYENTE VENCIDO")   # 19-20 → 26 de agosto
    t = tablero.armar(date(2026, 10, 8))
    tarea = next(x for x in t["tareas"] if x["codigo"] == "RENTA_VENCE" and x["cliente_id"] == cid)
    assert tarea["prioridad"] == "critica" and "hace 43 días" in tarea["titulo"]
    assert tarea["accion"]["ruta"] == f"/clientes/{cid}?seccion=renta"
    # Lejos del vencimiento no hay tarea.
    t = tablero.armar(date(2026, 7, 1))
    assert not any(x["codigo"] == "RENTA_VENCE" and x["cliente_id"] == cid for x in t["tareas"])


def test_sin_sesion_no_hay_renta(cliente_sin_sesion):
    assert cliente_sin_sesion.get("/api/renta/cartera").status_code == 401


def test_lectura_ruidosa_se_corrige_con_los_topes():
    """Errores típicos del OCR («$» leído como «5») se resuelven cuadrando con el encabezado."""
    filas = [["", "", "Tope 3 - Consumo TC", "$14,536,332.00", ""],
             ["", "", "Tope 5 - Compras", "$57,763,109.00", ""],
             ["BANCO A", "CONTRIBUYENTE A", "Total consumos o gastos con tarjeta", "$ 1,170,568.00", "Tope 3: Consumos TC"],
             ["COMPAÑIA", "CONTRIBUYENTE A", "Total consumos o gastos con tarjeta", "54,932,351.00", "Tope 3: Consumos TC"],
             ["BANCO A", "CONTRIBUYENTE A", "Total consumos o gastos con tarjeta", "$ 8,413,792.00", "Tope 3: Consumos TC"],
             ["BANCO A", "CONTRIBUYENTE A", "Total consumos o gastos con tarjeta", "$ 19,621.00", "Tope 3: Consumos TC"],
             ["DIAN", "CONTRIBUYENTE A", "Suma valor total facturas tras ajustes por notas", "$ 7,763,109.00",
              "Tope 5: Compras registradas"]]
    rep = interpretar_filas(filas, documento="foto.jpg")
    assert rep.topes[5] == "57763109"
    val = {v["tope"]: v for v in validar_sumas(rep)}
    assert val[3]["estado"] == "corregida" and val[5]["estado"] == "corregida"
    assert rep.topes[5] == "7763109"
    corregida = next(l for l in rep.lineas if l.corregida)
    assert corregida.valor == "4932351" and "54932351" in corregida.alternativas


def test_un_boton_descarga_los_tres(cliente_api):
    import zipfile

    cid = _cliente_natural(cliente_api)
    _subir(cliente_api, cid)
    _negocio(cliente_api, cid)
    r = cliente_api.get(f"/api/renta/{cid}/2025/descargar/todo")
    assert r.status_code == 200
    nombres = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert len(nombres) == 3 and any(n.endswith("borrador-210.pdf") for n in nombres)


def test_comerciante_pide_ingresos_y_costos_del_negocio(cliente_api):
    """Rescate H6: consignaciones 93,1 M y facturación 4,8 M → la app pide lo que solo el contador sabe."""
    cid = _cliente_natural(cliente_api)
    res = _subir(cliente_api, cid)["resultado"]
    assert res["incompleto"] is True and res["cifras"]["neto"] is None
    tipos = [p["tipo"] for p in res["pendientes"]]
    assert "negocio" in tipos
    caja = next(p for p in res["pendientes"] if p["tipo"] == "negocio")
    assert "Consignaciones 93,1 M" in caja["comparacion"] and "Facturación 4,8 M" in caja["comparacion"]
    # Con los datos del negocio y el 1 % de compras: las casillas del 210 presentado.
    _negocio(cliente_api, cid)
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio",
                         json={"tipo": "beneficio", "id": "compras_fe", "valor": True, "uno_por_ciento": "48000"})
    res = r.json()["resultado"]
    assert not any(p["tipo"] == "negocio" for p in res["pendientes"])
    casillas = {c["casilla"]: D(c["optimizada"]) for c in res["casillas"]}
    for n, valor in CASO_A["esperado_210"].items():
        assert casillas[int(n)] == D(valor), f"casilla {n}: {casillas[int(n)]} ≠ {valor}"
    # Las dos columnas explicadas en una línea.
    d74 = next(d for d in res["diferencias"] if d["casilla"] == 74)
    assert "facturación" in d74["por_que"] and "ingresos reales del negocio" in d74["por_que"]


def test_sin_negocio_se_puede_decir_y_no_bloquea(cliente_api):
    cid = _cliente_natural(cliente_api)
    _subir(cliente_api, cid)
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "negocio", "sin_negocio": True})
    assert not any(p["tipo"] == "negocio" for p in r.json()["resultado"]["pendientes"])


def test_ajuste_manual_de_casilla_con_nota(cliente_api):
    cid = _cliente_natural(cliente_api)
    _subir(cliente_api, cid)
    _negocio(cliente_api, cid)
    sin_nota = cliente_api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "ajuste_casilla", "casilla": 132, "valor": "9000", "nota": " "})
    assert sin_nota.status_code == 422 and "nota" in sin_nota.json()["detail"]["mensaje"]
    formula = cliente_api.post(f"/api/renta/{cid}/2025/cambio",
                               json={"tipo": "ajuste_casilla", "casilla": 93, "valor": "1", "nota": "x"})
    assert formula.status_code == 422 and "se calcula" in formula.json()["detail"]["mensaje"]
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio",
                         json={"tipo": "ajuste_casilla", "casilla": 132, "valor": "9000", "nota": "Certificado del banco B"})
    res = r.json()["resultado"]
    c132 = next(c for c in res["casillas"] if c["casilla"] == 132)
    assert D(c132["optimizada"]) == D("9000") and c132["ajuste"]["nota"] == "Certificado del banco B"
    assert res["ajustes_casilla"][0]["casilla"] == 132
    r = cliente_api.post(f"/api/renta/{cid}/2025/cambio", json={"tipo": "quitar_ajuste_casilla", "casilla": 132})
    c132 = next(c for c in r.json()["resultado"]["casillas"] if c["casilla"] == 132)
    assert D(c132["optimizada"]) == D("5000")


def test_pdf_borrador_solo_casillas_con_valor_y_anexo(cliente_api):
    """Rescate §4.5: el borrador completo muestra solo las casillas con valor; el anexo trae todas.
    El incompleto empieza por la lista de lo que falta."""
    import pdfplumber

    cid = _cliente_natural(cliente_api)
    _subir(cliente_api, cid)
    pdf = cliente_api.get(f"/api/renta/{cid}/2025/descargar/pdf").content
    with pdfplumber.open(io.BytesIO(pdf)) as d:
        primera = d.pages[0].extract_text()
    assert "Para terminar faltan" in primera and "Ingresos y costos del negocio" in primera
    _negocio(cliente_api, cid)
    pdf = cliente_api.get(f"/api/renta/{cid}/2025/descargar/pdf").content
    with pdfplumber.open(io.BytesIO(pdf)) as d:
        textos = [pg.extract_text() or "" for pg in d.pages]
    corte = next(i for i, t in enumerate(textos) if "Anexo · todas las casillas" in t)
    cuerpo, anexo = "\n".join(textos[:corte]), "\n".join(textos[corte:])
    assert "Resumen" in cuerpo and "Renta líquida gravable" in cuerpo and "Por qué su declaración" in cuerpo
    # Las ganancias ocasionales están en cero: no van en el cuerpo, sí en el anexo completo.
    assert "ganancias ocasionales" not in cuerpo.lower() and "ganancias ocasionales" in anexo.lower()


def test_facturas_dian_en_la_renta_alimentan_el_negocio(cliente_api):
    """Idea central: el Excel de facturas del portal de la DIAN se reconoce solo también en la renta."""
    from demo import dian_ficticio as F

    cid = _cliente_natural(cliente_api, nit=F.CLIENTE["nit"], nombre="COMERCIANTE FICTICIA")
    _subir(cliente_api, cid)
    r = cliente_api.post(f"/api/renta/{cid}/2025/documentos",
                         files=[("archivos", ("facturas_todas.xlsx", F.archivos()["facturas_todas_enero_2025.xlsx"],
                                              "application/octet-stream"))])
    assert r.status_code == 200, r.text
    res = r.json()["resultado"]
    assert D(res["facturas"]["ventas_base"]) == D("3300000") and D(res["facturas"]["compras_base"]) == D("1250000")
    caja = next(p for p in res["pendientes"] if p["tipo"] == "negocio")
    assert "Facturación" in caja["comparacion"]
    assert len(res["lineas"]) == 23          # la exógena sigue intacta


def test_exogena_en_excel_no_pide_confirmar_filas(cliente_api):
    """Un Excel trae el dato exacto: no se pide «confirmar filas» aunque no traiga los topes."""
    from demo.fotos_ficticias import exogena_persona_nueva

    cid = _cliente_natural(cliente_api, nit="10000009", nombre="PERSONA NUEVA FICTICIA")
    res = _subir(cliente_api, cid, "exogena.xlsx", exogena_persona_nueva())["resultado"]
    assert not any(p["tipo"] == "confirmar_filas" for p in res["pendientes"])
