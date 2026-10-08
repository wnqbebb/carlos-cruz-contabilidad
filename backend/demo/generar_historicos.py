"""Cinco clientes históricos de demostración, de enero de 2025 a septiembre de 2026 (spec v2.2 · fase 9).

Todo es ficticio: nombres, NIT (con su dígito de verificación válido), cédulas,
terceros y cifras. Cada cliente se crea y se carga POR LA API, como lo haría el
contador: subir archivos → identificar → calcular → cerrar, mes a mes.

    1  PANADERÍA LA ESPIGA DORADA S.A.S. (Tuluá)       plantilla oficial, kardex de insumos, nómina de 4,
                                                       IVA, pico de diciembre. Al día y cerrado.
    2  FERRETERÍA EL TORNILLO S.A.S. (Buga)            plantilla, ventas a crédito con cartera creciente,
                                                       proveedores, depreciación. Agosto 2026 calculado sin
                                                       cerrar; septiembre sin subir.
    3  CLÍNICA DENTAL SONRISA DEL VALLE S.A.S. (Cali)  cuenta T + hoja de trabajo, servicios de salud,
                                                       honorarios con retención, arriendo. Septiembre 2026
                                                       con un descuadre de $50.000 a propósito.
    4  TRANSPORTES RÍO CAUCA S.A.S. (Palmira)          un libro diario en CSV con los 21 meses: vehículos,
                                                       préstamo con intereses, depreciación, meses con
                                                       pérdida; el patrimonio cae por debajo del 50 % del capital.
    5  MARÍA ELENA ROJAS (persona natural, Guacarí)    registros auxiliares desordenados de su finca;
                                                       periodicidad bimestral; atrasada 3 meses (hasta junio 2026).

Los clientes 1 y 2 se dan de alta subiendo sus estatutos (.docx) y su RUT (.pdf),
generados aquí. Todos quedan con la etiqueta «Demostración» y la marca `demo`, que
es lo que usa «Eliminar clientes de demostración» (menú de la cuenta › Sistema).

Reproducible: azar con semilla fija. Algunos pagos (aportes, IVA, retenciones,
prestaciones) se toman del cierre del mes anterior que devuelve la propia
aplicación, como haría el contador; con la misma versión de la aplicación los
archivos salen idénticos.

Uso, desde la carpeta backend/:
    python -m demo.generar_historicos --url http://127.0.0.1:8001
    python -m demo.generar_historicos --url http://localhost:8000 --usuario carlos --clave '…'
    python -m demo.generar_historicos --solo 1,3          # algunos clientes
Los archivos quedan en docs/demo/<cliente>/.
"""
from __future__ import annotations

import argparse
import csv
import io
import random
import re
import sys
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import httpx
from openpyxl import Workbook

from app.contabilidad.puc import puc
from app.exportar.plantilla import HOJAS
from app.utils.nit import digito_verificacion
from app.utils.numeros import normalizar

RAIZ = Path(__file__).resolve().parents[2]
SALIDA = RAIZ / "docs" / "demo"
INICIO, FIN = (2025, 1), (2026, 9)
MESES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]
ETIQUETA = "Demostración"
NOTA = "Cliente de demostración: todos los datos son ficticios. Se borra desde el menú de la cuenta › Sistema."
CERO = Decimal(0)


class ErrorCarga(RuntimeError):
    pass


# ── utilidades ──────────────────────────────────────────────────────────────
def meses(desde=INICIO, hasta=FIN):
    a, m = desde
    while (a, m) <= hasta:
        yield a, m
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)


def ultimo_dia(a: int, m: int) -> date:
    return (date(a + (m == 12), m % 12 + 1, 1) - timedelta(days=1))


def r100(x) -> int:
    return int(round(float(x) / 100.0)) * 100


def repartir(total: int, partes: int) -> list[int]:
    base = total // partes
    return [base + (1 if i < total - base * partes else 0) for i in range(partes)]


def celda(v):
    """Decimal → número de Excel sin perder centavos (repr de un float de 2 decimales es exacto)."""
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    return v


def nombre_puc(codigo: str) -> str:
    return puc().nombre(codigo) or codigo


def yymm(a: int, m: int) -> str:
    return f"{a % 100:02d}{m:02d}"


def neto(saldos: list[dict]) -> dict[str, Decimal]:
    """saldos_siguiente de la API → {código: débito − crédito}."""
    return {s["codigo"]: Decimal(str(s["debito"] or 0)) - Decimal(str(s["credito"] or 0)) for s in saldos or []}


class Diario:
    """Asientos de la hoja MOVIMIENTOS de la plantilla. Cada asiento debe cuadrar."""

    def __init__(self):
        self.filas: list[list] = []

    def asiento(self, fecha: date, comp: str, tipo: str, desc: str, lineas: list[tuple], tercero=("", "")):
        d = sum((Decimal(str(l[1] or 0)) for l in lineas), CERO)
        c = sum((Decimal(str(l[2] or 0)) for l in lineas), CERO)
        if d != c:
            raise AssertionError(f"Asiento {comp} no cuadra: {d} ≠ {c}")
        for linea in lineas:
            cod, deb, cre = linea[:3]
            detalle = linea[3] if len(linea) > 3 and linea[3] else desc
            nit, nom = tercero
            self.filas.append([fecha, comp, tipo, cod, nombre_puc(cod), nit, nom, detalle,
                               celda(deb) if deb else None, celda(cre) if cre else None, None])


FECHA_FIJA = (2026, 1, 1, 0, 0, 0)


def fijar_zip(datos: bytes) -> bytes:
    """Excel y Word son zips con fechas adentro: se fijan para que el archivo salga idéntico byte a byte."""
    entrada, salida = zipfile.ZipFile(io.BytesIO(datos)), io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for item in entrada.infolist():
            contenido = entrada.read(item.filename)
            if item.filename == "docProps/core.xml":
                contenido = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", contenido)
            z.writestr(zipfile.ZipInfo(item.filename, date_time=FECHA_FIJA), contenido, compress_type=zipfile.ZIP_DEFLATED)
    return salida.getvalue()


def libro(hojas: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append([celda(v) for v in f])
    buf = io.BytesIO()
    wb.save(buf)
    return fijar_zip(buf.getvalue())


def plantilla(empresa: list[tuple], hojas: dict[str, list[list]]) -> bytes:
    """Plantilla oficial con las hojas que traiga el mes (encabezados de exportar/plantilla.py)."""
    salida = {"EMPRESA": [["Campo", "Valor"], *[list(x) for x in empresa]]}
    for nombre in ("SALDOS_INICIALES", "MOVIMIENTOS", "INVENTARIO_MOVS", "ACTIVOS_FIJOS", "NOMINA"):
        if hojas.get(nombre):
            salida[nombre] = [HOJAS[nombre], *hojas[nombre]]
    return libro(salida)


def guardar(cliente: str, nombre: str, datos: bytes) -> tuple[str, bytes]:
    carpeta = SALIDA / cliente
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / nombre).write_bytes(datos)
    return nombre, datos


# ── documentos de constitución (Word) y RUT (PDF) ───────────────────────────
def estatutos_docx(c: dict) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(f"ESTATUTOS DE {c['razon']}", level=1)
    doc.add_paragraph("Documento ficticio, generado para la demostración de la aplicación. No corresponde a ninguna empresa real.")
    doc.add_paragraph(
        f"En el municipio de {c['municipio']}, a los {c['dia_letras']} ({c['constitucion'].day}) días del mes de "
        f"{MESES[c['constitucion'].month]} del año {c['anio_letras']} ({c['constitucion'].year}), los accionistas abajo "
        f"firmantes constituyen, por DOCUMENTO PRIVADO No. {c['documento']}, una sociedad por acciones simplificada.")
    doc.add_paragraph(f"ARTÍCULO 1. Nombre. La sociedad se denominará {c['razon']} y podrá usar la sigla \"{c['sigla']}\".")
    doc.add_paragraph(f"ARTÍCULO 2. Domicilio. El domicilio principal de la sociedad será el municipio de {c['municipio']} Valle. "
                      f"La dirección para notificaciones judiciales será {c['direccion']}; correo {c['email']}, "
                      f"teléfono {c['telefono']}.")
    doc.add_paragraph(f"ARTÍCULO 3. Objeto social: {c['objeto']}. Actividades económicas CIIU {c['ciiu']} y {c['ciiu2']}.")
    doc.add_paragraph(f"ARTÍCULO 4. Capital. El capital autorizado de la sociedad es de ${c['autorizado']:,} M/cte. "
                      f"El capital suscrito y pagado es de ${c['capital']:,} M/cte, dividido en {c['acciones_letras']} "
                      f"({c['capital'] // 1000:,}) acciones con valor nominal de mil pesos ($1.000) cada una."
                      .replace(",", "."))
    doc.add_paragraph(
        f"ARTÍCULO 5. Representación legal. Los accionistas designan a {c['rep']}, identificada con la cédula de "
        f"ciudadanía No. {c['rep_cc_fmt']}, como representante legal principal, y {c['sup']} como representante legal "
        f"suplente con cédula No. {c['sup_cc_fmt']}.")
    doc.add_paragraph("ARTÍCULO 6. Revisoría fiscal. La revisoría fiscal solo será provista cuando la ley lo exija.")
    doc.add_paragraph("Composición accionaria:")
    t = doc.add_table(rows=1, cols=5)
    for i, txt in enumerate(["NOMBRE", "CÉDULA", "ACCIONES", "VALOR TOTAL", "%"]):
        t.rows[0].cells[i].text = txt
    for nombre, cc, acciones in c["socios"]:
        fila = t.add_row().cells
        pct = Decimal(acciones * 100) / Decimal(c["capital"] // 1000)
        for i, txt in enumerate([nombre, cc, f"{acciones:,}".replace(",", "."), f"{acciones * 1000:,}".replace(",", "."),
                                 f"{pct:.0f}%"]):
            fila[i].text = txt
    doc.add_paragraph("Firman los accionistas en constancia.")
    buf = io.BytesIO()
    doc.save(buf)
    return fijar_zip(buf.getvalue())


def rut_pdf(c: dict) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=letter, invariant=1)  # sin fecha ni identificador al azar
    y = 740
    lineas = [
        ("Helvetica-Bold", 13, "REGISTRO ÚNICO TRIBUTARIO — COPIA FICTICIA PARA DEMOSTRACIÓN"),
        ("Helvetica", 8, "Documento inventado para probar la aplicación. No es un formulario de la DIAN ni identifica a nadie."),
        ("Helvetica", 10, ""),
        ("Helvetica", 10, f"Número de Identificación Tributaria (NIT): {c['nit']} - {digito_verificacion(c['nit'])}"),
        ("Helvetica", 10, "Tipo de contribuyente: Persona jurídica"),
        ("Helvetica", 10, f"Razón social: {c['razon']}"),
        ("Helvetica", 10, f"Dirección principal: {c['direccion']}"),
        ("Helvetica", 10, f"Municipio / Ciudad: {c['municipio']}"),
        ("Helvetica", 10, "Departamento: Valle del Cauca"),
        ("Helvetica", 10, f"Correo electrónico: {c['email']}"),
        ("Helvetica", 10, f"Teléfono 1: {c['telefono']}"),
        ("Helvetica", 10, f"Actividad principal: {c['ciiu']}"),
        ("Helvetica", 10, f"Actividad secundaria: {c['ciiu2']}"),
        ("Helvetica", 10, "Responsabilidades:"),
        ("Helvetica", 10, "05 - Impuesto de renta y complementarios régimen ordinario"),
        ("Helvetica", 10, "07 - Retención en la fuente a título de renta"),
        ("Helvetica", 10, "48 - Impuesto sobre las ventas - IVA"),
        ("Helvetica", 10, "52 - Facturador electrónico"),
    ]
    for fuente, tam, texto in lineas:
        pdf.setFont(fuente, tam)
        pdf.drawString(50, y, texto)
        y -= 20
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


# ── API ─────────────────────────────────────────────────────────────────────
class Api:
    def __init__(self, url: str, usuario: str = "", clave: str = ""):
        self.c = httpx.Client(base_url=url.rstrip("/"), timeout=900)
        if usuario:
            self._ok(self.c.post("/api/sesion", json={"usuario": usuario, "clave": clave}))

    @staticmethod
    def _ok(r: httpx.Response):
        if r.status_code >= 400:
            raise ErrorCarga(f"{r.request.method} {r.request.url.path} → {r.status_code}: {r.text[:2000]}")
        return r.json()

    def get(self, ruta: str, **params):
        return self._ok(self.c.get(ruta, params=params))

    def post(self, ruta: str, cuerpo=None, **params):
        return self._ok(self.c.post(ruta, json=cuerpo if cuerpo is not None else {}, params=params))

    def subir(self, archivos: list[tuple[str, bytes]], cliente_id: str = ""):
        files = [("archivos", (n, b)) for n, b in archivos]
        return self._ok(self.c.post("/api/subir", files=files, params={"cliente_id": cliente_id}))


def _mapeo(items: list[dict], alias: dict[str, str]) -> dict[str, str]:
    """Lo que haría el contador en «Revisar»: confirma lo que el sistema ya sabe y elige el resto."""
    salida = {}
    for it in items:
        n = it["normalizado"]
        if n in alias:
            salida[n] = alias[n]
        elif it.get("codigo"):
            salida[n] = it["codigo"]
        else:
            raise ErrorCarga(f"Nombre de cuenta sin código: «{it.get('nombre') or n}»")
    return salida


def cargar(api: Api, archivos, *, cliente_id: str | None = None, crear: dict | None = None,
           alias: dict[str, str] | None = None, periodo: tuple[date, date] | None = None,
           por_periodo: bool = False, cerrar: bool = True) -> tuple[str, dict]:
    """subir → confirmar (crea el cliente la primera vez) → calcular → cerrar."""
    sub = api.subir(archivos, cliente_id or "")
    imp = api.post(f"/api/subir/{sub['subida_id']}/confirmar", {"cliente_id": cliente_id} if cliente_id else {"crear": crear})
    if imp.get("clase") != "contabilidad":
        raise ErrorCarga(f"La subida no se reconoció como contabilidad: {imp.get('clase')}")
    pet = {
        "sesion_id": imp["sesion_id"], "cliente_id": imp["cliente_id"], "incluir": {},
        "mapeo": _mapeo(imp["mapeo"], {normalizar(k): v for k, v in (alias or {}).items()}),
        "empresa": {"periodo_desde": periodo[0].isoformat(), "periodo_hasta": periodo[1].isoformat()} if periodo else {},
        "config": {}, "decisiones": {}, "recordar_alias": True,
        "respuestas": {p["id"]: p["defecto"] for p in imp.get("preguntas") or []},
    }
    if por_periodo:
        pet["periodizacion"], pet["cerrar_ultimo"] = "por_periodo", cerrar
    res = api.post("/api/calcular", pet)
    if cerrar and not por_periodo:
        api.post(f"/api/cierre/{imp['sesion_id']}")
    return imp["cliente_id"], res


def informe(nombre: str, a: int, m: int, res: dict, estado: str) -> None:
    p = res.get("periodo") or {}
    cuadra = "cuadra" if p.get("cuadra") else "NO CUADRA"
    print(f"  {nombre:<11} {a}-{m:02d}  utilidad {Decimal(str(res['resumen']['utilidad_neta'])):>16,.2f}  {cuadra:<9}  {estado}")


def ficha_base(nit: str, razon: str, **extra) -> dict:
    return {"nit": nit, "razon_social": razon, "demo": True, "etiquetas": [ETIQUETA], "notas": NOTA, **extra}


# ── pagos que salen del cierre anterior (como los haría el contador) ───────
def pagos_del_mes(diario: Diario, a: int, m: int, previos: dict[str, Decimal], nombres: dict[str, str],
                  cierre_dic: dict[str, Decimal]) -> None:
    if not previos:
        return
    pila = [(c, -v) for c, v in sorted(previos.items()) if c.startswith(("2370", "2380")) and v < 0]
    if pila:
        diario.asiento(date(a, m, 10), f"CE-{yymm(a, m)}-PILA", "Egreso", "Pago planilla PILA del mes anterior",
                       [(c, v, None) for c, v in pila] + [("111005", None, sum(v for _, v in pila))])
    ret = [(c, -v) for c, v in sorted(previos.items()) if c.startswith("2365") and v < 0]
    if ret:
        diario.asiento(date(a, m, 15), f"CE-{yymm(a, m)}-RTE", "Egreso", "Pago retención en la fuente (formulario 350)",
                       [(c, v, None) for c, v in ret] + [("111005", None, sum(v for _, v in ret))])
    if m in (1, 3, 5, 7, 9, 11):
        generado = -previos.get("240805", CERO)
        descontable = previos.get("240810", CERO)
        if generado > descontable > 0:
            diario.asiento(date(a, m, 18), f"CE-{yymm(a, m)}-IVA", "Egreso", "Pago IVA del bimestre anterior (formulario 300)",
                           [("240805", generado, None), ("240810", None, descontable),
                            ("111005", None, generado - descontable)])

    def pagar(prefijo_nombre: str, excluir: str, desde: dict, dia: int, texto: str):
        lineas = [(c, -v) for c, v in sorted(desde.items())
                  if c[:2] in ("25", "26") and v < 0 and prefijo_nombre in normalizar(nombres.get(c, nombre_puc(c)))
                  and not (excluir and excluir in normalizar(nombres.get(c, nombre_puc(c))))]
        if lineas:
            diario.asiento(date(a, m, dia), f"CE-{yymm(a, m)}-{prefijo_nombre[:3]}", "Egreso", texto,
                           [(c, v, None) for c, v in lineas] + [("111005", None, sum(v for _, v in lineas))])

    if m in (1, 7):
        pagar("PRIMA", "", previos, 8, "Pago prima de servicios del semestre")
    if m == 1 and cierre_dic:
        pagar("INTERESES", "", cierre_dic, 28, "Pago intereses sobre cesantías del año anterior")
    if m == 2 and cierre_dic:
        pagar("CESANTIAS", "INTERES", cierre_dic, 13, "Consignación de cesantías al fondo")


def nomina_diario(diario: Diario, a: int, m: int, empleados, aux: int) -> list[list]:
    """Pago de la nómina en el diario (devengado y deducciones del trabajador) y hoja NOMINA."""
    fin = ultimo_dia(a, m)
    sueldos = sum(e[3][a - 2025] for e in empleados)
    salud = sum(r100(e[3][a - 2025] * 0.04) for e in empleados)
    pension = salud
    total_aux = aux * len(empleados)
    diario.asiento(fin, f"CE-{yymm(a, m)}-NOM", "Egreso", f"Nómina de {MESES[m]} de {a}",
                   [("510506", sueldos, None, "Sueldos del mes"), ("510527", total_aux, None, "Auxilio de transporte"),
                    ("237005", None, salud, "Salud a cargo del trabajador (4 %)"),
                    ("238030", None, pension, "Pensión a cargo del trabajador (4 %)"),
                    ("111005", None, sueldos + total_aux - salud - pension, "Pago neto de la nómina")])
    return [[date(a, m, 1), e[0], e[1], e[2], e[3][a - 2025], None, None, 30, "SI", 0, 0, e[4]] for e in empleados]


def kardex_mes(inv: list[list], a: int, m: int, productos, stock: dict[str, int], ventas_q: dict[str, list[int]],
               compras_q: dict[str, int], dias_venta: list[int], doc_venta: str, tipo_salida: str,
               precio: bool) -> None:
    si = date(a, m, 1) - timedelta(days=1)
    for p in productos:
        cod, desc, costo = p[0], p[1], p[2]
        if stock[cod]:
            inv.append([si, "SI", cod, desc, "", "", None, "Inventario inicial", stock[cod], costo, None])
        if compras_q.get(cod):
            inv.append([date(a, m, 3), f"FC-{yymm(a, m)}", cod, desc, "", "", None, "Compra", compras_q[cod], costo, None])
        for i, (dia, q) in enumerate(zip(dias_venta, ventas_q[cod])):
            if q:
                inv.append([date(a, m, dia), f"{doc_venta}-{yymm(a, m)}-{i + 1}", cod, desc, "", "", None, tipo_salida, q,
                            None, p[3] if precio else None])
        stock[cod] += compras_q.get(cod, 0) - sum(ventas_q[cod])


def dias_semana(a: int, m: int) -> list[int]:
    return [7, 14, 21, ultimo_dia(a, m).day]


# ═══════════════════════════════════════════════════════════════════════════
# 1 · PANADERÍA LA ESPIGA DORADA S.A.S. — Tuluá
# ═══════════════════════════════════════════════════════════════════════════
PANADERIA = {
    "clave": "1-panaderia", "razon": "PANADERÍA LA ESPIGA DORADA S.A.S.", "sigla": "ESPIGA DORADA", "nit": "901482317",
    "municipio": "Tuluá", "direccion": "Calle 27 No. 24-15 barrio Centro", "email": "espigadorada.demo@example.com",
    "telefono": "3157778899", "ciiu": "1081", "ciiu2": "4724",
    "objeto": "la elaboración y venta de productos de panadería, pastelería y repostería, y la venta de bebidas",
    "constitucion": date(2019, 3, 14), "dia_letras": "catorce", "anio_letras": "dos mil diecinueve", "documento": "001",
    "autorizado": 100_000_000, "capital": 65_800_000, "acciones_letras": "sesenta y cinco mil ochocientas",
    "rep": "LUZ MARINA OSPINA RENDON", "rep_cc": "31456789", "rep_cc_fmt": "31.456.789",
    "sup": "JORGE ELIECER MONTOYA DIAZ", "sup_cc": "94367120", "sup_cc_fmt": "94.367.120",
    "socios": [("LUZ MARINA OSPINA RENDON", "31.456.789", 39_480), ("JORGE ELIECER MONTOYA DIAZ", "94.367.120", 26_320)],
}
INSUMOS = [  # código, descripción, costo, consumo base al mes, existencia que se busca al cierre
    ("H01", "Harina de trigo bulto 50 kg", 115_000, 40, 15),
    ("A01", "Azúcar refinada bulto 50 kg", 165_000, 12, 5),
    ("M01", "Margarina industrial caja 15 kg", 210_000, 8, 4),
    ("HU1", "Huevo AA cubeta x30", 14_000, 80, 30),
    ("L01", "Levadura fresca paquete 500 g", 9_000, 30, 10),
    ("E01", "Bolsas de papel paquete x100", 5_000, 120, 20),
]
INSUMOS_CON_IVA = {"M01", "E01"}
EMPLEADOS_PAN = [  # nombre, cédula, cargo, (salario 2025, salario 2026), clase de riesgo ARL
    ("WILSON ANDRES CAICEDO MURILLO", "1116254781", "Panadero", (1_900_000, 2_250_000), 2),
    ("YURANI PATRICIA GOMEZ LOAIZA", "1116260345", "Auxiliar de panadería", (1_423_500, 1_750_905), 2),
    ("DIANA MARCELA VALENCIA RUIZ", "1115078921", "Vendedora", (1_423_500, 1_750_905), 1),
    ("ANDREA CAROLINA PEREA LOAIZA", "1116241190", "Administradora", (2_300_000, 2_700_000), 1),
]
AUX = {2025: 200_000, 2026: 249_095}
PROV_PAN = ("900517262", "MOLINOS Y DISTRIBUCIONES DEL VALLE S.A.S. (demo)")
INMOB = ("805001234", "INMOBILIARIA EL PORVENIR (demo)")


def _empresa_hoja(c: dict, a: int, m: int, capital: int) -> list[tuple]:
    return [("Razón social", c["razon"]), ("Sigla", c.get("sigla", "")),
            ("NIT", f"{c['nit']}-{digito_verificacion(c['nit'])}"), ("Dirección", c["direccion"]),
            ("Municipio", c["municipio"]), ("Periodo desde", date(a, m, 1).isoformat()),
            ("Periodo hasta", ultimo_dia(a, m).isoformat()), ("Representante legal", c["rep"]),
            ("CC representante legal", c["rep_cc"]), ("Responsable de IVA", "SI"), ("Capital suscrito", capital),
            ("Demo", "SI")]


def panaderia(api: Api) -> str:
    c = PANADERIA
    rng = random.Random(101)
    stock = {p[0]: p[4] for p in INSUMOS}
    por_pagar = 0
    previos: dict[str, Decimal] = {}
    nombres: dict[str, str] = {}
    cierre_dic: dict[str, Decimal] = {}
    cliente_id = None
    print(c["razon"])
    for a, m in meses():
        f = 1.8 if m == 12 else 1.15 if m == 11 else 0.9 if m == 1 else rng.uniform(0.96, 1.06)
        fin = ultimo_dia(a, m)
        d = Diario()
        hojas: dict[str, list[list]] = {}
        if (a, m) == INICIO:
            hojas["SALDOS_INICIALES"] = [["110505", "Caja general", 2_000_000, None], ["111005", "Bancos", 25_000_000, None],
                                         ["143505", "Inventario de insumos", 4_000_000, None],
                                         ["152005", "Maquinaria y equipo", 28_800_000, None],
                                         ["152405", "Muebles y enseres", 6_000_000, None],
                                         ["310505", "Capital suscrito y pagado", None, c["capital"]]]
        pagos_del_mes(d, a, m, previos, nombres, cierre_dic)
        if por_pagar:
            d.asiento(date(a, m, 5), f"CE-{yymm(a, m)}-PRV", "Egreso", "Pago a proveedor de insumos (factura del mes anterior)",
                      [("220505", por_pagar, None), ("111005", None, por_pagar)], PROV_PAN)
        # Insumos: consumo del mes y compra para volver a la existencia buscada.
        consumo = {p[0]: max(1, round(p[3] * f * rng.uniform(0.95, 1.05))) for p in INSUMOS}
        compra = {p[0]: max(0, consumo[p[0]] + round(p[4] * (1.5 if m == 11 else 1)) - stock[p[0]]) for p in INSUMOS}
        inv: list[list] = []
        kardex_mes(inv, a, m, INSUMOS, stock, {k: repartir(v, 4) for k, v in consumo.items()}, compra,
                   dias_semana(a, m), "SP", "Salida", False)
        costo = sum(compra[p[0]] * p[2] for p in INSUMOS)
        iva = sum(compra[p[0]] * p[2] * 19 // 100 for p in INSUMOS if p[0] in INSUMOS_CON_IVA)
        d.asiento(date(a, m, 3), f"FC-{yymm(a, m)}", "Compra", "Compra de insumos a crédito",
                  [("143505", costo, None), ("240810", iva, None, "IVA 19 % margarina y empaques"),
                   ("220505", None, costo + iva)], PROV_PAN)
        por_pagar = costo + iva
        # Ventas de la semana: pan (excluido) y pastelería y bebidas (19 %), consignadas el mismo día.
        venta = 30_000_000 * f * (1.06 if a == 2026 else 1)
        for i, dia in enumerate(dias_semana(a, m)):
            pan, pas = r100(venta * 0.7 / 4 * rng.uniform(0.97, 1.03)), r100(venta * 0.3 / 4 * rng.uniform(0.97, 1.03))
            iva_v = pas * 19 // 100
            fecha = date(a, m, dia)
            d.asiento(fecha, f"RC-{yymm(a, m)}-{i + 1}", "Venta", "Ventas de mostrador de la semana",
                      [("110505", pan + pas + iva_v, None), ("412005", None, pan, "Venta de pan (excluido de IVA)"),
                       ("412005", None, pas, "Pastelería y bebidas gravadas 19 %"), ("240805", None, iva_v, "IVA generado 19 %")],
                      ("", "Clientes de mostrador"))
            d.asiento(fecha, f"NC-{yymm(a, m)}-{i + 1}", "Nota", "Consignación de las ventas de la semana",
                      [("111005", pan + pas + iva_v, None), ("110505", None, pan + pas + iva_v)])
        arriendo = 2_800_000 if a == 2025 else 3_000_000
        ret = arriendo * 35 // 1000
        d.asiento(date(a, m, 5), f"CE-{yymm(a, m)}-ARR", "Egreso", f"Arriendo del local, {MESES[m]}",
                  [("512010", arriendo, None), ("236530", None, ret, "Retención 3,5 % arrendamientos"),
                   ("111005", None, arriendo - ret)], INMOB)
        energia, gas, agua = r100(1_150_000 * f), r100(480_000 * f), r100(160_000 * rng.uniform(0.95, 1.05))
        d.asiento(date(a, m, 12), f"CE-{yymm(a, m)}-SP", "Egreso", "Servicios públicos",
                  [("513530", energia, None, "Energía eléctrica"), ("513555", gas, None, "Gas natural para hornos"),
                   ("513525", agua, None, "Acueducto y alcantarillado"), ("111005", None, energia + gas + agua)],
                  ("900200300", "EMPRESAS DE SERVICIOS DEL VALLE (demo)"))
        hojas["NOMINA"] = nomina_diario(d, a, m, EMPLEADOS_PAN, AUX[a])
        hon = 1_000_000 if a == 2025 else 1_100_000
        d.asiento(fin, f"CE-{yymm(a, m)}-HON", "Egreso", f"Honorarios contables de {MESES[m]}",
                  [("511025", hon, None), ("236515", None, hon * 11 // 100, "Retención 11 % honorarios"),
                   ("111005", None, hon - hon * 11 // 100)], ("", "Carlos Cruz — asesoría contable"))
        d.asiento(fin, f"NC-{yymm(a, m)}-BAN", "Nota", "Comisiones y cuota de manejo",
                  [("530505", 18_500, None), ("111005", None, 18_500)], ("", "Banco (demo)"))
        hojas["MOVIMIENTOS"] = d.filas
        hojas["INVENTARIO_MOVS"] = inv
        hojas["ACTIVOS_FIJOS"] = [["Horno rotatorio a gas de 18 bandejas", "152005", date(2024, 12, 10), 28_800_000, 120, 0, "Línea recta"],
                                  ["Vitrinas refrigeradas y estantería", "152405", date(2024, 12, 10), 6_000_000, 60, 0, "Línea recta"]]
        archivo = guardar(c["clave"], f"Espiga Dorada {a}-{m:02d} plantilla.xlsx",
                          plantilla(_empresa_hoja(c, a, m, c["capital"]), hojas))
        if cliente_id is None:
            docs = [guardar(c["clave"], "Estatutos Espiga Dorada.docx", estatutos_docx(c)),
                    guardar(c["clave"], "RUT Espiga Dorada.pdf", rut_pdf(c))]
            cliente_id, res = cargar(api, docs + [archivo], crear=ficha_base(c["nit"], c["razon"], honorarios_mes="1100000"))
        else:
            cliente_id, res = cargar(api, [archivo], cliente_id=cliente_id)
        informe("Panadería", a, m, res, "cerrado")
        previos = neto(res["saldos_siguiente"])
        nombres = {s["codigo"]: s["nombre"] for s in res["saldos_siguiente"]}
        if m == 12:
            cierre_dic = dict(previos)
    return cliente_id


# ═══════════════════════════════════════════════════════════════════════════
# 2 · FERRETERÍA EL TORNILLO S.A.S. — Buga
# ═══════════════════════════════════════════════════════════════════════════
FERRETERIA = {
    "clave": "2-ferreteria", "razon": "FERRETERÍA EL TORNILLO S.A.S.", "sigla": "EL TORNILLO", "nit": "901536208",
    "municipio": "Buga", "direccion": "Carrera 14 No. 5-62 barrio El Carmen", "email": "eltornillo.demo@example.com",
    "telefono": "3166543210", "ciiu": "4752", "ciiu2": "4663",
    "objeto": "el comercio al por menor y al por mayor de artículos de ferretería, pinturas y materiales de construcción",
    "constitucion": date(2017, 8, 22), "dia_letras": "veintidós", "anio_letras": "dos mil diecisiete", "documento": "017",
    "autorizado": 200_000_000, "capital": 160_000_000, "acciones_letras": "ciento sesenta mil",
    "rep": "HERNAN DARIO SALAZAR CORREA", "rep_cc": "14892334", "rep_cc_fmt": "14.892.334",
    "sup": "GLORIA AMPARO CORREA VELEZ", "sup_cc": "29123678", "sup_cc_fmt": "29.123.678",
    "socios": [("HERNAN DARIO SALAZAR CORREA", "14.892.334", 96_000), ("GLORIA AMPARO CORREA VELEZ", "29.123.678", 64_000)],
}
PRODUCTOS_FER = [  # código, descripción, costo, precio de venta, venta base al mes, existencia buscada
    ("C01", "Cemento gris bulto 50 kg", 28_000, 36_900, 870, 312),
    ("V01", "Varilla corrugada 1/2 pulgada x 6 m", 22_000, 29_500, 545, 208),
    ("P01", "Pintura vinilo tipo 1 galón", 55_000, 79_000, 125, 62),
    ("T01", "Tornillo drywall caja x100", 12_000, 18_500, 190, 83),
    ("PV1", "Tubo PVC sanitario 4 pulgadas x 6 m", 48_000, 67_000, 145, 52),
    ("A01", "Alambre negro calibre 18 (kg)", 7_000, 10_500, 625, 247),
]
EMPLEADOS_FER = [
    ("CRISTIAN CAMILO RENGIFO PAZ", "1115089123", "Vendedor de mostrador", (1_600_000, 1_950_000), 1),
    ("JHON FREDY ARANGO MEJIA", "1115092456", "Bodeguero", (1_423_500, 1_750_905), 3),
    ("OSCAR IVAN LOPEZ CARDONA", "94372811", "Conductor de reparto", (1_500_000, 1_850_000), 4),
]
CONSTRUCTORAS = [("901222333", "CONSTRUCTORA GUADALAJARA S.A.S. (demo)"), ("901444555", "OBRAS Y ACABADOS BUGA S.A.S. (demo)"),
                 ("901666777", "INGENIERÍA CIVIL DEL CENTRO S.A.S. (demo)")]
PROVEEDORES_FER = [("900888111", "CEMENTOS Y ACEROS DEL PACÍFICO S.A.S. (demo)"), ("900999222", "DISTRIBUIDORA FERRETERA NACIONAL S.A.S. (demo)")]


def ferreteria(api: Api) -> str:
    c = FERRETERIA
    rng = random.Random(202)
    stock = {p[0]: p[5] for p in PRODUCTOS_FER}
    cartera = {CONSTRUCTORAS[0]: 8_000_000, CONSTRUCTORAS[1]: 6_000_000, CONSTRUCTORAS[2]: 4_000_000}
    proveedores = {PROVEEDORES_FER[0]: 9_000_000, PROVEEDORES_FER[1]: 5_000_000}
    previos: dict[str, Decimal] = {}
    nombres: dict[str, str] = {}
    cierre_dic: dict[str, Decimal] = {}
    cliente_id = None
    print(c["razon"])
    ultimo = (2026, 8)
    for k, (a, m) in enumerate(meses(INICIO, ultimo)):
        g = (1 + 0.012 * k) * (1.12 if m in (3, 4, 10) else 0.85 if m == 12 else 1)
        fin = ultimo_dia(a, m)
        d = Diario()
        hojas: dict[str, list[list]] = {}
        if (a, m) == INICIO:
            hojas["SALDOS_INICIALES"] = [["110505", "Caja general", 3_000_000, None], ["111005", "Bancos", 35_057_000, None],
                                         ["130505", "Clientes nacionales", 18_000_000, None],
                                         ["143505", "Mercancías", 21_943_000, None],
                                         ["154005", "Flota y equipo de transporte", 84_000_000, None],
                                         ["152405", "Muebles y enseres", 12_000_000, None],
                                         ["220505", "Proveedores nacionales", None, 14_000_000],
                                         ["310505", "Capital suscrito y pagado", None, c["capital"]]]
        pagos_del_mes(d, a, m, previos, nombres, cierre_dic)
        # Recaudo de cartera (75 % de lo que se debía) y pago a proveedores (90 %).
        for t, saldo in list(cartera.items()):
            abono = r100(saldo * 0.75)
            if abono:
                d.asiento(date(a, m, 20), f"RC-{yymm(a, m)}-{t[0][-3:]}", "Ingreso", "Abono a factura de crédito",
                          [("111005", abono, None), ("130505", None, abono)], t)
                cartera[t] -= abono
        for t, saldo in list(proveedores.items()):
            pago = r100(saldo * 0.9)
            if pago:
                d.asiento(date(a, m, 25), f"CE-{yymm(a, m)}-{t[0][-3:]}", "Egreso", "Pago a proveedor",
                          [("220505", pago, None), ("111005", None, pago)], t)
                proveedores[t] -= pago
        ventas = {p[0]: max(1, round(p[4] * g * rng.uniform(0.9, 1.1))) for p in PRODUCTOS_FER}
        compras = {p[0]: max(0, ventas[p[0]] + p[5] - stock[p[0]]) for p in PRODUCTOS_FER}
        partes = {cod: repartir(q, 4) for cod, q in ventas.items()}
        inv: list[list] = []
        kardex_mes(inv, a, m, PRODUCTOS_FER, stock, partes, compras, dias_semana(a, m), "FV", "Venta", True)
        # Compras: cemento y varilla a un proveedor, el resto al otro.
        for prov, codigos in ((PROVEEDORES_FER[0], ("C01", "V01")), (PROVEEDORES_FER[1], ("P01", "T01", "PV1", "A01"))):
            costo = sum(compras[p[0]] * p[2] for p in PRODUCTOS_FER if p[0] in codigos)
            if costo:
                iva = costo * 19 // 100
                d.asiento(date(a, m, 3), f"FC-{yymm(a, m)}-{prov[0][-3:]}", "Compra", "Compra de mercancía a crédito",
                          [("143505", costo, None), ("240810", iva, None, "IVA descontable 19 %"), ("220505", None, costo + iva)], prov)
                proveedores[prov] += costo + iva
        # Ventas: dos semanas de contado, dos a crédito a constructoras.
        for i, dia in enumerate(dias_semana(a, m)):
            base = sum(partes[p[0]][i] * p[3] for p in PRODUCTOS_FER)
            iva = base * 19 // 100
            if i % 2 == 0:
                d.asiento(date(a, m, dia), f"FV-{yymm(a, m)}-{i + 1}", "Venta", "Ventas de contado de la semana",
                          [("110505", base + iva, None), ("4135", None, base), ("240805", None, iva)], ("", "Clientes de mostrador"))
                d.asiento(date(a, m, dia), f"NC-{yymm(a, m)}-{i + 1}", "Nota", "Consignación de las ventas de contado",
                          [("111005", base + iva, None), ("110505", None, base + iva)])
            else:
                t = CONSTRUCTORAS[(k + i) % 3]
                d.asiento(date(a, m, dia), f"FV-{yymm(a, m)}-{i + 1}", "Venta", "Venta a crédito a 30 días",
                          [("130505", base + iva, None), ("4135", None, base), ("240805", None, iva)], t)
                cartera[t] += base + iva
        arriendo = 4_200_000 if a == 2025 else 4_450_000
        ret = arriendo * 35 // 1000
        d.asiento(date(a, m, 5), f"CE-{yymm(a, m)}-ARR", "Egreso", f"Arriendo de la bodega, {MESES[m]}",
                  [("512010", arriendo, None), ("236530", None, ret), ("111005", None, arriendo - ret)],
                  ("805009876", "ARRENDAMIENTOS BUGA (demo)"))
        energia, combustible = r100(520_000 * rng.uniform(0.9, 1.1)), r100(950_000 * g * rng.uniform(0.9, 1.1))
        d.asiento(date(a, m, 12), f"CE-{yymm(a, m)}-SP", "Egreso", "Servicios y combustible de la camioneta",
                  [("513530", energia, None, "Energía eléctrica"), ("519535", combustible, None, "Combustible camioneta de reparto"),
                   ("111005", None, energia + combustible)])
        if m in (2, 6, 10):
            d.asiento(date(a, m, 16), f"CE-{yymm(a, m)}-MTO", "Egreso", "Mantenimiento de la camioneta",
                      [("514540", 680_000, None), ("111005", None, 680_000)], ("900123987", "TALLER AUTOMOTRIZ BUGA (demo)"))
        hojas["NOMINA"] = nomina_diario(d, a, m, EMPLEADOS_FER, AUX[a])
        hon = 1_200_000 if a == 2025 else 1_300_000
        d.asiento(fin, f"CE-{yymm(a, m)}-HON", "Egreso", f"Honorarios contables de {MESES[m]}",
                  [("511025", hon, None), ("236515", None, hon * 11 // 100), ("111005", None, hon - hon * 11 // 100)],
                  ("", "Carlos Cruz — asesoría contable"))
        d.asiento(fin, f"NC-{yymm(a, m)}-BAN", "Nota", "Comisiones bancarias",
                  [("530505", 42_000, None), ("111005", None, 42_000)], ("", "Banco (demo)"))
        hojas["MOVIMIENTOS"] = d.filas
        hojas["INVENTARIO_MOVS"] = inv
        hojas["ACTIVOS_FIJOS"] = [["Camioneta de reparto 1,5 toneladas", "154005", date(2024, 12, 20), 84_000_000, 60, 0, "Línea recta"],
                                  ["Estanterías metálicas de bodega", "152405", date(2024, 12, 20), 12_000_000, 120, 0, "Línea recta"]]
        archivo = guardar(c["clave"], f"El Tornillo {a}-{m:02d} plantilla.xlsx",
                          plantilla(_empresa_hoja(c, a, m, c["capital"]), hojas))
        cerrar = (a, m) != ultimo
        if cliente_id is None:
            docs = [guardar(c["clave"], "Estatutos El Tornillo.docx", estatutos_docx(c)),
                    guardar(c["clave"], "RUT El Tornillo.pdf", rut_pdf(c))]
            cliente_id, res = cargar(api, docs + [archivo], crear=ficha_base(c["nit"], c["razon"], honorarios_mes="1300000"))
        else:
            cliente_id, res = cargar(api, [archivo], cliente_id=cliente_id, cerrar=cerrar)
        informe("Ferretería", a, m, res, "cerrado" if cerrar else "calculado, sin cerrar")
        previos = neto(res["saldos_siguiente"])
        nombres = {s["codigo"]: s["nombre"] for s in res["saldos_siguiente"]}
        if m == 12:
            cierre_dic = dict(previos)
    return cliente_id


# ═══════════════════════════════════════════════════════════════════════════
# 3 · CLÍNICA DENTAL SONRISA DEL VALLE S.A.S. — Cali (cuenta T + hoja de trabajo)
# ═══════════════════════════════════════════════════════════════════════════
CLINICA = {"clave": "3-clinica", "razon": "CLÍNICA DENTAL SONRISA DEL VALLE S.A.S.", "nit": "901611459"}
# Nombre en la hoja → cuenta PUC (lo que el contador confirma la primera vez; luego el sistema lo recuerda).
CUENTAS_CLINICA = {
    "CAJA": "110505", "BANCOS": "111005", "CLIENTES": "130505", "EQUIPO MEDICO CIENTIFICO": "153205",
    "DEPRECIACION ACUMULADA": "159230", "PROVEEDORES": "220505", "RETENCION EN LA FUENTE": "2365",
    "CAPITAL SUSCRITO Y PAGADO": "310505", "UTILIDAD DEL EJERCICIO": "3605", "PERDIDA DEL EJERCICIO": "3610",
    "INGRESOS ODONTOLOGIA": "416505", "HONORARIOS": "5110", "ARRENDAMIENTOS": "512010", "SUELDOS": "510506",
    "SERVICIOS": "5135", "DIVERSOS": "5195", "DEPRECIACIONES": "516025", "GASTOS BANCARIOS": "530505",
}
COLUMNAS_T = ["CAJA", "BANCOS", "CLIENTES", "PROVEEDORES", "RETENCION EN LA FUENTE", "INGRESOS ODONTOLOGIA", "HONORARIOS",
              "ARRENDAMIENTOS", "SUELDOS", "SERVICIOS", "DIVERSOS", "DEPRECIACIONES", "DEPRECIACION ACUMULADA",
              "GASTOS BANCARIOS"]
EPS = [("800111222", "ASEGURADORA DENTAL DEL VALLE (demo)"), ("800333444", "PLAN COMPLEMENTARIO SALUD ORAL (demo)")]


def _cuentas_t(a: int, m: int, filas: list[tuple[str, dict[str, tuple[int, int]]]]) -> list[list]:
    """Cuentas en horizontal, cada una con Debe | Haber; una fila por operación, con su detalle a la izquierda."""
    ancho = len(COLUMNAS_T) * 2
    enc, sub = [None], ["DETALLE"]
    for n in COLUMNAS_T:
        enc += [n, None]
        sub += ["DEBE", "HABER"]
    salida = [[f"CLÍNICA DENTAL SONRISA DEL VALLE S.A.S. — CUENTAS T — {MESES[m].upper()} {a}"], enc, sub]
    for detalle, f in filas:
        fila = [detalle] + [None] * ancho
        for cuenta, (deb, cre) in f.items():
            i = 1 + COLUMNAS_T.index(cuenta) * 2
            fila[i], fila[i + 1] = deb or None, cre or None
        salida.append(fila)
    salida.append(["SUMAS"] + [sum((r[i] or 0) for r in salida[3:]) or None for i in range(1, ancho + 1)])
    return salida


def _hoja_trabajo(a: int, m: int, iniciales: dict[str, Decimal], movs: dict[str, list[int]]) -> list[list]:
    salida = [[f"HOJA DE TRABAJO — {MESES[m].upper()} DE {a}", None, None, None, None],
              ["CUENTAS", "BALANCE INICIAL", None, "MOVIMIENTO", None],
              [None, "DEBE", "HABER", "DEBE", "HABER"]]
    totales = [CERO] * 4
    for nombre in dict.fromkeys(list(iniciales) + list(movs)):
        v = iniciales.get(nombre, CERO)
        dm, cm = movs.get(nombre, (0, 0))
        fila = [nombre, v if v > 0 else None, -v if v < 0 else None, dm or None, cm or None]
        if any(x for x in fila[1:]):
            salida.append(fila)
            for i in range(4):
                totales[i] += Decimal(fila[i + 1] or 0)
    salida.append(["TOTALES", *totales])
    return salida


def clinica(api: Api) -> str:
    c = CLINICA
    rng = random.Random(303)
    inverso = {v: k for k, v in CUENTAS_CLINICA.items()}
    iniciales = {"CAJA": Decimal(1_500_000), "BANCOS": Decimal(22_000_000), "EQUIPO MEDICO CIENTIFICO": Decimal(66_000_000),
                 "CAPITAL SUSCRITO Y PAGADO": Decimal(-89_500_000)}
    por_cobrar = por_pagar = retencion = 0
    cliente_id = None
    print(c["razon"])
    for a, m in meses():
        filas: list[tuple[str, dict[str, tuple[int, int]]]] = []
        if por_cobrar:
            filas.append(("Pago de la aseguradora (mes anterior)", {"BANCOS": (por_cobrar, 0), "CLIENTES": (0, por_cobrar)}))
        if por_pagar:
            filas.append(("Pago al depósito dental", {"PROVEEDORES": (por_pagar, 0), "BANCOS": (0, por_pagar)}))
        if retencion:
            filas.append(("Pago retención mes anterior", {"RETENCION EN LA FUENTE": (retencion, 0), "BANCOS": (0, retencion)}))
        factor = (1.05 if a == 2026 else 1) * (0.85 if m in (1, 12) else rng.uniform(0.95, 1.08))
        particular = 0
        for semana in range(3):
            v = r100(5_600_000 * factor * rng.uniform(0.9, 1.1))
            particular += v
            filas.append((f"Consultas particulares, recibos semana {semana + 1}", {"CAJA": (v, 0), "INGRESOS ODONTOLOGIA": (0, v)}))
        filas.append(("Consignación del efectivo", {"BANCOS": (particular, 0), "CAJA": (0, particular)}))
        eps = r100(11_000_000 * factor * rng.uniform(0.9, 1.1))
        filas.append((f"Cuenta de cobro a {EPS[m % 2][1]}", {"CLIENTES": (eps, 0), "INGRESOS ODONTOLOGIA": (0, eps)}))
        por_cobrar = eps
        hon = r100(8_000_000 * factor * rng.uniform(0.92, 1.08))
        ret_h = hon * 11 // 100
        arr = 3_500_000 if a == 2025 else 3_700_000
        ret_a = arr * 35 // 1000
        filas.append(("Honorarios ortodoncista y endodoncista (ret. 11 %)",
                      {"HONORARIOS": (hon, 0), "RETENCION EN LA FUENTE": (0, ret_h), "BANCOS": (0, hon - ret_h)}))
        filas.append(("Arriendo del consultorio (ret. 3,5 %)",
                      {"ARRENDAMIENTOS": (arr, 0), "RETENCION EN LA FUENTE": (0, ret_a), "BANCOS": (0, arr - ret_a)}))
        retencion = ret_h + ret_a
        sueldos = 3_500_000 if a == 2025 else 4_100_000
        filas.append(("Sueldos auxiliar y recepcionista", {"SUELDOS": (sueldos, 0), "BANCOS": (0, sueldos)}))
        mat = r100(4_000_000 * factor * rng.uniform(0.9, 1.1))
        filas.append(("Materiales odontológicos a crédito", {"DIVERSOS": (mat, 0), "PROVEEDORES": (0, mat)}))
        por_pagar = mat
        serv = r100(900_000 * rng.uniform(0.9, 1.1))
        filas.append(("Energía, agua e internet", {"SERVICIOS": (serv, 0), "BANCOS": (0, serv)}))
        filas.append(("Depreciación de los equipos", {"DEPRECIACIONES": (1_100_000, 0), "DEPRECIACION ACUMULADA": (0, 1_100_000)}))
        filas.append(("Comisiones del banco", {"GASTOS BANCARIOS": (31_000, 0), "BANCOS": (0, 31_000)}))
        ultimo = (a, m) == FIN
        if ultimo:
            # Error de digitación a propósito: el pago del honorario quedó $50.000 por debajo.
            h = next(f for _, f in filas if "HONORARIOS" in f)
            h["BANCOS"] = (0, h["BANCOS"][1] - 50_000)
        movs: dict[str, list[int]] = {}
        for _, f in filas:
            for cuenta, (deb, cre) in f.items():
                acu = movs.setdefault(cuenta, [0, 0])
                acu[0] += deb
                acu[1] += cre
        datos = libro({"CUENTAS T": _cuentas_t(a, m, filas), "HOJA DE TRABAJO": _hoja_trabajo(a, m, iniciales, movs)})
        archivo = guardar(c["clave"], f"Sonrisa del Valle {MESES[m]} {a}.xlsx", datos)
        periodo = (date(a, m, 1), ultimo_dia(a, m))
        if cliente_id is None:
            crear = ficha_base(c["nit"], c["razon"], tipo_persona="juridica", tipo_sociedad="S.A.S.", municipio="Cali",
                               departamento="Valle del Cauca", direccion="Avenida 6 Norte No. 23-41 consultorio 302",
                               ciiu="8621", responsable_iva=False, regimen="no_responsable_iva",
                               rep_legal="Paola Andrea Quintero Lasso", rep_legal_cc="66987123",
                               capital_suscrito="89500000", honorarios_mes="950000",
                               email="sonrisadelvalle.demo@example.com", telefono="3178889900")
            cliente_id, res = cargar(api, [archivo], crear=crear, alias=CUENTAS_CLINICA, periodo=periodo)
        else:
            cliente_id, res = cargar(api, [archivo], cliente_id=cliente_id, alias=CUENTAS_CLINICA, periodo=periodo,
                                     cerrar=not ultimo)
        informe("Clínica", a, m, res, "calculado, sin cerrar (descuadre de $50.000)" if ultimo else "cerrado")
        iniciales = {inverso.get(s["codigo"], normalizar(s["nombre"])): Decimal(str(s["debito"] or 0)) - Decimal(str(s["credito"] or 0))
                     for s in res["saldos_siguiente"]}
    return cliente_id


# ═══════════════════════════════════════════════════════════════════════════
# 4 · TRANSPORTES RÍO CAUCA S.A.S. — Palmira (un libro diario en CSV)
# ═══════════════════════════════════════════════════════════════════════════
TRANSPORTES = {"clave": "4-transportes", "razon": "TRANSPORTES RÍO CAUCA S.A.S.", "nit": "901394772"}
CLIENTES_TR = [("900300400", "INGENIO LA PRIMAVERA DEL VALLE (demo)"), ("900765432", "COMERCIALIZADORA AGRÍCOLA PALMIRA (demo)"),
               ("901112131", "LOGÍSTICA DEL PACÍFICO S.A.S. (demo)")]
BANCO_TR = ("900100200", "BANCO DEMO DEL VALLE")


def transportes(api: Api) -> str:
    c = TRANSPORTES
    rng = random.Random(404)
    filas: list[list] = []
    banco = 0
    prestamo = 120_000_000
    por_cobrar: list[tuple[tuple[str, str], int]] = []
    ret_pend = 0

    def asiento(fecha: date, comp: str, desc: str, lineas: list[tuple], tercero=("", "")):
        nonlocal banco
        assert sum(l[1] for l in lineas) == sum(l[2] for l in lineas), comp
        for cod, deb, cre in lineas:
            filas.append([fecha.isoformat(), comp, cod, nombre_puc(cod), tercero[1], desc,
                          f"{deb:,}".replace(",", ".") if deb else "0", f"{cre:,}".replace(",", ".") if cre else "0"])
            if cod == "111005":
                banco += deb - cre

    asiento(date(2025, 1, 2), "CA-0001", "Asiento de apertura",
            [("111005", 30_000_000, 0), ("154005", 270_000_000, 0), ("210510", 0, 120_000_000), ("310505", 0, 180_000_000)])
    n = 0
    for a, m in meses():
        fin = ultimo_dia(a, m)
        k = yymm(a, m)
        if banco < 12_000_000:
            asiento(date(a, m, 1), f"RC-{k}-SOC", "Préstamo del socio para capital de trabajo",
                    [("111005", 15_000_000, 0), ("235505", 0, 15_000_000)], ("1113456789", "RICARDO ANTONIO VELASCO PRADO"))
        for t, v in por_cobrar:
            ret = v // 100
            asiento(date(a, m, 10), f"RC-{k}-{t[0][-3:]}", "Pago de factura de flete (retención 1 %)",
                    [("111005", v - ret, 0), ("135515", ret, 0), ("130505", 0, v)], t)
        por_cobrar = []
        if ret_pend:
            asiento(date(a, m, 15), f"CE-{k}-RTE", "Pago retención en la fuente del mes anterior",
                    [("236515", ret_pend, 0), ("111005", 0, ret_pend)])
        # Contratos grandes en 2025 dan meses con utilidad; en 2026 el negocio se cae.
        ventas = ((17_000_000 if a == 2025 else 11_000_000) * rng.uniform(0.82, 1.18)
                  + (7_000_000 if (a, m) in ((2025, 4), (2025, 9), (2025, 10)) else 0))
        for i, t in enumerate(CLIENTES_TR):
            v = r100(ventas * (0.5, 0.3, 0.2)[i])
            n += 1
            asiento(date(a, m, 8 + i * 7), f"FV-{n:04d}", "Flete de carga terrestre", [("130505", v, 0), ("414505", 0, v)], t)
            por_cobrar.append((t, v))
        comb = r100(ventas * 0.22)
        asiento(date(a, m, 6), f"CE-{k}-COM", "Combustible ACPM de los camiones", [("519535", comb, 0), ("111005", 0, comb)],
                ("900555666", "ESTACIÓN DE SERVICIO LA RECTA (demo)"))
        peajes = r100(1_200_000 * rng.uniform(0.85, 1.15))
        asiento(date(a, m, 6), f"CE-{k}-PEA", "Peajes del mes", [("519595", peajes, 0), ("111005", 0, peajes)])
        mant = r100(1_500_000 * rng.uniform(0.8, 1.2)) + (6_500_000 if (a, m) in ((2025, 5), (2025, 11), (2026, 3), (2026, 7)) else 0)
        asiento(date(a, m, 18), f"CE-{k}-MTO", "Mantenimiento y llantas de la flota", [("514540", mant, 0), ("111005", 0, mant)],
                ("900777888", "TECNIDIÉSEL PALMIRA (demo)"))
        sueldos = 2 * (2_200_000 if a == 2025 else 2_500_000)
        aux = 2 * AUX[a]
        aportes = r100(sueldos * 0.2296)
        asiento(fin, f"CE-{k}-NOM", "Nómina de conductores y aportes",
                [("510506", sueldos, 0), ("510527", aux, 0), ("510570", r100(sueldos * 0.12), 0),
                 ("510568", aportes - r100(sueldos * 0.12) - r100(sueldos * 0.04), 0), ("510572", r100(sueldos * 0.04), 0),
                 ("111005", 0, sueldos + aux + aportes)])
        asiento(date(a, m, 2), f"CE-{k}-SEG", "Póliza de seguros de los vehículos", [("513040", 1_400_000, 0), ("111005", 0, 1_400_000)],
                ("900500600", "ASEGURADORA DEMO S.A."))
        interes = r100(prestamo * 0.012)
        asiento(date(a, m, 26), f"CE-{k}-PRE", "Cuota del préstamo: intereses y abono a capital",
                [("530520", interes, 0), ("210510", 2_000_000, 0), ("111005", 0, interes + 2_000_000)], BANCO_TR)
        prestamo -= 2_000_000
        asiento(fin, f"NC-{k}-DEP", "Depreciación mensual de la flota", [("516035", 2_250_000, 0), ("159235", 0, 2_250_000)])
        asiento(fin, f"CE-{k}-HON", "Honorarios contables", [("511025", 900_000, 0), ("236515", 0, 99_000), ("111005", 0, 801_000)],
                ("", "Carlos Cruz — asesoría contable"))
        ret_pend = 99_000
        asiento(fin, f"NC-{k}-BAN", "Gastos bancarios", [("530505", 35_000, 0), ("111005", 0, 35_000)], BANCO_TR)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Fecha", "Comprobante", "Cuenta", "Nombre de la cuenta", "Tercero", "Detalle", "Débito", "Crédito"])
    w.writerows(filas)
    archivo = guardar(c["clave"], "Libro diario Transportes Rio Cauca 2025-2026.csv", buf.getvalue().encode("utf-8-sig"))
    print(c["razon"])
    crear = ficha_base(c["nit"], c["razon"], tipo_persona="juridica", tipo_sociedad="S.A.S.", municipio="Palmira",
                       departamento="Valle del Cauca", direccion="Kilómetro 3 vía Palmira-Cali, bodega 7", ciiu="4923",
                       responsable_iva=False, regimen="no_responsable_iva", rep_legal="Ricardo Antonio Velasco Prado",
                       rep_legal_cc="1113456789", capital_suscrito="180000000", honorarios_mes="900000",
                       email="riocauca.demo@example.com", telefono="3209991122")
    cliente_id, res = cargar(api, [archivo], crear=crear, por_periodo=True, cerrar=True)
    for p in res.get("periodos_procesados") or []:
        print(f"  Transportes {p['desde'][:7]}  utilidad {Decimal(str(p.get('utilidad') or 0)):>16,.2f}  "
              f"{'cuadra' if p.get('cuadra') else 'NO CUADRA':<9}  {p['estado']}")
    return cliente_id


# ═══════════════════════════════════════════════════════════════════════════
# 5 · MARÍA ELENA ROJAS — Guacarí (registros auxiliares desordenados, bimestral)
# ═══════════════════════════════════════════════════════════════════════════
MARIA = {"clave": "5-maria-elena", "razon": "MARÍA ELENA ROJAS", "nit": "29874553"}
COMPRADORES = ["CENTRAL DE ABASTOS - BODEGA 14", "COMERCIALIZADORA EL SOL", "DON AURELIO (VECINO)", "SUPERMERCADO EL VECINO",
               "FRUVER LA COSECHA"]
PRODUCTOS_FINCA = [("PLÁTANO HARTÓN", 2_400_000), ("LIMÓN TAHITÍ", 1_800_000), ("YUCA", 900_000), ("MARACUYÁ", 1_200_000)]
GASTOS_FINCA = [("JORNALES", 1_300_000), ("ABONO TRIPLE 15", 650_000), ("FUMIGACIÓN", 420_000), ("TRANSPORTE A LA CENTRAL", 380_000),
                ("ENERGÍA BOMBA DE RIEGO", 160_000), ("ESTACONES Y AMARRES", 140_000)]
ENCABEZADOS_VENTAS = [["FECHA", "COMPRADOR", "PRODUCTO", "VALOR"], ["Fecha", "Producto", "Cliente", "Total"],
                      ["VALOR", "FECHA", "DETALLE", "A QUIEN"], ["Fecha", "Comprador", "Qué se vendió", "Valor venta"]]
ENCABEZADOS_GASTOS = [["FECHA", "CONCEPTO", "VALOR"], ["Concepto", "Valor", "Fecha"], ["fecha", "gasto", "pagado"]]


def _fila(enc: list[str], valores: dict[str, object]) -> list:
    orden = {"FECHA": "fecha", "COMPRADOR": "tercero", "CLIENTE": "tercero", "A QUIEN": "tercero", "PRODUCTO": "producto",
             "DETALLE": "producto", "QUE SE VENDIO": "producto", "VALOR": "valor", "TOTAL": "valor", "VALOR VENTA": "valor",
             "CONCEPTO": "concepto", "GASTO": "concepto", "PAGADO": "valor"}
    return [valores.get(orden[normalizar(e)]) for e in enc]


def maria_elena(api: Api) -> str:
    c = MARIA
    rng = random.Random(505)
    ventas: list[list] = []
    gastos = {2025: [], 2026: []}
    hasta = (2026, 6)
    for i, (a, m) in enumerate(meses(INICIO, hasta)):
        fin = ultimo_dia(a, m)
        enc = ENCABEZADOS_VENTAS[i % len(ENCABEZADOS_VENTAS)]
        titulo = (f"VENTAS {MESES[m].upper()} {a}" if i % 3 == 0 else f"Ventas de {MESES[m]}" if i % 3 == 1
                  else f"VENTAS {MESES[m].upper()}/{a}")
        ventas += [[titulo], enc]
        total = 0
        dias = sorted(rng.sample(range(2, fin.day), 5))
        for j, dia in enumerate(dias):
            prod, base = PRODUCTOS_FINCA[(i + j) % len(PRODUCTOS_FINCA)]
            v = r100(base * rng.uniform(0.7, 1.3) * (1.1 if a == 2026 else 1))
            total += v
            fecha = date(a, m, dia)
            if i == 7 and j == 0:
                fecha = date(a, m - 1, dia)          # fecha copiada del mes anterior
            valor = f"$ {v:,}".replace(",", ".") if (i + j) % 6 == 0 else v
            ventas.append(_fila(enc, {"fecha": fecha, "tercero": COMPRADORES[(i * 2 + j) % len(COMPRADORES)],
                                      "producto": prod, "valor": valor}))
        if i % 2 == 0:
            ventas.append(_fila(enc, {"fecha": None, "tercero": "TOTAL", "producto": None, "valor": total})
                          if "A QUIEN" not in enc else ["TOTAL", None, None, None])
        ventas.append([])
        if i % 5 == 3:
            ventas.append([])
        enc_g = ENCABEZADOS_GASTOS[i % len(ENCABEZADOS_GASTOS)]
        hoja = gastos[a]
        hoja += [[f"GASTOS {MESES[m].upper()} {a}" if i % 2 else f"gastos {MESES[m]}"], enc_g]
        for j, (concepto, base) in enumerate(GASTOS_FINCA):
            if concepto == "ESTACONES Y AMARRES" and m % 3:
                continue
            dia = min(fin.day, 3 + j * 4 + rng.randint(0, 2))
            v = r100(base * rng.uniform(0.8, 1.2))
            hoja.append(_fila(enc_g, {"fecha": date(a, m, dia), "concepto": concepto, "valor": v}))
        if i % 4 == 1:
            hoja.append(_fila(enc_g, {"fecha": date(a, m, 28), "concepto": "JORNALES", "valor": 0}))
        hoja.append([])
    datos = libro({"VENTAS": ventas, "GASTOS 2025": gastos[2025], "GASTOS 2026": gastos[2026]})
    archivo = guardar(c["clave"], "MARIA ELENA ROJAS - cuentas finca La Esperanza 2025-2026.xlsx", datos)
    print(c["razon"])
    crear = ficha_base(c["nit"], c["razon"], tipo_persona="natural", municipio="Guacarí", departamento="Valle del Cauca",
                       direccion="Finca La Esperanza, vereda Sonso", ciiu="0122", responsable_iva=False,
                       regimen="no_responsable_iva", periodicidad="bimestral", honorarios_mes="350000",
                       telefono="3112223344")
    cliente_id, res = cargar(api, [archivo], crear=crear, por_periodo=True, cerrar=True)
    for p in res.get("periodos_procesados") or []:
        print(f"  María Elena {p['desde'][:7]}  utilidad {Decimal(str(p.get('utilidad') or 0)):>16,.2f}  "
              f"{'cuadra' if p.get('cuadra') else 'NO CUADRA':<9}  {p['estado']}")
    return cliente_id


CLIENTES = {"1": panaderia, "2": ferreteria, "3": clinica, "4": transportes, "5": maria_elena}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Genera y carga los cinco clientes de demostración por la API.")
    ap.add_argument("--url", required=True, help="Dirección de la aplicación, p. ej. http://127.0.0.1:8001")
    ap.add_argument("--usuario", default="", help="Usuario del contador, si la aplicación pide inicio de sesión")
    ap.add_argument("--clave", default="", help="Contraseña del contador")
    ap.add_argument("--solo", default="1,2,3,4,5", help="Clientes a cargar, p. ej. 1,3")
    args = ap.parse_args(argv)
    api = Api(args.url, args.usuario, args.clave)
    salud = api.get("/api/salud")
    print(f"Aplicación {salud.get('version')}")
    claves = [x.strip() for x in args.solo.split(",") if x.strip()]
    nits = {"1": PANADERIA, "2": FERRETERIA, "3": CLINICA, "4": TRANSPORTES, "5": MARIA}
    cargados = {x["nit"]: x["razon_social"] for x in api.get("/api/clientes/demostracion")["clientes"]}
    repetidos = [cargados[nits[k]["nit"]] for k in claves if nits[k]["nit"] in cargados]
    if repetidos:
        print("Estos clientes de demostración ya están cargados; bórrelos primero desde el menú de la cuenta › Sistema:")
        for r in repetidos:
            print(f"  · {r}")
        return 1
    for clave in claves:
        CLIENTES[clave](api)
    return 0


if __name__ == "__main__":
    sys.exit(main())
