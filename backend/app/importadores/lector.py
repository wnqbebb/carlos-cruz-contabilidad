"""Lectura uniforme de .xlsx, .xls y .csv en rejillas de celdas (índices base 0)."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from ..utils.numeros import normalizar


@dataclass
class Hoja:
    archivo: str
    nombre: str
    valores: list[list] = field(default_factory=list)
    formulas: dict[tuple[int, int], str] = field(default_factory=dict)

    @property
    def nfilas(self) -> int:
        return len(self.valores)

    @property
    def ncols(self) -> int:
        return max((len(f) for f in self.valores), default=0)

    def v(self, r: int, c: int):
        if 0 <= r < len(self.valores) and 0 <= c < len(self.valores[r]):
            val = self.valores[r][c]
            if isinstance(val, str) and not val.strip():
                return None
            return val
        return None

    def texto(self, r: int, c: int) -> str:
        val = self.v(r, c)
        return "" if val is None else str(val).strip()

    def norm(self, r: int, c: int) -> str:
        return normalizar(self.v(r, c))

    def formula(self, r: int, c: int) -> str | None:
        return self.formulas.get((r, c))

    def fila_textos(self, r: int) -> list[tuple[int, str]]:
        return [(c, normalizar(v)) for c, v in enumerate(self.valores[r]) if isinstance(v, str) and v.strip()] \
            if 0 <= r < self.nfilas else []

    def buscar_texto(self, patron: str, max_filas: int | None = None) -> tuple[int, int] | None:
        p = normalizar(patron)
        for r in range(min(self.nfilas, max_filas or self.nfilas)):
            for c, t in self.fila_textos(r):
                if p in t:
                    return r, c
        return None

    def origen(self, r: int, c: int | None = None) -> str:
        celda = f"{get_column_letter(c + 1)}{r + 1}" if c is not None else f"fila {r + 1}"
        return f"{self.archivo} › {self.nombre} › {celda}"


def _recortar(filas: list[list]) -> list[list]:
    while filas and all(v is None or (isinstance(v, str) and not v.strip()) for v in filas[-1]):
        filas.pop()
    return filas


def leer_xlsx(contenido: bytes, archivo: str) -> list[Hoja]:
    from openpyxl import load_workbook

    wb_v = load_workbook(io.BytesIO(contenido), data_only=True, read_only=False)
    wb_f = load_workbook(io.BytesIO(contenido), data_only=False, read_only=False)
    hojas = []
    for ws in wb_v.worksheets:
        valores = [list(fila) for fila in ws.iter_rows(values_only=True)]
        formulas = {}
        wsf = wb_f[ws.title]
        for fila in wsf.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    formulas[(celda.row - 1, celda.column - 1)] = celda.value
        hojas.append(Hoja(archivo, ws.title, _recortar(valores), formulas))
    return hojas


def leer_xls(contenido: bytes, archivo: str) -> list[Hoja]:
    import xlrd

    wb = xlrd.open_workbook(file_contents=contenido)
    hojas = []
    for sh in wb.sheets():
        filas = []
        for r in range(sh.nrows):
            fila = []
            for c in range(sh.ncols):
                celda = sh.cell(r, c)
                if celda.ctype == xlrd.XL_CELL_DATE:
                    fila.append(xlrd.xldate_as_datetime(celda.value, wb.datemode))
                elif celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    fila.append(None)
                elif celda.ctype == xlrd.XL_CELL_ERROR:
                    fila.append(None)
                else:
                    fila.append(celda.value)
            filas.append(fila)
        hojas.append(Hoja(archivo, sh.name, _recortar(filas), {}))
    return hojas


def leer_csv(contenido: bytes, archivo: str) -> list[Hoja]:
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t|")
        delim = dialecto.delimiter
    except csv.Error:
        delim = ";" if texto.count(";") > texto.count(",") else ","
    filas = [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delim)]
    nombre = archivo.rsplit(".", 1)[0]
    return [Hoja(archivo, nombre, _recortar(filas), {})]


def leer_archivo(contenido: bytes, archivo: str) -> list[Hoja]:
    ext = archivo.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm"):
        return leer_xlsx(contenido, archivo)
    if ext == "xls":
        return leer_xls(contenido, archivo)
    if ext in ("csv", "txt"):
        return leer_csv(contenido, archivo)
    # Se importan aquí para no cargar las librerías de PDF y Word si no hacen falta.
    if ext == "pdf":
        from .documentos import leer_pdf_completo

        return leer_pdf_completo(contenido, archivo)
    if ext == "docx":
        from .documentos import leer_docx

        return leer_docx(contenido, archivo)
    if ext == "doc":
        raise ValueError(
            f"«{archivo}» está en el formato viejo de Word (.doc). Ábralo en Word y guárdelo "
            "como .docx, o expórtelo a PDF."
        )
    raise ValueError(
        f"Formato no soportado: .{ext}. Se aceptan .xlsx, .xlsm, .xls, .csv, .txt, .pdf y .docx"
    )
