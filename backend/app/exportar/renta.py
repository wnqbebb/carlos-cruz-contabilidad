"""Salidas de la declaración de renta (spec v2.3 · 5.7).

- Borrador del 210 en PDF, con la numeración de casillas del oficial y la marca
  «BORRADOR — no válido para presentar».
- Papel de trabajo en Excel: cada casilla con su fórmula y sus líneas de origen.
- Resumen para el cliente, una página, en lenguaje sencillo.

Reciben la vista ya serializada (`a_json`): los importes llegan como texto decimal.
"""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from ..seguridad.archivos import blindar_libro
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TINTA = colors.HexColor("#141414")
LINEA = colors.HexColor("#DEDBD3")
HOJA2 = colors.HexColor("#F7F5F0")
AMBAR = colors.HexColor("#8A5A00")
_e = getSampleStyleSheet()
E_TIT = ParagraphStyle("t", parent=_e["Title"], fontSize=13, leading=16, textColor=TINTA, spaceAfter=2)
E_SUB = ParagraphStyle("s", parent=_e["Normal"], fontSize=9, leading=12, textColor=TINTA)
E_CEL = ParagraphStyle("c", parent=_e["Normal"], fontSize=7.5, leading=9)
E_H = ParagraphStyle("h", parent=_e["Heading3"], fontSize=10, textColor=TINTA, spaceBefore=8, spaceAfter=3)
E_GRANDE = ParagraphStyle("g", parent=_e["Normal"], fontSize=15, leading=19, textColor=TINTA, spaceAfter=4)


def _d(v) -> Decimal:
    try:
        return Decimal(str(v or "0"))
    except Exception:
        return Decimal("0")


def pesos(v) -> str:
    n = _d(v)
    texto = f"{abs(n):,.0f}".replace(",", ".")
    return f"-$ {texto}" if n < 0 else f"$ {texto}"


def _seguro_excel(texto) -> str:
    """Evita la inyección de fórmulas: un texto que empieza por = + - @ no se ejecuta."""
    t = str(texto or "")
    return "'" + t if t[:1] in ("=", "+", "-", "@", "\t", "\r") else t


def _marca_borrador(lienzo, doc) -> None:
    lienzo.saveState()
    lienzo.setFillColor(colors.Color(0.55, 0.1, 0.1, alpha=0.12))
    lienzo.setFont("Helvetica-Bold", 46)
    lienzo.translate(letter[0] / 2, letter[1] / 2)
    lienzo.rotate(35)
    lienzo.drawCentredString(0, 0, "BORRADOR")
    lienzo.restoreState()
    lienzo.setFont("Helvetica", 7.5)
    lienzo.setFillColor(AMBAR)
    lienzo.drawString(1.5 * cm, 1 * cm, "BORRADOR — no válido para presentar. Diligencie la declaración en el portal "
                                        "de la DIAN con la firma electrónica del contribuyente.")
    lienzo.drawRightString(letter[0] - 1.5 * cm, 1 * cm, f"Página {doc.page}")


def _documento(buf) -> SimpleDocTemplate:
    return SimpleDocTemplate(buf, pagesize=letter, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.4 * cm,
                             bottomMargin=1.6 * cm, title="Borrador formulario 210", author="Carlos Cruz",
                             creator="Carlos Cruz", producer="Carlos Cruz", subject="", keywords="")


def borrador_pdf(v: dict) -> bytes:
    r = v["resultado"]
    buf = io.BytesIO()
    doc = _documento(buf)
    c = v["contribuyente"]
    hist = [Paragraph(f"Formulario 210 · Año gravable {v['anio']} · BORRADOR", E_TIT),
            Paragraph(f"{c['nombre']} · documento {c['nit']}", E_SUB)]
    oblig = r.get("obligacion") or {}
    if oblig:
        hist.append(Paragraph(oblig.get("veredicto", ""), E_SUB))
    hist.append(Spacer(1, 6))
    seccion = None
    filas = [["Casilla", "Concepto", "Propuesta DIAN", "Declaración"]]
    estilo = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5), ("BACKGROUND", (0, 0), (-1, 0), HOJA2),
              ("LINEBELOW", (0, 0), (-1, -1), 0.25, LINEA), ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
              ("FONTSIZE", (0, 1), (-1, -1), 7.5), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for cas in r["casillas"]:
        if cas["seccion"] != seccion:
            seccion = cas["seccion"]
            filas.append(["", Paragraph(f"<b>{seccion}</b>", E_CEL), "", ""])
            estilo.append(("BACKGROUND", (0, len(filas) - 1), (-1, len(filas) - 1), HOJA2))
        nombre = cas["nombre"] + (f" — {cas['columna']}" if cas.get("columna") else "")
        filas.append([str(cas["casilla"]), Paragraph(nombre, E_CEL), pesos(cas["dian"]), pesos(cas["optimizada"])])
    t = Table(filas, colWidths=[1.4 * cm, 10.6 * cm, 3 * cm, 3 * cm], repeatRows=1)
    t.setStyle(TableStyle(estilo))
    hist.append(t)
    doc.build(hist, onFirstPage=_marca_borrador, onLaterPages=_marca_borrador)
    return buf.getvalue()


def resumen_cliente(v: dict) -> bytes:
    r = v["resultado"]
    buf = io.BytesIO()
    doc = _documento(buf)
    c = v["contribuyente"]
    cif = r["cifras"]
    neto = _d(cif["neto"])
    oblig = r.get("obligacion") or {}
    venc = r.get("vencimiento") or {}
    hist = [Paragraph(f"Su declaración de renta {v['anio']}", E_TIT), Paragraph(c["nombre"], E_SUB), Spacer(1, 10),
            Paragraph("<b>¿Debe declarar?</b>", E_H), Paragraph(oblig.get("veredicto", "Falta información."), E_SUB),
            Paragraph("<b>¿Paga o le devuelven?</b>", E_H),
            Paragraph(f"Paga {pesos(neto)}" if neto > 0 else
                      (f"Le devuelven {pesos(-neto)}" if neto < 0 else "No paga ni le devuelven."), E_GRANDE),
            Paragraph("<b>Ahorro frente a la propuesta de la DIAN</b>", E_H),
            Paragraph(pesos(cif.get("ahorro")) if _d(cif.get("ahorro")) > 0 else
                      "La propuesta de la DIAN no tenía beneficios que aplicar.", E_SUB),
            Paragraph("<b>Fecha límite</b>", E_H), Paragraph(venc.get("texto", ""), E_SUB)]
    faltan = r.get("documentos_faltantes") or []
    if faltan:
        hist += [Paragraph("<b>Documentos que debe entregar</b>", E_H)] + [Paragraph(f"• {f}", E_SUB) for f in faltan]
    beneficios = [b for b in r.get("beneficios") or [] if not b.get("valor") and _d(b.get("ahorro_hasta")) > 0]
    if beneficios:
        hist.append(Paragraph("<b>Documentos que bajarían el impuesto</b>", E_H))
        hist += [Paragraph(f"• {b['soporte'].capitalize()}: ahorro posible hasta {pesos(b['ahorro_hasta'])}", E_SUB)
                 for b in beneficios]
    hist += [Spacer(1, 14), Paragraph("Este resumen se basa en la información disponible a la fecha. La declaración la "
                                      "presenta su contador en el portal de la DIAN con su firma electrónica.", E_SUB)]
    doc.build(hist, onFirstPage=_marca_borrador, onLaterPages=_marca_borrador)
    return buf.getvalue()


def papel_trabajo(v: dict) -> bytes:
    r = v["resultado"]
    wb = Workbook()
    wb.properties.creator = "Carlos Cruz"
    ws = wb.active
    ws.title = "Formulario 210"
    negrita = Font(bold=True)
    fondo = PatternFill("solid", fgColor="F7F5F0")
    ws.append([f"Papel de trabajo · formulario 210 · año gravable {v['anio']} · BORRADOR"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([_seguro_excel(f"{v['contribuyente']['nombre']} · documento {v['contribuyente']['nit']}")])
    ws.append([])
    ws.append(["Casilla", "Sección", "Concepto", "Fórmula", "Propuesta DIAN", "Declaración", "Explicación"])
    for celda in ws[4]:
        celda.font = negrita
        celda.fill = fondo
    for cas in r["casillas"]:
        ws.append([cas["casilla"], cas["seccion"], _seguro_excel(cas["nombre"] + (f" — {cas['columna']}" if cas.get("columna") else "")),
                   _seguro_excel(cas.get("formula", "")), int(_d(cas["dian"])), int(_d(cas["optimizada"])),
                   _seguro_excel(cas.get("explicacion", ""))])
    for fila in ws.iter_rows(min_row=5, min_col=5, max_col=6):
        for celda in fila:
            celda.number_format = '#,##0'
    for col, ancho in zip("ABCDEFG", (8, 22, 60, 30, 16, 16, 60)):
        ws.column_dimensions[col].width = ancho

    lo = wb.create_sheet("Líneas del reporte")
    lo.append(["Documento", "Página", "Fila", "Entidad", "Titular", "Detalle", "Valor", "Uso sugerido",
               "Clasificación", "Incluida", "Leída con confianza", "Marcada en el papel"])
    for celda in lo[1]:
        celda.font = negrita
        celda.fill = fondo
    for l in r["lineas"]:
        lo.append([_seguro_excel(l["documento"]), l["pagina"], l["fila"], _seguro_excel(l["entidad"]),
                   _seguro_excel(l["titular"]), _seguro_excel(l["detalle"]), int(_d(l["valor"])),
                   _seguro_excel(l["uso"]), l["categoria_texto"], "Sí" if l["incluida"] else "No",
                   f"{float(l['confianza']):.0%}", "Sí" if l["resaltada"] else ""])
        lo.cell(row=lo.max_row, column=7).number_format = '#,##0'
    for col, ancho in zip("ABCDEFGHIJKL", (22, 8, 6, 28, 22, 50, 16, 32, 32, 9, 12, 12)):
        lo.column_dimensions[col].width = ancho

    og = wb.create_sheet("De dónde sale cada peso")
    og.append(["Campo de la declaración", "Línea", "Documento", "Página", "Fila", "Valor", "Descripción"])
    for celda in og[1]:
        celda.font = negrita
        celda.fill = fondo
    for o in r.get("origen") or []:
        og.append([o["campo"], _seguro_excel(o["linea"]), _seguro_excel(o["documento"]), o["pagina"], o["fila"],
                   int(_d(o["valor"])), _seguro_excel(o.get("descripcion", ""))])
        og.cell(row=og.max_row, column=6).number_format = '#,##0'
    for col, ancho in zip("ABCDEFG", (26, 40, 26, 8, 6, 16, 40)):
        og.column_dimensions[col].width = ancho

    av = wb.create_sheet("Avisos")
    av.append(["Aviso"])
    av["A1"].font = negrita
    for m in (r.get("marcas") or []):
        av.append([_seguro_excel(m["texto"])])
    for a in (r.get("avisos") or []):
        av.append([_seguro_excel(a)])
    av.column_dimensions["A"].width = 120
    for fila in av.iter_rows(min_row=2):
        fila[0].alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    blindar_libro(wb)
    wb.save(buf)
    return buf.getvalue()
